"""원본 작성 일보를 날짜별 저장본으로 보존하는 일회성 가져오기."""
from pathlib import Path
from datetime import date, datetime
import hashlib
import json
import re
import shutil
import openpyxl
from night_shift import daily_report as dr

ROOT = Path(__file__).resolve().parent
SOURCE = Path(r'C:\Users\LJH\Desktop\업무\26년 근태관리 일보(7월~12월).xlsx')
STORE = ROOT / 'history'
wb = openpyxl.load_workbook(SOURCE, data_only=True)
pending = []
for ws in wb:
    match = re.fullmatch(r'(\d{2})\.(\d{2})\.(\d{2})', ws.title.strip())
    if not match:
        raise ValueError(f'확인되지 않은 시트: {ws.title}')
    y, m, d = map(int, match.groups())
    day = date(2000+y, m, d)
    report = dr.build(day, dr.entries_for(STORE, day), dr.load_roster(STORE))
    rows = []
    for row in report['rows']:
        line = row['행']
        to, accident, note = [ws.cell(line, c).value for c in (7, 8, 10)]
        if not isinstance(to, (int, float)) or not isinstance(accident, (int, float)):
            raise ValueError(f'{ws.title} {line}: 숫자 확인 필요')
        assert int(to) == to and int(accident) == accident
        rows.append({**row, 'T/O': int(to), '사고': int(accident), '비고': note or ''})
    writer_raw = ws['M6'].value or ''
    writer = re.sub(r'\s*\(인\)\s*$', '', writer_raw).replace(' ', '')
    pending.append((day, ws, rows, writer, writer_raw))
assert len({p[0] for p in pending}) == len(pending)
backup = ROOT / 'backup' / ('daily_report_original_' + datetime.now().strftime('%Y%m%d_%H%M%S'))
backup.mkdir(parents=True)
shutil.copy2(SOURCE, backup / SOURCE.name)
if (STORE / dr.REPORT_DIR).exists():
    shutil.copytree(STORE / dr.REPORT_DIR, backup / dr.REPORT_DIR)
existing = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (STORE / dr.REPORT_DIR).glob('*.json')}
for day, ws, rows, writer, writer_raw in pending:
    data = dr.save_report(STORE, day, rows, writer)
    data['source_file'] = SOURCE.name
    data['source_sheet'] = ws.title
    data['source_writer_raw'] = writer_raw
    data['source_cells'] = {c.coordinate: c.value for row in ws for c in row if c.value is not None}
    path = STORE / dr.REPORT_DIR / f'{day.isoformat()}.json'
    path.write_text(json.dumps(data, ensure_ascii=False, indent=1, default=str), encoding='utf-8')
    saved = dr.load_saved(STORE, day)
    assert saved['writer'] == writer
    for row in rows:
        original = saved['rows'][str(row['행'])]
        assert all(original[k] == row[k] for k in ('T/O', '사고', '비고'))
    assert saved['source_cells'] == data['source_cells']
for name, digest in existing.items():
    if date.fromisoformat(name[:-5]) not in {p[0] for p in pending}:
        assert hashlib.sha256((STORE / dr.REPORT_DIR / name).read_bytes()).hexdigest() == digest
print(json.dumps({'imported': len(pending), 'from': min(p[0] for p in pending).isoformat(),
                  'through': max(p[0] for p in pending).isoformat(), 'verified_rows': len(pending)*6,
                  'backup': str(backup)}, ensure_ascii=False))
