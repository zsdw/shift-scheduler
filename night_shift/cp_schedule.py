"""CP 근무자 월간 근무표 (2026-10-05 사용자 요청).

- 칸 값 '주1'·'야2' = 그날 주간 CP1 / 야간 CP2. '연차' 등 다른 글자는 그날 근무 없음.
- 날짜를 고르면 그날 CP1·CP2·CP3 이름이 기본값으로 정해지고, 야간 현장 양식 이미지와 주간 화면에 표시된다.
- 야간 시간별 CP 칸은 매일 같은 패턴(사용자 확정): 18시 CP1.CP3 → 19시 CP1.CP2 → 20시 CP2.CP3 반복,
  07시 셔틀은 CP3. 근무표 계산(Solver)에는 넣지 않는다.
- 다음 달부터는 원본 엑셀 업로드로 읽는다(parse_excel). 저장 위치: history/cp_schedule.json.
"""
from __future__ import annotations

import io
import json
import re
from datetime import date
from pathlib import Path

CP_FILE = "cp_schedule.json"
HOURLY = ((1, 3), (1, 2), (2, 3))       # 18시부터 3시간마다 반복 (현장 양식 그대로)
SHUTTLE_07 = 3
_CELL = re.compile(r"^\s*(주|야)\s*([123])\s*$")


def hourly_label(slot: int) -> str:
    a, b = HOURLY[slot % 3]
    return f"CP{a}.CP{b}"


NIGHT_HOURS = [str(h) for h in (18, 19, 20, 21, 22, 23, 0, 1, 2, 3, 4, 5, 6, 7)]   # 참조 자료의 시간 키
DAY_HOURS = [str(h) for h in range(8, 18)]


def image_override(store, day, shift):
    path = Path(store) / "cp_image_overrides.json"
    if not path.exists():
        return None
    return json.loads(path.read_text(encoding="utf-8")).get(day.isoformat(), {}).get(shift)


def save_image_override(store, day, shift, rows, names, filename):
    """확인 완료한 A조 이미지의 CP만 저장. 보안근무 기록은 변경하지 않는다."""
    hours = NIGHT_HOURS if shift == "night" else DAY_HOURS
    cells = {str(int(str(r['시간']).split(':')[0]) % 24): r for r in rows}
    if not names or not all(h in cells and 'CP' in cells[h] for h in hours):
        return False
    def cp_text(value):
        return '.'.join(re.findall(r'CP[123]', str(value).upper()))
    value = {'names': {str(k).replace('CP', ''): v for k, v in names.items()},
             'hours': [cp_text(cells[h]['CP']) for h in hours],
             'shuttle': cp_text(cells[hours[-1]].get('셔틀', '')), 'file': filename}
    path = Path(store) / 'cp_image_overrides.json'
    data = json.loads(path.read_text(encoding='utf-8')) if path.exists() else {}
    if data.get(day.isoformat(), {}).get(shift) == value:
        return False
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        import shutil
        from datetime import datetime
        backup = Path(store) / '.trash'
        backup.mkdir(exist_ok=True)
        shutil.copy2(path, backup / ('cp_image_overrides_' + datetime.now().strftime('%Y%m%d%H%M%S%f') + '.json'))
    data.setdefault(day.isoformat(), {})[shift] = value
    tmp = path.with_suffix('.tmp')
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding='utf-8')
    tmp.replace(path)
    return True


def standard_pattern(store: Path, shift: str, numbers, day: date | None = None, keep=None) -> dict | None:
    """엑셀에 실제로 있는 같은 인원 CP 근무 형태를 표준으로 (2026-10-11 사용자 확정: CP 연차로 2명인 날 등).

    올린 편성표 참조(모든 조)에서 CP 인원 수가 같은 날 중 기준일에 가장 가까운(이전 우선) 시트의 시간별 CP 칸.
    keep(날짜) → False인 날짜는 건너뜀 (휴일 주간 양식은 휴일 시트끼리, 2026-10-12).
    반환 {'hours': [칸 글자…], 'shuttle': 셔틀 칸 CP 글자(야간 07시·주간 17시) 또는 ''}. 번호는 오늘 CP 번호로 바꿔 준다.
    """
    override = image_override(store, day, shift) if day is not None else None
    if override is not None:
        return {'hours': override['hours'], 'shuttle': override['shuttle']}
    path = Path(store) / "order_source" / "uploaded_sources.json"
    try:
        sources = json.loads(path.read_text(encoding="utf-8")).get("sources", {})
    except (OSError, ValueError):
        return None
    want = sorted(str(n) for n in numbers)
    hours = NIGHT_HOURS if shift == "night" else DAY_HOURS
    found = []
    for iso, shifts in sources.items():
        src = (shifts or {}).get(shift) or {}
        have = sorted(str(k) for k in (src.get("cp") or {}))
        cells = src.get("cells") or {}
        if len(have) != len(want) or not all(str((cells.get(h) or {}).get("CP") or "").strip() for h in hours):
            continue
        d = date.fromisoformat(iso)
        if keep is not None and not keep(d):
            continue
        rank = (0 if day is None or d <= day else 1, abs((day - d).days) if day else -d.toordinal(), src.get("team") != "A")
        found.append((rank, have, cells))
    if not found:
        return None
    _, have, cells = min(found, key=lambda x: x[0])
    relabel = dict(zip(have, want))
    def fix(text):
        return re.sub(r"CP\s*([123])", lambda m: f"CP{relabel.get(m[1], m[1])}", str(text or "").upper().replace(" ", "")
                      .replace(",", "."))
    last = cells.get(hours[-1]) or {}
    shuttle = next((fix(v) for k, v in last.items() if k in ("셔틀", "순회") and "CP" in str(v).upper()), "")
    return {"hours": [fix(cells[h]["CP"]) for h in hours], "shuttle": shuttle}


def hourly(store: Path, day: date, shift: str, on: dict) -> dict:
    """그날 CP 시간별 칸·셔틀 (야간 14칸·주간 10칸). 엑셀에 같은 인원 형태가 있으면 그것, 없으면 기본 패턴."""
    if not on:
        return {"hours": [""] * (14 if shift == "night" else 10), "shuttle": ""}
    found = standard_pattern(store, shift, list(on), day)
    if found:
        return found
    if shift == "night":
        return {"hours": [".".join(f"CP{n}" for n in HOURLY[t % 3] if str(n) in on) for t in range(14)],
                "shuttle": f"CP{SHUTTLE_07}" if str(SHUTTLE_07) in on else ""}
    return {"hours": [""] * 10, "shuttle": ""}


def short(name: str) -> str:
    """'이덕영' → '덕영' (현장 양식 표기)."""
    return name[-2:] if len(name) >= 3 else name


def parse_rows(year: int, month: int, rows: dict[str, dict[int, str]]) -> dict[str, dict]:
    """{이름: {일: '주1'|'야2'|...}} → {'YYYY-MM-DD': {'day': {'1': 이름}, 'night': {...}}}."""
    out: dict[str, dict] = {}
    for name, cells in rows.items():
        for day_number, value in cells.items():
            m = _CELL.match(str(value or ""))
            if not m:
                continue
            shift = "day" if m[1] == "주" else "night"
            key = date(year, month, int(day_number)).isoformat()
            slot = out.setdefault(key, {"day": {}, "night": {}})[shift]
            if m[2] in slot and slot[m[2]] != name:
                raise ValueError(f"{key} {'주간' if shift == 'day' else '야간'} CP{m[2]}가 두 명입니다: {slot[m[2]]}, {name}")
            slot[m[2]] = name
    return out


def parse_excel(data: bytes, year: int | None = None) -> dict[str, dict]:
    """원본 엑셀(제목 'CP 10월 근무', 날짜 줄 1~31, 이름 줄마다 주N/야N)에서 읽는다."""
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(data), data_only=True)
    found: dict[str, dict] = {}
    recognized = False
    for sheet in book.worksheets:
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
        month = next((int(m[1]) for r in rows for v in r if v
                      and (m := re.search(r"CP\s*(\d{1,2})\s*월", str(v)))), None)
        if month is None:
            continue
        header = None
        for i, r in enumerate(rows):
            numbers = {j: int(v) for j, v in enumerate(r) if isinstance(v, (int, float)) and 1 <= v <= 31
                       and float(v).is_integer()}
            if len(numbers) >= 28 and list(numbers.values())[:3] == [1, 2, 3]:
                header = (i, numbers)
                break
        if header is None:
            continue
        index, columns = header
        people: dict[str, dict[int, str]] = {}
        for r in rows[index + 1:]:
            name = next((str(v).strip() for j, v in enumerate(r) if j < min(columns)
                         and v and re.fullmatch(r"[가-힣]{2,4}", str(v).strip())
                         and not str(v).strip().endswith("조")), None)
            if name:
                people[name] = {d: r[j] for j, d in columns.items() if j < len(r)}
        found.update(parse_rows(year or date.today().year, month, people))
        recognized = recognized or bool(people)
    book.close()
    if not found and not recognized:
        raise ValueError("CP 근무표(제목 'CP N월', 날짜 1~31 줄, 주N/야N 칸)를 찾지 못했습니다.")
    return found


def is_cp_workbook(data: bytes) -> bool:
    try:
        parse_excel(data)
        return True
    except Exception:
        return False


def load(store: Path) -> dict[str, dict]:
    path = Path(store) / CP_FILE
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def save(store: Path, entries: dict[str, dict]) -> Path:
    """같은 날짜는 새 값으로 덮어쓴다 (월 단위로 다시 올려도 됨)."""
    merged = {**load(store), **entries}
    path = Path(store) / CP_FILE
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(dict(sorted(merged.items())), ensure_ascii=False, indent=1), encoding="utf-8")
    return path


def save_uploaded(store: Path,data: bytes,filename: str,year: int):
    """업로드에 포함된 월을 통째로 교체하여 연차·빈 날짜의 이전 CP 배정도 지운다."""
    from openpyxl import load_workbook
    entries=parse_excel(data,year)
    book=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
    try:
        months={f'{year}-{int(m[1]):02d}' for ws in book for row in ws.iter_rows(values_only=True)
                for value in row if value and (m:=re.search(r'CP\s*(\d{1,2})\s*월',str(value)))}
    finally:book.close()
    return replace_uploaded(store,entries,months,data,filename)


def replace_uploaded(store,entries,months,data,filename):
    merged={key:value for key,value in load(store).items() if key[:7] not in months}
    merged.update(entries)
    store=Path(store);store.mkdir(parents=True,exist_ok=True)
    path=store/CP_FILE;tmp=path.with_suffix('.tmp')
    tmp.write_text(json.dumps(dict(sorted(merged.items())),ensure_ascii=False,indent=1),encoding='utf-8');tmp.replace(path)
    marker=store/'cp_upload_source.json'
    try:old=json.loads(marker.read_text(encoding='utf-8')).get('months',[])
    except (OSError,ValueError):old=[]
    marker.write_text(json.dumps({'months':sorted(set(old)|months),'file':filename},ensure_ascii=False),encoding='utf-8')
    from .source_attachment import stash
    stash(store,data,filename)
    return len(entries)


def covers(store,day):
    try:return day.strftime('%Y-%m') in json.loads((Path(store)/'cp_upload_source.json').read_text(encoding='utf-8'))['months']
    except (OSError,ValueError,KeyError):return False


def for_day(store: Path, day: date, shift: str) -> dict[str, str]:
    """{'1': 이름, '2': 이름, '3': 이름} (없는 번호는 빠짐)."""
    store=Path(store)
    names=load(store).get(day.isoformat(), {}).get(shift, {})
    if not covers(store,day):
        # CP 전용 업로드가 없는 날짜는 최신 원본의 명단을 사용한다.
        # reference_sources는 for_day를 부르므로 가공 전 참조 JSON을 직접 읽는다.
        path=store/'order_source'/'uploaded_sources.json'
        if path.exists():
            source=json.loads(path.read_text(encoding='utf-8'))['sources'].get(day.isoformat(),{}).get(shift)
        else:
            from .order_source import latest_file,read_day
            original=latest_file()
            source=read_day(original,day,store).get(shift) if original else None
        if source is not None and 'cp' in source:names=source['cp']
    override = image_override(store, day, shift)
    if override is not None:
        names = override['names']
    from . import daily_report as dr
    roster=dr.active(dr.load_roster(store),day)
    known=[p['이름'] for p in roster if p['직책']=='CP']
    absent={e['name'] for e in dr.entries_for(store,day) if any(kind=='absent' for kind,_ in dr.classify_entry(e))}
    normalized={}
    for number,name in names.items():
        matched=[full for full in known if full==name or full.endswith(name)]
        full=matched[0] if len(matched)==1 else name
        if not any(full==off or full.endswith(off) or off.endswith(full) for off in absent):normalized[number]=full
    return normalized


def caption(names: dict[str, str]) -> str:
    return " · ".join(f"CP{n} {names[n]}" for n in ("1", "2", "3") if n in names)
