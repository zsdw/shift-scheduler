"""근태관리 일보 자동 작성 (2026-10-06 사용자 요청).

연차일정(앱에 올린 history/leave_schedule.txt) + 명단(이름 → 소속) + 6일 교대 + 공휴일로
일보의 T/O·사고·현인원·사고내용(비고)을 채운다. 과거 일보 7~9월 대조로 확인한 기준:

- 사고 = 결원 + 그날 출근하지 않은 사람(연차·휴가·병가·결근·대휴·휴무·예비군·훈련·경조사).
- 출근은 했으면 사고가 아니다: 반차·교육·검진·입사·퇴사는 비고에만.
- 지원·대근: 원래 소속에서 출근으로 센다(사고 아님), 비고에 '야간지원 1명(이름)' 등. 받는 조는 건드리지 않는다.
- 주말·공휴일: 주간고정·안내·주차는 정원 전체가 사고, 지원·대근 나온 사람만 현인원.
- 휴무조는 정원 전체가 사고.
- 조 교대: A조 달력(주·주·야·야·휴·휴) 기준 — A 주간이면 야간 B·휴무 C, A 야간이면 주간 C·휴무 B,
  A 휴무면 주간 B·야간 C (근무명령서 7~9월 6일 주기와 같음).
"""
from __future__ import annotations

import json
import re
from datetime import date, timedelta
from pathlib import Path

from .calendar_info import is_off_day, shift_kind

ROSTER_FILE = "daily_roster.json"
TEMPLATE = Path(__file__).resolve().parent / "templates" / "daily_report.xlsx"
GROUPS = ("주간고정", "A조", "B조", "C조", "안내", "주차")
TO = {"주간고정": 11, "A조": 10, "B조": 10, "C조": 10, "안내": 5, "주차": 3}

# 2026-10-06 사용자 확인 명단 (모든 조 풀 T/O). 시작·종료는 그 날짜 포함.
DEFAULT_ROSTER = [
    *({"이름": n, "소속": "주간고정", "직책": j} for n, j in (
        ("최찬우", "실장"), ("김성덕", "팀장"), ("김영복", "부팀장"), ("진기운", "통제D"), ("최정식", "통제E"),
        ("김수환", "행정"), ("김도운", "대원"), ("이준형", "대원"), ("조용주", "대원"), ("홍성욱", "대원"))),
    {"이름": "윤여진", "소속": "주간고정", "직책": "대원", "시작": "2026-09-28"},
    {"이름": "박인철", "소속": "A조", "직책": "통제A", "종료": "2026-09-15"},   # 9/15 퇴사(그날 야간까지 근무)
    {"이름": "백원석", "소속": "A조", "직책": "통제A", "시작": "2026-09-21"},
    *({"이름": n, "소속": "A조", "직책": "대원"} for n in ("이성호", "김광희", "한만년", "이준호", "유원종", "김제욱")),
    *({"이름": n, "소속": "A조", "직책": "CP"} for n in ("권혁준", "이덕영", "안철룡")),
    {"이름": "김현", "소속": "B조", "직책": "통제B"},
    *({"이름": n, "소속": "B조", "직책": "대원"} for n in ("이경빈", "김광현", "박재우", "정현욱", "강병종", "박나현")),
    *({"이름": n, "소속": "B조", "직책": "CP"} for n in ("서용준", "김준하", "전안성")),
    {"이름": "이채용", "소속": "C조", "직책": "통제C"},
    *({"이름": n, "소속": "C조", "직책": "대원"} for n in ("허승회", "김상민", "강석규", "이세명")),
    {"이름": "박정호", "소속": "C조", "직책": "대원", "종료": "2026-09-18"},
    {"이름": "김대현", "소속": "C조", "직책": "대원", "시작": "2026-09-15"},
    {"이름": "이지훈", "소속": "C조", "직책": "대원", "시작": "2026-10-06"},
    *({"이름": n, "소속": "C조", "직책": "CP"} for n in ("이석산", "조백호", "송승현")),
    {"이름": "김주희", "소속": "안내", "직책": "통제I"},
    *({"이름": n, "소속": "안내", "직책": "대원"} for n in ("김은예", "유민", "최지은", "박지연")),
    *({"이름": n, "소속": "주차", "직책": "대원"} for n in ("이대성", "이정민", "조호준")),
]


def load_roster(store: Path) -> list[dict]:
    try:
        data = json.loads((Path(store) / ROSTER_FILE).read_text(encoding="utf-8"))
        return data["roster"] if isinstance(data, dict) else data
    except (OSError, ValueError, KeyError):
        return [dict(r) for r in DEFAULT_ROSTER]


def save_roster(store: Path, roster: list[dict]) -> None:
    Path(store).mkdir(parents=True, exist_ok=True)
    (Path(store) / ROSTER_FILE).write_text(json.dumps({"roster": roster}, ensure_ascii=False, indent=1),
                                           encoding="utf-8")


def active(roster: list[dict], day: date) -> list[dict]:
    out = []
    for r in roster:
        start, end = str(r.get("시작") or ""), str(r.get("종료") or "")
        if (not start or day >= date.fromisoformat(start[:10])) and (not end or day <= date.fromisoformat(end[:10])):
            out.append(r)
    return out


def teams_for(day: date) -> dict[str, str]:
    """{'주간조': 'A조', '야간조': 'B조', '휴무조': 'C조'} — A조 달력 기준 6일 교대."""
    a = shift_kind(day)
    day_team, night_team, off_team = {"주": ("A", "B", "C"), "야": ("C", "A", "B"), "휴": ("B", "C", "A")}[a]
    return {"주간조": day_team + "조", "야간조": night_team + "조", "휴무조": off_team + "조"}


def classify(kind: str) -> list[tuple[str, str]]:
    """연차일정 내용 → [(구분, 비고 이름)]. 구분: 'absent'(사고) | 'present'(출근, 비고만) | 'skip' | 'unknown'."""
    k = kind.replace(" ", "")
    out = []
    if k.startswith("출근:"):
        return [("present", kind)]
    if k.startswith("미출근:"):
        return [("absent", kind)]
    if "지원" in k:
        label = "야간지원" if "야간" in k else "주간지원" if "주간" in k else "교대조 지원"
        return [("present", label)]
    if "대근" in k:
        return [("present", "대근")]
    if "반차" in k or "반야" in k:                # 원본 일보 표기: 오전반차·오후반차, 야간은 전반야 반차·후반야 반차
        label = ("오전반차" if "오전" in k else "오후반차" if "오후" in k else "전반야 반차" if "전반" in k
                 else "반차")                      # 원본: 후반야 반차는 '반차 1명(허승회)'
        return [("present", label)]
    if "대휴" in k:
        out.append(("absent", "대휴"))
    elif "연차" in k:
        out.append(("absent", "연차"))
    elif "휴가" in k:                             # 원본 9월 일보: 하계휴가도 '휴가 1명(이준호)'
        out.append(("absent", "휴가"))
    elif "병가" in k:
        out.append(("absent", "병가"))
    elif "결근" in k:
        out.append(("absent", "결근"))
    elif "휴무" in k:
        out.append(("absent", "휴무"))
    elif "예비군" in k or "민방위" in k or "훈련" in k:   # 원본: '훈련 2명(이경빈, 박재우)'
        out.append(("absent", "훈련"))
    elif "경조" in k:
        out.append(("absent", "경조사"))
    if "퇴사" in k:
        out.append(("present", "퇴사"))
    if "입사" in k:                               # 원본: '입사 및 주간교육 1명(백원석)'
        out.append(("present", "입사 및 주간교육" if "교육" in k and "주간" in k else "입사"))
    if not out:
        if "신임경비교육" in k:
            out.append(("present", "신임경비교육"))
        elif "교육" in k:
            out.append(("present", "주간교육" if "주간" in k else "교육"))
        elif "검진" in k:
            out.append(("present", "검진"))
        elif "편입" in k or "명단" in k or "전입" in k or "전출" in k:
            out.append(("skip", ""))
        else:
            out.append(("unknown", kind))
    return out


# 원본 일보 사고내용 순서 (7~10월 대조): 결원 → 연차 → 휴가 → 대휴·휴무 → 훈련·경조사 → 반차 → 교육 → 지원·대근 → 입사 → 퇴사
NOTE_ORDER = ("연차", "하계휴가", "휴가", "병가", "결근", "대휴", "휴무", "예비군 훈련", "훈련", "경조사",
              "오전반차", "오후반차", "전반야 반차", "후반야 반차", "반차",
              "신임경비교육", "주간교육", "교육", "검진",
              "야간지원", "주간지원", "교대조 지원", "대근", "입사 및 주간교육", "입사", "퇴사")
# 원본 표기에서 이름 뒤 호칭을 붙인 사람 (예: '연차 1명(최찬우 이사님)')
HONORIFIC = {"최찬우": "이사님"}


def _note(items: list[tuple[str, str]], vacancy: int = 0, seat: dict | None = None) -> str:
    """[(비고 이름, 사람)] → '결원 1명, 연차 2명(최정식, 김수환), 오후반차 1명(박인철)' (원본 일보 기재 양식).
    seat: 이름 → 명단 순서 (원본은 이름을 직책·명단 순서로 적음: 최정식 파트장 → 김수환 행정)."""
    parts = [f"결원 {vacancy}명"] if vacancy else []
    order: dict[str, list[str]] = {}
    seat = seat or {}
    for label, who in sorted(items, key=lambda x: seat.get(x[1], len(seat))):
        name = f"{who} {HONORIFIC[who]}" if who in HONORIFIC else who
        if name not in order.setdefault(label, []):
            order[label].append(name)
    rank = {label: i for i, label in enumerate(NOTE_ORDER)}
    for label in sorted(order, key=lambda x: rank.get(x, len(rank))):
        parts.append(f"{label} {len(order[label])}명({', '.join(order[label])})")
    return ", ".join(parts)


def _counts(items: list[tuple[str, str]]) -> list[tuple[str, int]]:
    """[(종류, 사람)] → [(종류, 인원)] 원본 순서대로."""
    rank = {label: i for i, label in enumerate(NOTE_ORDER)}
    tally: dict[str, set] = {}
    for label, who in items:
        tally.setdefault(label, set()).add(who)
    return [(k, len(tally[k])) for k in sorted(tally, key=lambda x: rank.get(x, len(rank)))]


def spaced(name: str) -> str:
    """일보 작성자 칸 표기: '박 인 철 (인)', '김  현 (인)'."""
    return (" ".join(name) if len(name) >= 3 else "  ".join(name)) + " (인)"


def classify_entry(entry: dict) -> list[tuple[str, str]]:
    result = classify(entry["kind"])
    if entry.get("manual"):
        presence = entry.get("presence", "자동")
        if presence in ("출근", "미출근"):
            result = [("present" if presence == "출근" else "absent", f"{presence}:{entry['kind']}")]
        elif entry["kind"] == "출근":
            result = []
        if entry.get("note"):
            note = re.sub(r"[(),\n\r]", " ", entry["note"])
            result = [(what, f"{label} · {note}") for what, label in result]
    return result


def build(day: date, entries: list[dict], roster: list[dict], prev_entries: list[dict] | None = None) -> dict:
    """하루 일보. entries: leave_schedule.parse(...)[day] 형태 [{'team','name','kind'}].
    prev_entries: 전날 연차일정 (작성자인 전날 야간조 조장이 빠졌는지 확인용)."""
    members = active(roster, day)
    group_of = {r["이름"]: r["소속"] for r in members}
    seat = {r["이름"]: i for i, r in enumerate(roster)}
    teams = teams_for(day)
    off_day = is_off_day(day)
    working = {"주간고정": not off_day, "안내": not off_day, "주차": not off_day,
               teams["주간조"]: True, teams["야간조"]: True, teams["휴무조"]: False}
    absent = {g: [] for g in GROUPS}       # [(라벨, 이름)]
    present = {g: [] for g in GROUPS}      # 출근(비고만)
    checks = []
    for e in entries:
        name = e["name"]
        group = group_of.get(name)
        if group is None:
            checks.append(f"명단에 없는 이름: {e.get('team', '')} {name} - {e['kind']}")
            continue
        for what, label in classify_entry(e):
            if what == "absent":
                if working[group]:
                    absent[group].append((label, name))
                else:
                    checks.append(f"{name}({group})은 이날 근무가 아닌데 '{e['kind']}' 기록이 있습니다.")
            elif what == "present":
                present[group].append((label, name))
            elif what == "unknown":
                checks.append(f"분류 못 함(사고로 세지 않음): {name} - {e['kind']}")
    rows = []

    def row(label, group, slot=None):
        to = TO[group]
        vacancy = max(0, to - sum(1 for r in members if r["소속"] == group))
        if working[group]:
            accident = vacancy + len(absent[group])
            note = _note(absent[group] + present[group], vacancy, seat)
            counted = ([f"결원 {vacancy}"] if vacancy else []) + [
                f"{label_} {n}" for label_, n in _counts(absent[group])]
            reason = (" + ".join(counted) + f" = {accident}") if counted else "빠진 사람 없음 = 0"
            if present[group]:
                reason += " (반차·교육·지원 등은 출근)"
        else:
            # 쉬는 날 나온 사람(지원·대근·교육)은 출근 — 원본: 휴무조 주간교육 1명이면 사고 9 (7/1~3, 8/26)
            came = [x for x in present[group] if x[0].startswith("출근:") or x[0].split(" · ")[0] in (
                "야간지원", "주간지원", "교대조 지원", "대근", "교육",
                "주간교육", "신임경비교육", "입사 및 주간교육")]
            accident = to - len({who for _, who in came})
            note = _note(present[group], vacancy, seat)
            why = "휴무조" if group == teams["휴무조"] else "주말·공휴일 휴무"
            reason = f"{why} 전원 {to}" + (f" − 출근 {to - accident}(지원·대근·교육) = {accident}" if accident != to else "")
        rows.append({"구분": label, "조": group if slot else "", "T/O": to, "사고": accident, "비고": note,
                     "근거": reason, "행": {"주간고정": 12, "안내": 20, "주차": 21}.get(group, slot)})
    row("주간고정", "주간고정")
    row(f"주간조 ({teams['주간조']})", teams["주간조"], 13)
    row(f"야간조 ({teams['야간조']})", teams["야간조"], 15)
    row(f"휴무조 ({teams['휴무조']})", teams["휴무조"], 16)
    row("안내", "안내")
    row("주차", "주차")
    # 작성자 = 전날 밤 야간조 조장 (2026-10-06 사용자 확인·원본 일보 대조: 그 밤 근무하며 다음날 일보를 씀)
    prev = day - timedelta(days=1)
    night = teams_for(prev)["야간조"]
    writer = next((r["이름"] for r in active(roster, prev)
                   if r["소속"] == night and str(r.get("직책", "")).startswith("통제")), "")
    if writer and any(e["name"] == writer and "absent" in [w for w, _ in classify_entry(e)]
                      for e in (prev_entries or [])):
        checks.append(f"작성자: 전날({prev:%m/%d}) 야간조 조장 {writer}이(가) 빠진 날이라 대신 쓴 사람(최정식·선임 등)으로 고치세요.")
        writer = ""
    return {"day": day, "teams": teams, "off_day": off_day, "rows": rows, "checks": checks, "writer": writer,
            "writer_from": f"{prev:%m/%d} 야간조({night}) 조장"}


def to_excel(report: dict, rows: list[dict] | None = None) -> bytes:
    """원본 일보 양식(템플릿)에 값만 채운 엑셀. rows: 화면에서 고친 표(없으면 계산값)."""
    from io import BytesIO
    import openpyxl
    day = report["day"]
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb.active
    ws.title = f"{day:%y.%m.%d}"
    ws["G6"] = f"{day.year}년 {day.month:02d}월 {day.day:02d}일({'월화수목금토일'[day.weekday()]})"
    if report["writer"]:
        ws["M6"] = spaced(report["writer"])
    ws["G11"], ws["H11"], ws["J11"] = 0, 0, None
    for r in rows or report["rows"]:
        line = int(r["행"])
        ws.cell(line, 7).value = int(r["T/O"])
        ws.cell(line, 8).value = int(r["사고"])
        ws.cell(line, 10).value = (r.get("비고") or None)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def team_members(day: date, roster: list[dict]) -> dict[str, dict]:
    """그날 명단의 조별 구성 {'A조': {'조장': 이름, '대원': [...], 'CP': [...]}, ...} (A·B·C조·주간고정)."""
    out = {}
    for group in ("주간고정", "A조", "B조", "C조"):
        people = [r for r in active(roster, day) if r["소속"] == group]
        lead = "" if group == "주간고정" else next(          # 주간고정은 조장 없이 직책과 함께 전원
            (r["이름"] for r in people if str(r.get("직책", "")).startswith("통제")), "")
        out[group] = {
            "조장": lead,
            "대원": [r["이름"] + (f"({r['직책']})" if group == "주간고정" and r.get("직책") not in ("대원", "", None) else "")
                   for r in people if r["이름"] != lead and r.get("직책") != "CP"],
            "CP": [r["이름"] for r in people if r.get("직책") == "CP"],
            "인원": len(people), "T/O": TO[group]}
    return out


def sheet_values(report: dict, rows: list[dict] | None = None) -> dict[int, tuple]:
    """엑셀 행 번호 → (T/O, 사고, 현인원, 비고). 원본 수식(현인원 = T/O−사고, 계·소계·합계)과 같은 계산."""
    v = {11: (0, 0, "")}
    for r in rows or report["rows"]:
        v[int(r["행"])] = (int(r["T/O"]), int(r["사고"]), r.get("비고") or "")
    s = lambda keys, i: sum(v[k][i] for k in keys)
    v[14] = (s((12, 13), 0), s((12, 13), 1), "")
    v[17] = (s((14, 15, 16), 0), s((14, 15, 16), 1), "")
    v[24] = (s((17, 20, 21), 0), s((17, 20, 21), 1), "")
    for line, to in ((18, 7), (19, 0), (22, 12), (23, 0)):   # 원본의 고정 줄 (합계에 안 들어감)
        v[line] = (to, None, "")
    return {k: (to, acc, to - (acc or 0), note) for k, (to, acc, note) in v.items()}


def sheet_html(report: dict, rows: list[dict] | None = None) -> str:
    """원본 엑셀 '근태관리 일보' 모양 그대로 (합친 칸·결재란·회색 머리 칸). 흰 바탕(다크모드 무관)."""
    from html import escape
    day = report["day"]
    v = sheet_values(report, rows)
    writer = spaced(report["writer"]) if report["writer"] else ""
    G = "background:#d9d9d9;"
    td = lambda text="", span=1, rows_=1, style="": (
        f"<td colspan='{span}' rowspan='{rows_}' style='{style}'>{text}</td>")

    def num(x):
        return "" if x is None else str(x)

    # 사고 계산 근거 (앱 화면만, 2026-10-06 사용자 요청): 사고내용 칸을 세로로 나눠 오른쪽에 짧게
    computed = {int(r["행"]): r for r in report["rows"]}
    reasons = {line_: r.get("근거", "") for line_, r in computed.items()}
    for r in rows or []:
        base = computed.get(int(r["행"]))
        if base and int(r["사고"]) != int(base["사고"]):
            reasons[int(r["행"])] = f"직접 고침 (계산값 {base['사고']})"
    reasons.update({14: "주간고정 + 주간조", 17: "계 + 야간조 + 휴무조", 24: "소계 + 안내 + 주차",
                    18: "고정 줄 (합계에 안 들어감)", 19: "고정 줄 (합계에 안 들어감)",
                    22: "고정 줄 (합계에 안 들어감)", 23: "고정 줄 (합계에 안 들어감)"})

    def line(r, label_cells, bold=False):
        to, acc, now, note = v[r]
        b = "font-weight:700;" if bold else ""
        gray = G if bold or r in (14, 17) else ""
        return ("<tr>" + label_cells + td(num(to), style=b + gray) + td(num(acc), style=b + gray)
                + td(num(now), style=b + gray) + td(escape(note), 3, style="text-align:left;padding-left:6px;" + gray)
                + td(escape(reasons.get(r, "")), 3,
                     style="text-align:left;padding-left:6px;font-size:11px;color:#4b5563;" + gray) + "</tr>")
    date_text = f"{day.year}년 {day.month:02d}월 {day.day:02d}일({'월화수목금토일'[day.weekday()]})"
    widths = [34, 30, 40, 40, 58, 58, 82, 82, 82, 62, 62, 62, 92, 62, 62]
    h = ["<style>table.drs td{border:1px solid #6b7280;height:30px;padding:2px;color:#111827}</style>"
         "<div style='overflow-x:auto'><table class='drs' style='border-collapse:collapse;table-layout:fixed;"
         "min-width:760px;width:100%;background:#fff;color:#111827;font-size:13px;text-align:center'><colgroup>"
         + "".join(f"<col style='width:{w}px'>" for w in widths) + "</colgroup><tbody>"]
    h.append("<tr>" + td("H D S", 5) + td("", 10, 4, "font-size:20px;font-weight:700") + "</tr>")
    h.append("<tr>" + td("결<br>재", 1, 3) + td("담 당", 2) + td("팀 장", 2) + "</tr>")
    h.append("<tr>" + td("", 2, 2, "height:44px") + td("", 2, 2) + "</tr><tr></tr>")
    h.append("<tr>" + td("사 업 자 명", 6, style=G) + td("에이치디에스 주식회사", 4) + td("부 서", 2, style=G)
             + td("보안사업MU FM사업팀(판교)", 3) + "</tr>")
    h.append("<tr>" + td("일 자", 6, style=G) + td(date_text, 4) + td("", 2, style=G) + td(escape(writer), 3) + "</tr>")
    h.append("<tr><td colspan='15' style='border:0;height:14px'></td></tr>")
    h.append("<tr>" + td("구 분", 6, style=G) + td("T/O", style=G) + td("사고", style=G) + td("현인원", style=G)
             + td("사고내용(비고)", 3, style=G) + td("사고 계산 근거", 3, style=G + "font-size:12px") + "</tr>")
    # 원본 11행 '관리/시스템'(항상 0)은 화면에서 뺀다 (2026-10-06 사용자 요청)
    h.append(line(12, td("보안", 2, 6) + td("주간조", 2, 3) + td("주간고정", 2)))
    h.append(line(13, td("주간조", 2)))
    h.append(line(14, td("계", 2, style=G)))
    h.append(line(15, td("야간조", 4)))
    h.append(line(16, td("휴무조", 4)))
    h.append(line(17, td("소계", 4, style=G)))
    h.append(line(20, td("안내", 2) + td("주간조", 4)))
    h.append(line(21, td("주차", 2) + td("주간조", 4)))
    h.append(line(24, td("합 계", 6, style=G + "font-weight:700"), bold=True))
    return "".join(h) + "</tbody></table></div>"


# ---------------------------------------------------------------- 날짜별 저장 (2026-10-06 사용자 요청)
REPORT_DIR = "daily_reports"


def _report_path(store: Path, day: date) -> Path:
    return Path(store) / REPORT_DIR / f"{day.isoformat()}.json"


def load_saved(store: Path, day: date) -> dict | None:
    """저장한 일보 {'rows': {행: {'사고', '비고'}}, 'writer', 'saved_at'} 또는 None."""
    try:
        return json.loads(_report_path(store, day).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None


def _to_trash(store: Path, day: date) -> None:
    path = _report_path(store, day)
    if path.exists():
        from datetime import datetime
        trash = path.parent / ".trash"
        trash.mkdir(parents=True, exist_ok=True)
        path.replace(trash / f"{day.isoformat()}_{datetime.now():%Y%m%d_%H%M%S_%f}.json")


def save_report(store: Path, day: date, rows: list[dict], writer: str, computed: list[dict] | None = None) -> dict:
    """화면에서 확정한 사고·비고·작성자를 날짜별로 저장(기존 저장본은 휴지통에 보관). 근무명령서도 이 값을 쓴다.
    computed: 저장 당시 연차일정 계산값 (나중에 연차일정이 바뀌면 알려 주려고 같이 남김)."""
    from datetime import datetime
    calc = {int(r["행"]): r for r in computed or []}
    _to_trash(store, day)
    data = {"day": day.isoformat(), "saved_at": datetime.now().isoformat(timespec="seconds"), "writer": writer,
            "rows": {str(int(r["행"])): {"구분": r["구분"], "T/O": int(r["T/O"]), "사고": int(r["사고"]),
                                         "비고": r.get("비고") or "",
                                         "계산사고": int(calc[int(r["행"])]["사고"]) if int(r["행"]) in calc else None,
                                         "계산비고": calc[int(r["행"])]["비고"] if int(r["행"]) in calc else None}
                     for r in rows}}
    path = _report_path(store, day)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    tmp.replace(path)
    return data


def delete_saved(store: Path, day: date) -> None:
    """저장본을 휴지통으로 (계산값으로 되돌리기)."""
    _to_trash(store, day)


def apply_saved(report: dict, saved: dict | None) -> tuple[list[dict], list[str]]:
    """계산한 줄에 저장본 사고·비고를 덮어쓴 줄, 그리고 저장 뒤 계산값이 달라진 줄 안내."""
    if not saved:
        return [dict(r) for r in report["rows"]], []
    rows, changed = [], []
    for r in report["rows"]:
        s = saved.get("rows", {}).get(str(int(r["행"])))
        rows.append({**r, "사고": s["사고"], "비고": s["비고"]} if s else dict(r))
        if s and s.get("계산사고") is not None and (s["계산사고"] != r["사고"] or (s.get("계산비고") or "") != r["비고"]):
            changed.append(r["구분"])
    return rows, changed


ATTENDANCE_FIX_FILE = "attendance_fix.json"   # 날짜별 출근 상태 고침 (연차일정 원본 위에 덧씌움, 2026-10-07)
TEAM_LETTER = {"A조": "A", "B조": "B", "C조": "C", "주간고정": "D"}


def load_attendance_fixes(store: Path) -> dict:
    try:
        return json.loads((Path(store) / ATTENDANCE_FIX_FILE).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def save_attendance_fix(store: Path, day: date, name: str, status: str) -> None:
    """그 날짜 그 사람의 상태를 고친다. '출근' = 연차일정 항목을 그날만 없앤다. 연차일정 원본 txt는 그대로."""
    data = load_attendance_fixes(store)
    group = next((r["소속"] for r in load_roster(store) if r["이름"] == name), "")
    data.setdefault(day.isoformat(), {})[name] = {"kind": status, "team": TEAM_LETTER.get(group, "")}
    (Path(store) / ATTENDANCE_FIX_FILE).write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")


def entries_for(store: Path, day: date) -> list[dict]:
    from .leave_schedule import load
    entries = list(load(store, day.year).get(day, []))
    # 명단의 마지막 근무일(종료) 당일 = 일보 비고 '퇴사' (2026-10-12 사용자 확정). 이미 퇴사 일정이 있으면 그대로.
    for r in load_roster(store):
        if str(r.get("종료") or "")[:10] == day.isoformat() and not any(
                "퇴사" in e.get("kind", "") and (e["name"] == r["이름"] or r["이름"].endswith(e["name"])) for e in entries):
            entries.append({"name": r["이름"], "team": TEAM_LETTER.get(r.get("소속", ""), ""), "kind": "퇴사"})
    return entries


def short_name(name: str) -> str:
    return name[-2:] if len(name) >= 3 else name


def retired(store: Path, day: date) -> set[str]:
    """그 날짜에 이미 퇴사한 사람(명단 마지막 근무일 < 날짜)의 이름·짧은 이름 (근무 편성 명단에서 뺀다)."""
    out = set()
    for r in load_roster(store):
        end = str(r.get("종료") or "")[:10]
        if end and day > date.fromisoformat(end):
            out |= {r["이름"], short_name(r["이름"])}
    return out


def retiring(store: Path, day: date, days: int = 7) -> list[tuple[str, date]]:
    """마지막 근무일이 오늘~7일 뒤인 사람 [(이름, 마지막 근무일)] — 퇴사 예정 안내."""
    out = []
    for r in load_roster(store):
        end = str(r.get("종료") or "")[:10]
        if end and 0 <= (date.fromisoformat(end) - day).days <= days:
            out.append((r["이름"], date.fromisoformat(end)))
    return sorted(out, key=lambda x: x[1])


def set_last_day(store: Path, name: str, last: date | None) -> bool:
    """명단의 마지막 근무일(종료)을 정한다(없애려면 None). 이름이 명단에 없으면 False."""
    roster = load_roster(store)
    hit = [r for r in roster if r["이름"] == name or short_name(r["이름"]) == name]
    if len(hit) != 1:
        return False
    if last is None:
        hit[0].pop("종료", None)
    else:
        hit[0]["종료"] = last.isoformat()
    save_roster(store, roster)
    return True
