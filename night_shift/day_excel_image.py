"""평일 주간 현장 양식 이미지를 '진짜 엑셀'로 그린다 (2026-10-08 사용자 요청: 프로그램이 만든 티가 안 나게).

원본 편성표의 주간 시트(서식·조건부 번호 색·글꼴·병합·메모 표시)를 양식으로 복사해 두고
(templates/day_schedule_form.xlsx, tools/make_day_form.ps1), 값만 채운 뒤 설치된 엑셀이
'그림으로 복사'한 화면 그대로 PNG로 내보낸다(tools/render_day_form.ps1). 원본 이미지와 같은 1177×686.
엑셀이 없거나 실패하면 None → 호출한 쪽이 기존 그림 방식(day_image)으로 대신한다.
"""
from __future__ import annotations

import hashlib
import json
import re
import subprocess
from pathlib import Path

from .day_shift import (CONTROL_SLOTS, EDU, EDU_SLOTS, PATROL, REST, SHUTTLE, STANDBY, control_active,
                        control_patrol_at)

ROOT = Path(__file__).resolve().parent.parent
FORM = Path(__file__).resolve().parent / "templates" / "day_schedule_form.xlsx"
SCRIPT = ROOT / "tools" / "render_day_form.ps1"
CP_PATTERN = {3: ['CP1.CP2', 'CP2.CP3', 'CP3', 'CP1.CP2', 'CP2.CP3', 'CP1.CP3', 'CP1', 'CP2.CP3', 'CP1.CP3', 'CP1.CP2'],
              2: ['CP1.CP2', 'CP1', 'CP2', 'CP1.CP2', 'CP1', 'CP2', 'CP1', 'CP1', 'CP2', 'CP1']}
WEEKDAYS = "월화수목금토일"
KNOWN = ("G1", "G2", "G3", "4C", "L1", "L2", "C1", "C2", "CP", "1C", "셔틀", "순찰", "대기", "휴게")


def _norm(value) -> str:
    return re.sub(r"\s+", "", str(value or "")).upper()


_LAYOUT: dict = {}


def layout(form: Path = FORM) -> dict:
    key = (str(form), form.stat().st_mtime_ns)
    if key not in _LAYOUT:
        _LAYOUT[key] = _read_layout(form)
    return _LAYOUT[key]


def _read_layout(form: Path) -> dict:
    """양식에서 칸 위치를 읽는다: 제목 날짜 칸, 명단 두 줄, 머리글(포스트)별 열, 시간별 행."""
    import openpyxl
    from openpyxl.utils import get_column_letter
    ws = openpyxl.load_workbook(form).worksheets[0]
    header = next(r for r in range(1, 20) if _norm(ws.cell(r, 1).value) == "근무시간")
    columns = {}
    for c in range(2, ws.max_column + 1):
        label = _norm(ws.cell(header, c).value)
        if label and label not in columns and c <= ws.max_column and label in KNOWN:   # 인쇄 밖 메모 열 제외
            columns[label] = c
    hours = {}
    for r in range(header + 1, header + 30):
        m = re.fullmatch(r"(\d{1,2}):00", str(ws.cell(r, 1).value or "").strip())
        if m:
            hours[int(m[1])] = r
    date_cell = next(f"{get_column_letter(c.column)}{c.row}" for row in ws.iter_rows(min_row=1, max_row=header)
                     for c in row if isinstance(c.value, str) and "근무조" in c.value)
    return {"header": header, "names": (header - 2, header - 1), "columns": columns, "hours": hours,
            "date_cell": date_cell}


def _cell(row: int, col: int) -> str:
    from openpyxl.utils import get_column_letter
    return f"{get_column_letter(col)}{row}"


def cells_for(schedule, store, cp: dict | None) -> tuple[list[str], list[str], dict[int, dict[str, str]]]:
    """(명단 위 줄, 아래 줄, {시: {머리글: 글자}}). 업로드 저장본은 원본 칸 그대로, 생성표는 근무표 칸으로."""
    from .day_control import control_e_slots
    from .schedule_image_import import normalize_cell
    from . import daily_report as dr
    from .security_order import present
    layout_ = schedule.source_layout or {}
    raw = layout_.get("applied_rows") or layout_.get("rows")
    roster = dr.active(dr.load_roster(Path(store)), schedule.day)
    report = dr.build(schedule.day, dr.entries_for(Path(store), schedule.day), dr.load_roster(Path(store)))
    rows, _ = dr.apply_saved(report, dr.load_saved(Path(store), schedule.day))
    staff = {p["직책"]: p["이름"] for p in roster if p["직책"].startswith("통제")}
    on = {c for c in "DE" if staff.get(f"통제{c}") and present(staff[f"통제{c}"], rows, off=dr.is_off_day(schedule.day))}
    from .cp_schedule import short
    if raw and len(raw) == 10:
        names = {int(k): v for k, v in (layout_.get("mapping") or {}).items() if str(k).isdigit() and v}
        cps = {k[2:]: v for k, v in (layout_.get("extra_names") or {}).items() if k.startswith("CP") and v} or (cp or {})
        hours = {}
        for t, r in enumerate(raw):
            hour = {}
            for key, value in r.items():
                if key == "시간" or not value:
                    continue
                text = str(value).strip()
                hour[_norm(key)] = text if key == "셔틀" or "(" in text or text.startswith("교육") else normalize_cell(text)
            hours[8 + t] = hour
        controls = [c for c in layout_.get("controls") or [] if c in "DE"] or sorted(on)
    else:
        number = {p.pid: p.number for p in schedule.people if p.number}
        names = {n: pid for pid, n in number.items()}
        cps = cp or {}
        e_slots = set(control_e_slots(Path(store), schedule.day)) if "E" in on else set()
        hours = {}
        cp_count = len(cps)
        from .cp_schedule import hourly                 # 엑셀 표준 형태(같은 인원) 우선 (2026-10-11 사용자 확정)
        standard = hourly(Path(store), schedule.day, "day", cps) if cps else None
        for t in range(10):
            col = {pid: row[t] for pid, row in schedule.grid.items() if pid in number}

            def join(labels, extra=()):
                found = sorted((number[p] for p, label in col.items() if label in labels))
                return ".".join([*extra, *map(str, found)])
            hour = {post: join({post}, [w for p, w in schedule.fill.get(t, {}).items() if p == post])
                    for post in ("G1", "G2", "G3", "L1", "L2", "C1", "C2")}
            a_here = control_active(schedule, t)
            patrol = a_here and control_patrol_at(schedule, t)
            four = (["A"] if a_here and t in CONTROL_SLOTS else []) + (["E"] if t in e_slots else [])
            hour["4C"] = join({"통제"}, four)
            hour["순찰"] = join({PATROL}, ["A"] if patrol else [])
            hour["대기"] = join({STANDBY})
            cp_text = (standard["hours"][t] if standard and any(standard["hours"]) else
                       CP_PATTERN.get(cp_count, [""] * 10)[t])
            hour["CP"] = cp_text
            on_post = {int(n) for n in re.findall(r"CP(\d)", cp_text)}
            if standard and any(standard["hours"]):
                m = re.search(r"CP(\d)", standard["shuttle"]) if t == 9 else None
                shuttle_cp = int(m[1]) if m else None
            else:
                shuttle_cp = cp_count if t == 9 and cp_count in (2, 3) else None
            cp_rest = [f"CP{n}" for n in range(1, cp_count + 1) if n not in on_post and n != shuttle_cp]
            resting = a_here and t not in CONTROL_SLOTS and not patrol and "A" not in schedule.fill.get(t, {}).values()
            lead = (["E"] if "E" in on and t not in e_slots and t not in EDU_SLOTS else []) + (["A"] if resting else [])
            if t in EDU_SLOTS:          # 10·14시 교육 칸: '교육: 1.3.CP1' (CP는 쉬는 첫 사람만, 원본 작성 방식)
                edu = join({EDU, REST}, lead)
                hour["휴게"] = "교육: " + ".".join(x for x in (edu, cp_rest[0] if cp_rest else "") if x)
            else:
                hour["휴게"] = ".".join(x for x in (join({REST}, lead), ".".join(cp_rest)) if x)
            shuttle = schedule.shuttle.get(t)
            hour["셔틀"] = f"{number[shuttle]}({SHUTTLE[t]})" if shuttle in number else (f"CP{shuttle_cp}" if shuttle_cp else "")
            hours[8 + t] = hour
        controls = sorted(on)
    tags = [f"{n}{names[n]}" for n in sorted(names)]
    cp_tags = [f"CP{k}{short(v)}" for k, v in sorted(cps.items())]
    top = [f"통제{c}" for c in controls] + tags[:5]
    bottom = (["통제A"] if schedule.control else []) + tags[5:] + cp_tags
    return top, bottom, hours


def build_spec(schedule, store, cp: dict | None, team: str = "A", form: Path = FORM) -> dict:
    lay = layout(form)
    groups = sorted(lay["columns"].values())
    clear, values = [], {}
    for r in lay["names"]:
        clear += [_cell(r, c) for c in groups]
    for r in lay["hours"].values():
        clear += [_cell(r, c) for c in groups]
    d = schedule.day
    values[lay["date_cell"]] = f"{d.year}년 {d.month:02d}월 {d.day:02d}일 {WEEKDAYS[d.weekday()]}요일 근무조 : {team}조"
    top, bottom, hours = cells_for(schedule, store, cp)
    for r, names in zip(lay["names"], (top, bottom)):
        for c, name in zip(groups, names):
            values[_cell(r, c)] = name
    for hour, row in lay["hours"].items():
        for label, text in hours.get(hour, {}).items():
            col = lay["columns"].get(label)
            if col and text:
                if label == "휴게" and text.startswith("교육") and "대기" in lay["columns"]:
                    col = lay["columns"]["대기"]          # 10·14시 교육은 대기·휴게 합친 칸
                values[_cell(row, col)] = text
    return {"clear": clear, "values": values}


def to_xlsx(schedule, store, cp: dict | None = None, team: str = "A", form: Path = FORM) -> bytes:
    """원본 편성표 주간 시트와 같은 엑셀 파일 (2026-10-11 사용자 요청). 양식의 서식·조건부 번호 색·특이사항
    문구는 그대로 두고 이미지와 같은 칸 내용(build_spec)만 바꾼다. 엑셀 프로그램 없이 만든다."""
    from io import BytesIO
    import openpyxl
    from openpyxl.cell.cell import MergedCell
    spec = build_spec(schedule, store, cp, team, form)
    wb = openpyxl.load_workbook(form)
    ws = wb.worksheets[0]
    d = schedule.day
    ws.title = f"{d:%y}년{d:%m}월{d:%d}일(주)"
    for address in spec["clear"]:
        if not isinstance(ws[address], MergedCell):
            ws[address] = None
    for address, text in spec["values"].items():
        ws[address] = int(text) if text.isdigit() else text   # 원본처럼 번호 하나는 숫자로
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _target(schedule, store, cp, team):
    spec = build_spec(schedule, store, cp, team)
    key = hashlib.sha1((json.dumps(spec, ensure_ascii=False, sort_keys=True)
                        + str(FORM.stat().st_mtime_ns) + str(SCRIPT.stat().st_mtime_ns)).encode("utf-8")).hexdigest()[:20]
    cache = Path(store) / "day_images"
    return spec, key, cache, cache / f"{schedule.day.isoformat()}_{key}.png"


def cached(schedule, store, cp: dict | None = None, team: str = "A") -> bytes | None:
    """이미 같은 내용으로 그려 둔 그림이 있으면 그 PNG (엑셀을 다시 띄우지 않음)."""
    if not FORM.exists() or not SCRIPT.exists():
        return None
    out = _target(schedule, store, cp, team)[3]
    return out.read_bytes() if out.exists() else None


def render(schedule, store, cp: dict | None = None, team: str = "A", timeout: int = 90) -> bytes | None:
    """엑셀로 그린 PNG (같은 내용이면 저장해 둔 그림을 다시 씀). 엑셀이 없거나 실패하면 None."""
    if not FORM.exists() or not SCRIPT.exists():
        return None
    spec, key, cache, out = _target(schedule, store, cp, team)
    cache.mkdir(parents=True, exist_ok=True)
    if out.exists():
        return out.read_bytes()
    spec_path = cache / f"{key}.json"
    spec_path.write_text(json.dumps(spec, ensure_ascii=False), encoding="utf-8")
    try:
        subprocess.run(["powershell", "-NoProfile", "-STA", "-ExecutionPolicy", "Bypass", "-File", str(SCRIPT),
                        "-Form", str(FORM), "-Spec", str(spec_path), "-Output", str(out)],
                       capture_output=True, timeout=timeout, check=True)
    except (OSError, subprocess.SubprocessError):
        return None
    finally:
        spec_path.unlink(missing_ok=True)
    return out.read_bytes() if out.exists() else None
