"""Printable XLSX projection of one validated ServiceResult.

All values come from the service result. This module neither solves nor checks
rules and does not derive work, rest, fatigue or consecutive-run statistics.
"""

from io import BytesIO
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from .models import AssignmentKind
from .presentation import PATROL_LABELS, assignment_view, run_text
from .service import ServiceResult, rulebook_label


NAVY = "23395B"
PALE = "F1F5F9"
INK = "172033"
LINE = "CBD5E1"
WHITE = "FFFFFF"
FONT_NAME = "Arial"
THIN_BORDER = Border(bottom=Side(style="thin", color=LINE))
GRID_BORDER = Border(left=Side(style="thin", color=LINE), right=Side(style="thin", color=LINE),
                     top=Side(style="thin", color=LINE), bottom=Side(style="thin", color=LINE))


def _setup(ws, *, landscape=True):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.35, bottom=0.35,
                                  header=0.15, footer=0.15)
    ws.print_options.horizontalCentered = True


def _title(ws, title, last_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    cell = ws.cell(1, 1, title)
    cell.font = Font(name=FONT_NAME, size=15, bold=True, color=INK)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32


def _header(ws, row, labels):
    for column, label in enumerate(labels, 1):
        cell = ws.cell(row, column, label)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name=FONT_NAME, size=10, color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = GRID_BORDER
    ws.row_dimensions[row].height = 27


def _body(cell, *, fill=None, center=False, wrap=False):
    cell.font = Font(name=FONT_NAME, size=10, color=INK)
    cell.alignment = Alignment(horizontal="center" if center else "left",
                               vertical="center", wrap_text=wrap)
    cell.border = GRID_BORDER if center else THIN_BORDER
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)


def _sorted_people(result):
    return sorted(result.employee_stats.values(),
                  key=lambda person: (not person.is_manager, person.display_number or 0))


def _person_label(stats):
    if stats.employee_id == "A":
        return "통제"
    return f"{stats.display_number or 'A'} {stats.name}"


def _schedule_sheet(workbook, result):
    ws = workbook.active
    ws.title = "근무표"
    _setup(ws)
    _title(ws, "야간 근무표", 15)
    ws["A2"] = "기준일"
    ws["B2"] = result.context.shift_start_date
    ws["B2"].number_format = "yyyy-mm-dd"
    ws["D2"] = "N유형"
    ws["E2"] = result.context.night_type.value
    ws["G2"] = "운영 모드"
    ws["H2"] = result.context.effective_mode.value
    ws["J2"] = "관리자"
    manager = next((p for p in _sorted_people(result) if p.is_manager), None)
    ws["K2"] = _person_label(manager) if manager else "없음"
    ws["M2"] = "룰북"
    ws["N2"] = rulebook_label(result.rules_version)
    for key in ("A2", "D2", "G2", "J2", "M2"):
        ws[key].font = Font(name=FONT_NAME, size=10, bold=True, color=NAVY)
    for key in ("B2", "E2", "H2", "K2", "N2"):
        ws[key].font = Font(name=FONT_NAME, size=10, color=INK)
    ws.row_dimensions[2].height = 25
    ws["A3"] = "검증"
    ws["A3"].font = Font(name=FONT_NAME, size=10, bold=True, color=NAVY)
    ws.merge_cells("B3:O3")
    ws["B3"] = "하드룰 검증 통과 · 위반 0건"
    ws["B3"].font = Font(name=FONT_NAME, size=10, bold=True, color="166534")
    ws.row_dimensions[3].height = 24

    slots = sorted({row.slot for row in result.record.assignments}, key=lambda slot: slot.index)
    rows = {(row.employee_id, row.slot.index): row for row in result.record.assignments}
    _header(ws, 5, ["직원"] + [f"{slot.hour:02d}시" for slot in slots])
    for excel_row, person in enumerate(_sorted_people(result), start=6):
        name_cell = ws.cell(excel_row, 1, _person_label(person))
        _body(name_cell, fill=PALE, center=True)
        name_cell.font = Font(name=FONT_NAME, size=10, color=INK, bold=True)
        for column, slot in enumerate(slots, start=2):
            assignment = rows[(person.employee_id, slot.index)]
            label, notes, color, _ = assignment_view(assignment)
            cell = ws.cell(excel_row, column,
                           label + ("\n" + " · ".join(notes) if notes else ""))
            _body(cell, fill=color, center=True, wrap=True)
        ws.row_dimensions[excel_row].height = 42
    legend_row = 7 + len(result.employee_stats)
    ws.cell(legend_row, 1, "색상 구분")
    legends = (("휴게", "E5E7EB"), ("보안실", "DBEAFE"), ("G1", "BBF7D0"),
               ("L3", "BFDBFE"), ("I1", "FDE68A"), ("남문", "FECACA"),
               ("순찰", "DDD6FE"), ("차단봉", "FBCFE8"))
    for column, (label, fill) in enumerate(legends, start=2):
        cell = ws.cell(legend_row, column, label)
        _body(cell, fill=fill, center=True)
    ws.column_dimensions["A"].width = 16
    for column in range(2, 16):
        ws.column_dimensions[get_column_letter(column)].width = 12.5
    ws.freeze_panes = "B6"
    ws.print_title_rows = "1:5"
    ws.print_area = f"A1:O{legend_row}"
    return ws


def _summary_sheet(workbook, result):
    ws = workbook.create_sheet("직원별 요약")
    _setup(ws)
    _title(ws, "직원별 요약", 12)
    ws["A2"] = "기준일"
    ws["B2"] = result.context.shift_start_date
    ws["B2"].number_format = "yyyy-mm-dd"
    ws["D2"] = "유형/모드"
    ws["E2"] = f"{result.context.night_type.value} / {result.context.effective_mode.value}"
    headers = ("직원", "실근무", "실휴게", "피로도", "입식", "좌식", "순찰 횟수",
               "순찰 피로도", "표시 포스트", "최대 연속 실근무", "최대 연속 룰근무", "최대 연속 실휴게")
    _header(ws, 4, headers)
    for row_number, stats in enumerate(_sorted_people(result), start=5):
        values = (_person_label(stats), stats.actual_work_time_value,
                  stats.actual_rest_time_value, stats.fatigue_value, stats.post_standing,
                  stats.post_seated, stats.patrol_count, stats.patrol_fatigue,
                  stats.displayed_post_total, stats.max_actual_work_run,
                  stats.max_rule_work_run, stats.max_rest_run)
        for column, value in enumerate(values, 1):
            cell = ws.cell(row_number, column, value)
            _body(cell, center=column > 1)
            if column > 1:
                cell.number_format = "0.##"
        ws.row_dimensions[row_number].height = 25

    details_header = 7 + len(result.employee_stats)
    _header(ws, details_header,
            ("직원", "표시 포스트별 횟수", "포스트별 입식/좌식", "순찰구역",
             "최장 실근무 구간", "최장 실휴게 구간"))
    for row_number, stats in enumerate(_sorted_people(result), start=details_header + 1):
        counts = ", ".join(f"{label} {count}회" for label, count in stats.displayed_counts.items()
                           if label in ("G1", "L3", "I1", "G2", "G3", "G2(1)") and count)
        posture = ", ".join(f"{label} 입식 {values['standing']}·좌식 {values['seated']}"
                            for label, values in stats.posture_by_post.items()
                            if values["standing"] or values["seated"])
        patrols = ", ".join(PATROL_LABELS.get(row.displayed_assignment.label,
                                              row.displayed_assignment.label)
                            for row in result.record.assignments
                            if row.employee_id == stats.employee_id
                            and row.displayed_assignment.kind == AssignmentKind.PATROL)
        values = (_person_label(stats), counts, posture, patrols or "없음",
                  run_text(stats.longest_actual_work), run_text(stats.longest_rest))
        for column, value in enumerate(values, 1):
            cell = ws.cell(row_number, column, value)
            _body(cell, wrap=True)
        ws.row_dimensions[row_number].height = 57
    ws.column_dimensions["A"].width = 16
    for column in range(2, 13):
        ws.column_dimensions[get_column_letter(column)].width = 15
    ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 27
    ws.column_dimensions["E"].width = 34
    ws.column_dimensions["F"].width = 34
    ws.freeze_panes = "B5"
    ws.print_area = f"A1:L{details_header + len(result.employee_stats)}"
    return ws


def _validation_sheet(workbook, result):
    ws = workbook.create_sheet("검증 결과")
    _setup(ws)
    _title(ws, "검증 결과", 5)
    entries = (("Validator", "하드룰 검증 통과"), ("위반 건수", 0),
               ("Solver 상태", result.solve_result.status),
               ("목표 충족", "예" if result.solve_result.targets_met else "아니요"),
               ("시도 횟수", len(result.solve_result.attempts)),
               ("룰북", rulebook_label(result.rules_version)))
    for row_number, (label, value) in enumerate(entries, start=2):
        ws.cell(row_number, 1, label)
        ws.cell(row_number, 2, value)
        _body(ws.cell(row_number, 1), fill=PALE)
        _body(ws.cell(row_number, 2))
    _header(ws, 9, ("시도", "상태", "목표 탐색", "실행 초", "기타 Solver 상세"))
    for row_number, attempt in enumerate(result.solve_result.attempts, start=10):
        details = {key: value for key, value in attempt.items()
                   if key not in ("status", "balance_target_probe", "seconds")}
        values = (row_number - 9, attempt.get("status"),
                  "예" if attempt.get("balance_target_probe") else "아니요",
                  attempt.get("seconds"), json.dumps(details, ensure_ascii=False, default=str))
        for column, value in enumerate(values, 1):
            cell = ws.cell(row_number, column, value)
            _body(cell, wrap=column == 5)
        ws.row_dimensions[row_number].height = 38
    note_row = 11 + len(result.solve_result.attempts)
    ws.cell(note_row, 1, "Validator 위반")
    ws.cell(note_row, 2, "없음")
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 85
    ws.print_area = f"A1:E{note_row}"
    return ws


def _history_sheet(workbook, result):
    ws = workbook.create_sheet("이력 분석")
    _setup(ws)
    _title(ws, "이력 분석", 8)
    prepared = result.prepared_inputs
    ws["A2"] = "확정 실제 이력"
    ws["B2"] = prepared.actual_record_count
    ws["A3"] = "가정 누적값"
    ws["B3"] = "사용" if prepared.assumed_history_used else "없음"
    last_row = 4
    if prepared.historical_totals is not None:
        labels = ("직원", "과거 출근 횟수", "누적 피로도", "누적 입식", "누적 좌식",
                  "누적 실근무", "누적 실휴게", "누적 순찰 피로도")
        _header(ws, 5, labels)
        keys = ("fatigue_value", "standing", "seated", "actual_work_time_value",
                "actual_rest_time_value", "patrol_fatigue")
        names = {person.employee_id: person.name for person in result.context.employees}
        for row_number, employee_id in enumerate(prepared.historical_nights or {}, start=6):
            values = (f"{employee_id} {names.get(employee_id, '')}",
                      prepared.historical_nights[employee_id],
                      *(prepared.historical_totals.get(key, {}).get(employee_id, 0) for key in keys))
            for column, value in enumerate(values, 1):
                cell = ws.cell(row_number, column, value)
                _body(cell, center=column > 1)
            last_row = row_number
    else:
        ws["A5"] = "입력된 누적 이력 없음"
        last_row = 5
    prior = (("직전 설치자", prepared.previous_install_ids),
             ("직전 철수자", prepared.previous_remove_ids),
             ("이번 설치 제외자", prepared.forbid_install_ids),
             ("직전 07시 실휴게자", prepared.forbid_last_rest_ids),
             ("직전 순찰 제외자", prepared.forbid_patrol_exclusion_ids),
             ("순찰 회전 우선자", prepared.patrol_rotation_priority_ids))
    for row_number, (label, ids) in enumerate(prior, start=last_row + 2):
        ws.cell(row_number, 1, label)
        ws.cell(row_number, 2, ", ".join(ids) or "없음")
        _body(ws.cell(row_number, 1), fill=PALE)
        _body(ws.cell(row_number, 2))
        last_row = row_number
    if prepared.soft_debt:
        _header(ws, last_row + 2, ("soft rule ID", "직원 ID", "debt"))
        last_row += 2
        for rule, debts in prepared.soft_debt.items():
            for employee_id, value in debts.items():
                last_row += 1
                for column, item in enumerate((rule, employee_id, value), 1):
                    cell = ws.cell(last_row, column, item)
                    _body(cell)
    ws.column_dimensions["A"].width = 22
    for column in range(2, 9):
        ws.column_dimensions[get_column_letter(column)].width = 19
    ws.freeze_panes = "B6" if prepared.historical_totals is not None else "A1"
    ws.print_area = f"A1:H{last_row}"
    return ws


def export_xlsx(result: ServiceResult) -> bytes:
    """Export only a generated plan with zero Validator errors."""
    if not result.valid:
        raise ValueError("검증을 통과한 근무표만 Excel로 내보낼 수 있습니다.")
    workbook = Workbook()
    _schedule_sheet(workbook, result)
    _summary_sheet(workbook, result)
    _validation_sheet(workbook, result)
    _history_sheet(workbook, result)
    workbook.properties.title = "야간 근무표"
    workbook.properties.subject = f"{result.context.shift_start_date} {result.context.night_type.value}"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
