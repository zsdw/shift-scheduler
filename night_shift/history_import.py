"""Local, review-first night history import. No solver or source-sheet execution."""
from __future__ import annotations
import base64, hashlib, io, json, os, re, subprocess, tempfile
from dataclasses import dataclass, replace
from datetime import date, datetime, time
from pathlib import Path
from .models import *
from .profiles import make_record
from .service import StaffInput, ShiftRequest, build_context, calculate_all_stats
from .calendar_info import night_type_for
from . import history_store

ROSTER = {'A': '통제A', '1': '만년', '2': '원종', '3': '제욱', '4': '광희', '5': '성호', '6': '준호'}
POSTS = ('G1', 'L3', 'I1', 'G3', 'G2', 'G2(1)', 'G2(2)', '순찰', '보안실')
CHOICES = ('휴게', '보안실 대기', 'G1', 'L3', 'I1', 'G2', 'G3', 'G2(1)',
           '상층순찰', '중층순찰', '하층순찰', '근무불가', '확인 필요')
PATROLS = {3: '상층순찰', 4: '중층순찰', 5: '하층순찰', 7: '상층순찰', 8: '중층순찰', 9: '하층순찰'}

@dataclass
class ImportDraft:
    day: date | None
    sheets: str
    mapping: dict
    grid: dict
    warnings: list
    raw: list

def detected_date(text):
    year = re.search(r'(20\d{2})\s*(?:년|[-/.])', text)
    md = re.search(r'(\d{1,2})\s*월\s*(\d{1,2})\s*일', text)
    iso = re.search(r'(20\d{2})[-/.](\d{1,2})[-/.](\d{1,2})', text)
    try:
        if iso: return date(*(int(v) for v in iso.groups()))
        if year and md: return date(int(year[1]), int(md[1]), int(md[2]))
    except ValueError: pass
    return None

def hour_of(value):
    if isinstance(value, (datetime, time)): return value.hour
    m = re.fullmatch(r'\s*(\d{1,2})\s*[:시]\s*(?:00)?\s*', str(value))
    return int(m[1]) % 24 if m and int(m[1]) in (*range(8), *range(18, 25)) else None

def normalize_header(value):
    return re.sub(r'\s+', '', str(value or '')).replace('휴게', '휴게').upper()

def tokens(value):
    value = str(value or '').strip()
    if not value: return [], ''
    # One token = one nightly display number; never map CP personnel into the roster.
    value = value.upper().replace('통제A', 'A')
    value = re.sub(r'([1-6])\.0$', r'\1', value) if re.fullmatch(r'[1-6]\.0', value) else value
    parts = [v for v in re.split(r'[.,·、/\s]+', value) if v]
    if any(v not in ROSTER for v in parts): return [], value
    return parts, ''

def draft_from_cells(cells, day=None, sheet='현장 양식', mapping=None, warnings=None):
    grid = {e: ['휴게'] * 14 for e in ROSTER}
    warnings = list(warnings or [])
    seen = set()
    for row in cells:
        hour = row['hour']; t = (hour - 18) % 24
        if t not in range(14) or t in seen: raise ValueError('시간이 중복되거나 야간 범위를 벗어납니다.')
        seen.add(t)
        assigned = {}; uncertain=False
        for post in POSTS:
            persons, unknown = tokens(row.get(post, ''))
            if unknown:
                uncertain=True
                warnings.append(f'{hour:02}시 {post}: 읽을 수 없는 값 {unknown!r}. 원본과 비교해 수정하세요.')
            if post in ('G1', 'L3', 'I1') and not persons:
                uncertain=True
                warnings.append(f'{hour:02}시 {post}: 배정을 읽지 못했습니다. 확인이 필요합니다.')
            if post=='순찰' and t in PATROLS and not persons:
                uncertain=True
                warnings.append(f'{hour:02}시 순찰: 담당자를 읽지 못했습니다.')
            if post == 'G2(2)' and persons:
                warnings.append(f'{hour:02}시 G2(2): 현재 야간 룰북의 기록 포스트에 없어 직접 확인하세요.')
            label = PATROLS.get(t, '확인 필요') if post == '순찰' else '보안실 대기' if post == '보안실' else post
            for e in persons:
                if e in assigned:
                    grid[e][t] = '확인 필요'
                    who = (mapping or ROSTER).get(e, ROSTER.get(e, e))   # 그 시트의 번호 이름 (예전엔 고정 명단이라 2번=원종으로 잘못 표시)
                    warnings.append(f'{hour:02}시 {e}번 {who}: {assigned[e]}·{post} 중복. 실제 배정을 골라주세요.')
                else:
                    grid[e][t] = label if label in CHOICES else '확인 필요'
                assigned[e] = post
        if uncertain:
            for e in ROSTER:
                if e not in assigned:grid[e][t]='확인 필요'
    if seen != set(range(14)): raise ValueError('18~07시 14개 시간 행을 모두 찾지 못했습니다.')
    return ImportDraft(day, sheet, dict(mapping or ROSTER), grid, warnings, cells)

def excel_sheets(data):
    from openpyxl import load_workbook
    if len(data) > 20 * 1024 * 1024: raise ValueError('20MB 이하 파일을 올려주세요.')
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try: return book.sheetnames
    finally: book.close()

def parse_excel(data, sheet):
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = book[sheet]
        if ws.max_row > 500 or ws.max_column > 200: raise ValueError('한 시트당 500행·200열 이하의 근무표를 올려주세요.')
        values = list(ws.values)
    finally: book.close()
    alltext = ' '.join(str(v) for row in values for v in row if v is not None)
    mapping = dict(ROSTER)
    for row in values:
        for value in row:
            m = re.fullmatch(r'\s*([1-6])\s*([가-힣]{2,4})\s*', str(value or ''))
            if m: mapping[m[1]] = m[2]
    for index,row in enumerate(values):
        if str(row[0] or '').strip()!='직원':continue
        hours={hour_of(v):j for j,v in enumerate(row[1:],1) if hour_of(v) is not None}
        if set(hours)!=set((*range(18,24),*range(8))):continue
        grid={e:['근무불가']*14 for e in ROSTER};seen=set();warnings=[]
        for body in values[index+1:]:
            name=str(body[0] or '').strip()
            if name.startswith('통제'):token='A'
            else:
                m=re.match(r'^([1-6])\s+([가-힣]+)',name)
                if not m:continue
                token=m[1];mapping[token]=m[2]
            if token in seen:raise ValueError('직원 행이 중복됩니다.')
            seen.add(token)
            for hour,col in hours.items():
                value=str(body[col] or '').split('\n')[0].strip()
                value={'상층':'상층순찰','중층':'중층순찰','하층':'하층순찰','×':'근무불가'}.get(value,value)
                grid[token][(hour-18)%24]=value if value in CHOICES else '확인 필요'
        if not seen:raise ValueError('직원 배정 행을 찾지 못했습니다.')
        return ImportDraft(detected_date(alltext),sheet,mapping,grid,
            ['직원별 근무표입니다. 셀 메모는 자동 반영하지 않으므로 실제 상태와 차단봉 담당자를 확인하세요.'],
            [{'번호':e,**{f'{(18+t)%24:02}시':v for t,v in enumerate(grid[e])}} for e in seen])
    for index, row in enumerate(values):
        headers = {normalize_header(v): j for j, v in enumerate(row) if v is not None}
        if not {'G1', 'L3', 'I1'}.issubset(headers): continue
        timecol = headers.get('근무시간', headers.get('시간', 0))
        cells = []
        for body in values[index+1:]:
            hour = hour_of(body[timecol] if timecol < len(body) else None)
            if hour is None: continue
            item = {'hour': hour}
            for post in POSTS:
                col = headers.get(post.upper())
                item[post] = body[col] if col is not None and col < len(body) else ''
            cells.append(item)
            if len(cells) == 14: break
        return draft_from_cells(cells, detected_date(alltext), sheet, mapping,
                                ['휴게 명단과 아래 특이사항은 가져오지 않습니다. 포스트·순찰·대기 외에는 휴게로 처리합니다.'])
    raise ValueError('G1·L3·I1 열과 근무시간이 있는 현장 근무표 시트를 찾지 못했습니다.')

def image_ocr(data, *, clean=True):
    from PIL import Image, ImageOps
    with Image.open(io.BytesIO(data)) as source:
        image = ImageOps.exif_transpose(source).convert('RGB')
    if image.width * image.height > 25_000_000: raise ValueError('이미지가 너무 큽니다. 2,500만 화소 이하로 줄여주세요.')
    # 2026-10-04: 폰 사진·캡처(JPG 압축·흐림·작은 크기)도 읽도록 확대 → 선명화 → 흑백 정리 후 문자 인식
    from PIL import ImageFilter
    factor = min(3, 3000 / max(image.size))
    image = image.resize((round(image.width*factor), round(image.height*factor)), Image.LANCZOS)
    if clean:
        gray = ImageOps.autocontrast(ImageOps.grayscale(image)).filter(ImageFilter.UnsharpMask(radius=3, percent=200, threshold=0))
        image = gray.point(lambda v: 0 if v < 140 else 255).convert('RGB')
    if os.name != 'nt': raise ValueError('이미지 자동 인식은 Windows PC에서 지원합니다. 엑셀 파일을 이용해 주세요.')
    with tempfile.TemporaryDirectory(prefix='night_ocr_') as folder:
        path = Path(folder)/'source.png'; image.save(path)
        script=Path(__file__).with_name('image_ocr.ps1').read_text(encoding='utf-8')
        script='$Path=$env:NIGHT_OCR_FILE\n'+script.split('\n',1)[1]
        encoded=base64.b64encode(script.encode('utf-16le')).decode('ascii')
        try:
            proc = subprocess.run(['powershell.exe', '-NoProfile', '-EncodedCommand',encoded],
                                  env={**os.environ,'NIGHT_OCR_FILE':str(path)},capture_output=True, timeout=60,
                                  creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
        except subprocess.TimeoutExpired as exc:
            raise ValueError('Windows 문자 인식 시간이 초과됐습니다. 파일 다시 분석을 누르거나 더 작은 원본을 이용하세요.') from exc
    if proc.returncode: raise ValueError('Windows 문자 인식에 실패했습니다. Windows 한국어 언어팩을 확인해 주세요.')
    result = json.loads(proc.stdout.decode('utf-8-sig'))
    return image, result

def parse_image(data):
    import numpy as np
    image, ocr = image_ocr(data)
    pixels = np.asarray(image)
    black = (pixels.max(axis=2) < 100)
    def centers(indices):
        groups=[]
        for n in indices:
            if not groups or n > groups[-1][-1]+2: groups.append([int(n)])
            else: groups[-1].append(int(n))
        return [sum(g)/len(g) for g in groups]
    horizontal = centers(np.flatnonzero(black.mean(axis=1) > .7))
    words = ocr['words']
    times = [w for w in words if w['x'] < image.width*.08 and re.match(r'18\s*[:시]?\s*00', w['text'])]
    if not times: raise ValueError('18:00 행을 찾지 못했습니다. 표 전체가 선명하게 보이는 이미지를 올려주세요.')
    cy = times[0]['y']+times[0]['h']/2
    tops = [v for v in horizontal if v < cy]
    if not tops: raise ValueError('표 테두리를 찾지 못했습니다.')
    top=max(tops); bodylines=[v for v in horizontal if v>=top]
    if len(bodylines)<15: raise ValueError('18~07시의 가로 테두리를 모두 찾지 못했습니다.')
    bodylines=bodylines[:15]
    vertical=centers(np.flatnonzero(black[round(top)+3:round(bodylines[-1])-3].mean(axis=0)>.85))
    if vertical and vertical[0]>image.width*.02:vertical.insert(0,0)
    if vertical and vertical[-1]<image.width*.98:vertical.append(image.width-1)
    if len(vertical)!=14: raise ValueError('현재 이미지는 보내주신 13열 현장 양식을 지원합니다. 표를 잘리지 않게 올려주세요.')
    def cell(x0,y0,x1,y1):
        selected=[w for w in words if x0<w['x']+w['w']/2<x1 and y0<w['y']+w['h']/2<y1]
        return ' '.join(w['text'] for w in sorted(selected,key=lambda w:w['x']))
    cells=[]
    columns={1:'G1',2:'L3',3:'I1',4:'G3',5:'G2',6:'G2(1)',7:'G2(2)',10:'순찰',11:'보안실'}
    boxes={}
    for t in range(14):
        item={'hour':(18+t)%24}
        for col,post in columns.items():
            boxes[t,post]=(vertical[col],bodylines[t],vertical[col+1],bodylines[t+1])
            item[post]=cell(*boxes[t,post])
        cells.append(item)
    # Windows 한국어 문자 인식은 칸에 'A'처럼 글자가 하나만 있으면 아예 읽지 않는다 (2026-10-04 확인).
    # 처음 읽기에서 비었는데 글씨가 있는 칸만, 이미 읽은 숫자 글자 옆에 붙여 다시 읽는다.
    missing=[(t,post) for t in range(14) for post in columns.values()
             if not cells[t][post].strip() and _ink(black,boxes[t,post]) is not None]
    time_boxes={t:(vertical[0],bodylines[t],vertical[1],bodylines[t+1]) for t in range(14)}
    for (t,post),text in _reread_with_context(image,black,cells,boxes,missing,time_boxes).items():
        cells[t][post]=text
    for item in cells:
        for post in columns.values():
            # OCR's I/l are common numeric-1 confusions inside numbered cells only.
            item[post]=re.sub(r'(?<![A-Z])[Il](?![A-Z])','1',item[post].replace('•','.').replace('·','.'))
    mapping=dict(ROSTER)
    return draft_from_cells(cells,detected_date(ocr['text']),'이미지',mapping,
        ['이미지 인식 결과입니다. 빈칸·숫자·번호별 이름을 반드시 원본과 대조하세요.',
         '아래 특이사항과 휴게 명단은 제외했습니다. CP·셔틀은 야간 인원 이력에 넣지 않습니다.'])

def _ink(black, box, pad=5):
    """칸 안 글씨 영역 (x0,y0,x1,y1). 테두리 선은 pad만큼 빼고 본다. 글씨가 없으면 None."""
    x0,y0,x1,y1=(round(v) for v in box)
    import numpy as np
    region=black[y0+pad:y1-pad, x0+pad:x1-pad]
    ys,xs=np.nonzero(region)
    if len(xs)<12: return None
    return (x0+pad+xs.min(), y0+pad+ys.min(), x0+pad+xs.max()+1, y0+pad+ys.max()+1)

def _reread_with_context(image, black, cells, boxes, missing, time_boxes=None):
    """빈칸으로 읽힌 글씨 칸을 칸마다 다시 읽는다: 칸만 → 이미 읽은 숫자 두 개 옆 → 같은 줄 시간 칸(예: 18:00) 옆.
    흐린 폰 사진은 숫자 칸이 하나도 안 읽혀 기준 숫자가 없으므로 시간 칸을 기준 글자로 쓴다 (2026-10-04)."""
    if not missing: return {}
    from PIL import Image
    ref=None
    for t,item in enumerate(cells):                       # 이미 한 글자 숫자로 읽힌 칸을 기준 글자로
        for post,value in item.items():
            if post!='hour' and re.fullmatch(r'[1-6]',str(value).strip()) and (t,post) in boxes:
                ink=_ink(black,boxes[t,post])
                if ink: ref=(value.strip(),image.crop(ink)); break
        if ref: break
    if ref is None and not time_boxes: return {}
    digit,glyph=ref if ref else (None,None)

    def ocr_text(canvas):
        buffer=io.BytesIO(); canvas.save(buffer,'PNG')
        return ' '.join(w['text'] for w in sorted(image_ocr(buffer.getvalue())[1]['words'],key=lambda w:w['x']))

    def tokens_only(text):                               # 번호 칸 형식(1~6·A·점)일 때만 받는다
        text=re.sub(r'(?<![A-Z])[Il](?![A-Z])','1',text.replace(' ','').replace('•','.').replace('·','.').replace(',','.'))
        return text if re.fullmatch(r'[1-6A](\.[1-6A])*\.?',text) else ''

    def read(key):
        # 칸마다 한 장씩: 여러 칸을 한 장에 모으면 엔진이 'A'를 다시 빠뜨리거나 'A.1'을 '4.1'로 읽었다.
        x0,y0,x1,y1=(round(v) for v in boxes[key])
        cell_img=image.crop((x0+4,y0+4,x1-4,y1-4))
        if (text:=tokens_only(ocr_text(cell_img))):      # 'A.1'처럼 글자가 여럿이면 칸만으로 읽힌다
            return text
        piece=image.crop(_ink(black,boxes[key]))
        for gap in ((20,40) if glyph is not None else ()):   # 'A' 한 글자는 기준 숫자 두 개 옆에 붙여야 읽힌다
            h=max(glyph.height,piece.height)+40
            canvas=Image.new('RGB',(glyph.width*2+piece.width+60+gap*2,h),'white')
            canvas.paste(glyph,(20,20)); canvas.paste(glyph,(40+glyph.width,20))
            canvas.paste(piece,(40+gap+glyph.width*2,20))
            line=ocr_text(canvas.resize((canvas.width*2,canvas.height*2)))
            rest=re.sub(rf'^\s*{digit}\s*{digit}\s*','',line,count=1)   # 앞의 기준 숫자 두 개를 뗀 나머지
            if rest!=line and (text:=tokens_only(rest)):
                return text
        t=key[0]
        stamp=_ink(black,time_boxes[t]) if time_boxes and t in time_boxes else None
        if stamp:                                        # 같은 줄 시간 칸('18:00') 옆에 붙여 읽기
            clock=image.crop(stamp)
            for gap in (30,60):
                h=max(clock.height,piece.height)+40
                canvas=Image.new('RGB',(clock.width+piece.width+40+gap,h),'white')
                canvas.paste(clock,(20,20)); canvas.paste(piece,(20+clock.width+gap,20))
                line=ocr_text(canvas)
                rest=re.sub(r'^\s*\d{1,2}\s*[:.;]?\s*\d{2}\s*','',line,count=1)
                if rest!=line and (text:=tokens_only(rest)):
                    return text
        return ''
    from concurrent.futures import ThreadPoolExecutor
    with ThreadPoolExecutor(max_workers=4) as pool:
        texts=list(pool.map(read,missing))
    return {key:text for key,text in zip(missing,texts) if text}

def apply_leave(grid, token_names, statuses):
    """연차일정의 반차·연차 시간 중 '휴게'로 읽힌 칸을 '근무불가'로 바꾼다 (룰북 21-7, 2026-10-05).

    현장 양식은 근무불가 시간을 빈칸으로 두므로 그림·엑셀만으로는 휴게와 구분할 수 없다.
    그 시간에 실제 배정이 적혀 있으면 기록을 우선해 그대로 두고 안내만 한다.
    token_names = {번호: 명단 이름 또는 'A'}, statuses = {이름: (상태, 비고)}. 반환 (새 grid, 안내들)."""
    from .leave_schedule import night_unavailable_slots
    out = {token: list(cells) for token, cells in grid.items()}
    notes = []
    for token, name in token_names.items():
        status = statuses.get(name, (None,))[0]
        slots = night_unavailable_slots(status) if status else set()
        if token not in out or not slots:
            continue
        busy = [t for t in sorted(slots) if out[token][t] not in ('휴게', '근무불가')]
        if busy:                                  # 반차 시간에 근무 기록 = 반차 취소·계획표일 수 있음 → 바꾸지 않음
            notes.append(f'{token}번 {name}: 연차일정은 {status}인데 {", ".join(f"{(18+t)%24:02}시" for t in busy)}에 '
                         '근무가 적혀 있어 기록대로 둡니다(근무불가 표시 안 함). 실제와 다르면 확인하세요.')
            continue
        for t in slots:
            if out[token][t] == '휴게':
                out[token][t] = '근무불가'
        notes.append(f'{token}번 {name}: 연차일정 {status} → 해당 시간 근무불가(×)')
    return out, notes


def leave_fix_record(record, slots_by_id):
    """저장된 기록에서 반차·연차 시간의 '휴게' 칸을 근무불가로 바로잡는다 (2026-10-05, 룰북 21-7).
    slots_by_id = {직원 ID: 근무불가 슬롯 집합}. 실제 배정이 적힌 칸은 건드리지 않는다. (새 기록, 바꾼 칸 목록)."""
    from .models import AssignmentKind as K
    from .profiles import options_for
    inside = [(eid, r) for r in record.assignments for eid, slots in slots_by_id.items()
              if r.employee_id == eid and r.slot.index in slots]
    busy = {eid for eid, r in inside if r.actual_post_or_work_type.kind not in (K.REST, K.UNAVAILABLE)}
    fix = {(eid, r.slot.index) for eid, r in inside if eid not in busy   # 반차 시간에 근무 기록 있으면 그 사람은 손대지 않음
           and r.displayed_assignment.kind == K.REST and r.actual_post_or_work_type.kind == K.REST}
    if not fix:
        return record, []
    attendance = []
    for a in record.context.attendance:
        gone = {t for e, t in fix if e == a.employee_id}
        if gone:
            keep = tuple(t for t in (a.available_slot_indices if a.available_slot_indices is not None else range(14))
                         if t not in gone)
            a = replace(a, available_slot_indices=keep, present=bool(keep))
        attendance.append(a)
    ctx = replace(record.context, attendance=tuple(attendance))
    rows = tuple(options_for(ctx, r.employee_id, r.slot.index, patrol_standing=True)[0]
                 if (r.employee_id, r.slot.index) in fix else r for r in record.assignments)
    return replace(record, context=ctx, assignments=rows), sorted(fix, key=lambda x: (x[0], x[1]))


def build_import_record(day, mode, grid, identities, availability, installs=(), removes=()):
    if len(set(identities.values())) != len(identities): raise ValueError('번호별 실제 직원이 중복됩니다.')
    staff=[]
    for token,identity in identities.items():
        slots=tuple(t for t,label in enumerate(grid[token]) if label!='근무불가')
        staff.append(StaffInput(identity, availability[identity],bool(slots),token!='A',
                                None if token=='A' else int(token),token=='A',slots))
    if not any(p.present and p.acts_as_manager for p in staff):
        # 준호가 실제로 관리자였다는 확인 없이 자동 승계하지 않는다.
        raise ValueError('통제A가 없는 이력은 관리자 대행을 명시할 수 있는 양식으로 확인이 필요합니다.')
    ctx=build_context(ShiftRequest(day,NightType(night_type_for(day)),mode,tuple(staff)))
    rows=[]
    for token,identity in identities.items():
        for t,label in enumerate(grid[token]):
            if label not in CHOICES or label=='확인 필요': raise ValueError(f'{token}번 {(18+t)%24:02}시 배정을 확인하세요.')
            if label=='근무불가':
                if not next(a for a in ctx.attendance if a.employee_id==identity).present: continue
                from .profiles import options_for
                row=options_for(ctx,identity,t,patrol_standing=True)[0]
            else:
                label={'상층순찰':'upper','중층순찰':'middle','하층순찰':'lower'}.get(label,label)
                row=make_record(ctx,identity,t,label,install=(token,t) in installs,
                                remove=(token,t) in removes,patrol_standing=True)
                # Import only explicitly confirmed barrier activity, never infer actuals.
                row=replace(row,barrier_remove_flag=(token,t) in removes)
            rows.append(row)
    return ShiftRecord(f'import-{day}-{datetime.now().strftime("%H%M%S%f")}',RecordKind.PLAN,ctx,tuple(rows))

def fingerprint(store,day):
    path=Path(store)/f'{day}.json'
    return hashlib.sha256(path.read_bytes()).hexdigest() if path.exists() else None

def save_import(store,record,source,note,expected=None,extra=None):
    """Atomic replace with recoverable original and optimistic conflict check.
    extra: 함께 저장할 값 (예: 편성표 일괄 가져오기의 원본 시트 지문 source_hash)."""
    store=Path(store); store.mkdir(parents=True,exist_ok=True)
    day=record.context.shift_start_date; path=store/f'{day}.json'
    if fingerprint(store,day)!=expected: raise ValueError('이 날짜의 이력이 변경됐습니다. 화면을 새로고침한 뒤 다시 확인하세요.')
    extra=dict(extra or {});original_bytes=extra.pop('source_bytes',None)
    payload={'version':history_store.STORE_VERSION,'rules_version':'2.20',
             'confirmed_at':datetime.now().isoformat(timespec='seconds'),
             'plan':history_store.record_to_json(record),'edits':[], 'note':note,'source':source,**(extra or {})}
    from .import_compare import identical,mark_overwrite
    if identical(path,payload):return path
    mark_overwrite(path,payload)
    if original_bytes is not None:
        from .source_attachment import stash
        payload['source_attachment']=stash(store,original_bytes,payload['source_file'])
    encoded=json.dumps(payload,ensure_ascii=False,indent=1).encode('utf-8')
    fd,temp=tempfile.mkstemp(prefix='.import-',suffix='.tmp',dir=store)
    try:
        with os.fdopen(fd,'wb') as f: f.write(encoded)
        if path.exists():
            import shutil
            trash=store/'.trash';trash.mkdir(exist_ok=True)
            shutil.copy2(path,trash/f'{day}__deleted-import-{datetime.now().strftime("%Y%m%d-%H%M%S%f")}.json')
        os.replace(temp,path)
    finally:
        Path(temp).unlink(missing_ok=True)
    return path
