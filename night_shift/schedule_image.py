"""Exact, deterministic PNG of a recorded night, in the field form layout."""
from io import BytesIO
from pathlib import Path
from .models import AssignmentKind

def arrange_cells(headers,buckets):
    """한 시간 줄의 칸 정리 (2026-10-06 사용자 확정): 여러 명이 같이 적히면 A를 항상 맨 앞에(CP·셔틀 제외),
    G2에 두 명(추가 남문 운영)이면 A 아닌 사람은 G3 칸에."""
    skip={headers.index('CP'),headers.index('셔틀')}
    for j in range(len(headers)):
        if j not in skip:
            buckets[j]=sorted(buckets[j],key=lambda v:(v!='A',int(v) if v.isdigit() else 99,v))
    g2,g3=headers.index('G2'),headers.index('G3')
    if len(buckets[g2])>1 and not buckets[g3]:
        buckets[g3]=buckets[g2][1:];buckets[g2]=buckets[g2][:1]
    return buckets


HEADERS=['근무시간','G1','L3','I1','G3','G2','G2(1)','G2(2)','CP','셔틀','순찰','보안실','휴 게']


def field_cells(record, cp=None):
    """현장 양식 칸 내용 (이미지·엑셀 공통, 2026-10-11). 반환 {'names': 명단 7칸(통제A,1~6),
    'cps': CP 줄 3칸, 'rows': 14시간 × 13칸(첫 칸 시간), 'notes': (왼쪽, 오른쪽)}."""
    c=record.context
    people={e.employee_id:e.name for e in c.employees}
    present=[a for a in c.attendance if a.present]
    numbered={str(a.display_number) if a.display_number is not None else 'A':a for a in present}
    def token(e):
        a=next(a for a in present if a.employee_id==e)
        return str(a.display_number) if a.display_number is not None else 'A'
    names=[]
    for j in range(7):
        key='A' if j==0 else str(j)
        a=numbered.get(key)
        names.append('' if not a else (('통제A' if people[a.employee_id]=='통제' else people[a.employee_id]) if key=='A'
                                        else key+people[a.employee_id]))     # 현장 양식 표기 '통제A' (2026-10-05)
    from .cp_schedule import short, hourly
    cps=[f'CP{n}{short(cp[str(n)])}' if cp and str(n) in cp else '' for n in (1,2,3)]   # CP1덕영 · CP2철룡 · CP3혁준
    import os
    store=Path(os.environ.get('NIGHT_SHIFT_HISTORY_DIR') or Path(__file__).resolve().parent.parent/'history')
    pattern=hourly(store,c.shift_start_date,'night',cp or {})   # 엑셀 표준 형태(같은 인원) 우선 (2026-10-11)
    headers=HEADERS;rows=[]
    for t in range(14):
        buckets=[[] for _ in headers]
        for row in record.assignments:
            if row.slot.index!=t:continue
            label=row.displayed_assignment.label;who=token(row.employee_id)
            if row.displayed_assignment.kind==AssignmentKind.UNAVAILABLE:continue
            if label in headers:index=headers.index(label)
            elif row.displayed_assignment.kind==AssignmentKind.PATROL:index=10
            elif row.displayed_assignment.kind==AssignmentKind.SECURITY_STANDBY:index=11
            elif row.displayed_assignment.kind==AssignmentKind.REST:index=12
            else:raise ValueError(f'이미지에 표시할 수 없는 배정: {label}')
            buckets[index].append(who)
        if cp:                                           # 시간별 CP 칸·07시 셔틀 (엑셀 표준 형태)
            buckets[headers.index('CP')]=[x for x in pattern['hours'][t].split('.') if x]
            if t==13 and pattern['shuttle']:buckets[headers.index('셔틀')].append(pattern['shuttle'])
        arrange_cells(headers,buckets)
        rows.append([f'{24 if t==6 else (18+t)%24:02}:00']+['.'.join(b) for b in buckets[1:]])
    return {'names':names,'cps':cps,'rows':rows,'notes':field_duty_notes(record,token)}


def render_png(record, cp=None, *, team='A', cells=None):
    """cp = {'1': 이름, '2': 이름, '3': 이름} (CP 근무표 기본값, 없으면 CP 칸 빈칸)."""
    from PIL import Image,ImageDraw,ImageFont
    regular=Path('C:/Windows/Fonts/malgun.ttf');bold=Path('C:/Windows/Fonts/malgunbd.ttf')
    if not regular.exists():
        raise ValueError('이미지 출력용 한글 글꼴을 찾지 못했습니다.')
    scale=2;W,H=1112,790
    im=Image.new('RGB',(W*scale,H*scale),'white');d=ImageDraw.Draw(im)
    def box(a,b,c,e,fill='white'):
        d.rectangle(tuple(round(v*scale) for v in (a,b,c,e)),fill=fill,outline='black',width=scale)
    def text(a,b,c,e,value,size=14,strong=False,left=False):
        f=ImageFont.truetype(str(bold if strong else regular),size*scale)
        bb=d.textbbox((0,0),value,font=f)
        x=(a+6)*scale if left else (a+c)*scale/2-(bb[2]-bb[0])/2
        y=(b+e)*scale/2-(bb[3]-bb[1])/2-bb[1]
        d.text((x,y),value,font=f,fill='#001c48')
    c=record.context;day=c.shift_start_date
    box(0,0,880,57);text(0,0,880,57,'야간 근무 편성표',28,True)
    box(0,57,880,108);text(0,57,880,108,f'{day:%Y년 %m월 %d일} {"월화수목금토일"[day.weekday()]}요일 근무조 : {team}조',23,True)
    box(880,0,929,108,'#c0c0c0');text(880,0,929,108,'결\n재',15)
    for a,b,label in [(929,1017,'담 당'),(1017,1111,'팀 장')]:
        box(a,0,b,32,'#c0c0c0');text(a,0,b,32,label,15,True);box(a,32,b,108)
    xs=[0,80,160,240,320,400,480,560,640,720,800,880,960,1111]
    box(0,108,80,170,'#c0c0c0');text(0,108,80,170,'근무자명')
    cells=field_cells(record,cp) if cells is None else cells
    for j in range(12):
        box(xs[j+1],108,xs[j+2],139,'#c0c0c0');text(xs[j+1],108,xs[j+2],139,cells['names'][j] if j<7 else '',12)
        box(xs[j+1],139,xs[j+2],170,'#c0c0c0')
        if j<3 and cells['cps'][j]:                      # 둘째 줄: CP1덕영 · CP2철룡 · CP3혁준 (현장 양식)
            text(xs[j+1],139,xs[j+2],170,cells['cps'][j],12)
    for j,label in enumerate(HEADERS):
        box(xs[j],170,xs[j+1],200,'#c0c0c0');text(xs[j],170,xs[j+1],200,label)
    for t,values in enumerate(cells['rows']):
        y=200+t*33
        for j,value in enumerate(values):
            box(xs[j],y,xs[j+1],y+33);text(xs[j],y,xs[j+1],y+33,value)
    y=662;box(0,y,80,H-1);text(0,y,80,H-1,'특이사항');box(80,y,640,H-1);box(640,y,1111,H-1)
    left,right=cells['notes']                        # 현장 특이사항 (2026-10-05 사용자 확정 표기 규칙)
    for k,line in enumerate(left):text(80,y+3+k*24,400,y+27+k*24,line,14,True,left=True)
    for k,line in enumerate(right):text(395,y+3+k*24,640,y+27+k*24,line,14,True,left=True)
    # 오른쪽 칸의 운영 모드·차단봉 담당·근무불가 안내는 표시하지 않는다 (2026-10-05 사용자 요청)
    buffer=BytesIO();im.save(buffer,format='PNG');return buffer.getvalue()

def field_duty_notes(record,token):
    """현장 특이사항 칸 (2026-10-05 사용자 확정). 근무표 칸에서 번호를 읽는다. 반환 (왼쪽 5줄, 오른쪽 2줄).
    - 흡연장 파라솔: parasol() 참고 (2026-10-11 변경)
    - 21시 폐문 및 차단봉 설치: 20시 G2 근무자, G2가 없으면 20시 I1 근무자 (2026-10-06 사용자 확정)
    - 21시 라이브러리 시건: 21시 L3 / 24시 폐문(2차) 및 차단봉 설치: 24시 I1
    - 05시 개문 및 차단봉 공조실 이동: 04시 I1(철수자) / 05시 0호기 1층 이동: 항상 A / 06시 라이브러리 시건 해제: 06시 L3"""
    cells={(r.employee_id,r.slot.index):r for r in record.assignments}
    def who(slot,label):
        found=[token(r.employee_id) for (e,t),r in cells.items() if t==slot and r.displayed_assignment.label==label]
        found.sort(key=lambda v:(v!='A',int(v) if v.isdigit() else 99))   # A를 앞에 (2026-10-06)
        return '.'.join(found) or '-'
    closing=who(2,'G2')
    if closing=='-':closing=who(2,'I1')
    left=[f'흡연장 파라솔 - {parasol(record,token,who)}',f'21시 폐문 및 차단봉 설치 - {closing}',
          f'21시 라이브러리 시건 - {who(3,"L3")}',f'24시 폐문 및 차단봉 설치 - {who(6,"I1")}',
          f'05시 개문 및 차단봉 공조실 이동 - {who(10,"I1")}']
    right=['05시 0호기 1층 이동 - A',f'06시 라이브러리 시건 해제 - {who(12,"L3")}']
    return left,right


def parasol(record,token,who):
    """흡연장 파라솔 담당 (2026-10-11 사용자 확정).
    - 19시 G2 근무자가 있으면 그 사람
    - 없으면 선임(준호)이 19시 또는 20시에 휴게·보안실 대기면 선임, 19·20시 모두 근무 중이면 A
    - A가 연차(미출근)면 선임"""
    g2=who(1,'G2')
    if g2!='-':return g2
    present={a.employee_id for a in record.context.attendance if a.present}
    senior=next((e.employee_id for e in record.context.employees if e.name=='준호'),None)
    if senior not in present:return 'A' if 'A' in present else '-'
    if 'A' not in present:return token(senior)
    off=(AssignmentKind.REST,AssignmentKind.SECURITY_STANDBY)
    free=any(r.employee_id==senior and r.slot.index in (1,2) and r.displayed_assignment.kind in off
             for r in record.assignments)
    return token(senior) if free else 'A'


FORM_COLUMNS=['A','F','K','P','U','Z','AE','AJ','AO','AT','AY','BD','BI']   # 원본 야간 시트의 칸 (근무시간~휴게)


def render_xlsx(record, cp=None, *, team='A', cells=None):
    """원본 편성표 야간 시트와 같은 양식의 엑셀 (2026-10-11 사용자 요청). 양식(templates/field_night.xlsx)은
    원본 시트의 병합·글꼴·테두리·열 너비를 그대로 둔 빈 틀이고, 이미지와 같은 칸 내용만 채운다."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment
    from .field_exporter import TEMPLATE
    wb=load_workbook(TEMPLATE);ws=wb.active
    day=record.context.shift_start_date
    ws.title=f'{day:%y}년{day:%m}월{day:%d}일(야)'
    ws.freeze_panes=None                             # 원본 시트엔 틀 고정이 없다 (양식에 남은 F8 고정 해제)
    from openpyxl.worksheet.views import Selection
    ws.sheet_view.selection=[Selection(activeCell='A1',sqref='A1')]
    # 다운로드에는 본표(A:BP)만 남긴다. 원본의 우측 보조 양식은 제거한다.
    from openpyxl.utils import column_index_from_string
    last_col=column_index_from_string('BP')
    for merged in list(ws.merged_cells.ranges):
        if merged.min_col>last_col:
            ws.unmerge_cells(str(merged))
    if ws.max_column>last_col:
        ws.delete_cols(last_col+1,ws.max_column-last_col)
    for col in list(ws.column_dimensions):
        if column_index_from_string(col)>last_col:
            del ws.column_dimensions[col]
    ws.print_area='A1:BP24'
    cells=field_cells(record,cp) if cells is None else cells
    ws['A1']='야간 근무 편성표';ws['BD1']='결 제';ws['BG1']='담 당';ws['BL1']='팀 장'
    ws['A3']=f'{day:%Y년 %m월 %d일} {"월화수목금토일"[day.weekday()]}요일 근무조 : {team}조'
    ws['A5']='근무자명'
    names=cells['names']
    if not names[0]:
        names=[name for name in names if name]+['']*names.count('')
    for col,name in zip(FORM_COLUMNS[1:8],names):
        ws[f'{col}5']=name or None
    for col,name in zip(FORM_COLUMNS[1:4],cells['cps']):
        ws[f'{col}6']=name or None
    for col,label in zip(FORM_COLUMNS,HEADERS):
        ws[f'{col}7']=label
    for t,values in enumerate(cells['rows']):
        for col,value in zip(FORM_COLUMNS,values):
            ws[f'{col}{8+t}']=(int(value) if value.isdigit() and col!='A' else value) or None
    left,right=cells['notes']                        # 원본처럼 왼쪽 줄 뒤에 오른쪽 줄을 띄워 쓴다
    lines=[(l.ljust(40)+r) if k<len(right) and (r:=right[k]) else l for k,l in enumerate(left)]
    ws['A22']='특이사항';ws['F22']='\n'.join(lines)
    ws['F22'].alignment=Alignment(horizontal='left',vertical='center',wrap_text=True)
    buffer=BytesIO();wb.save(buffer);return buffer.getvalue()


def download_button(record,key,cp=None):
    """현장 양식 이미지·엑셀 다운로드 (PC·폰 공통, 엑셀은 2026-10-11 추가)."""
    import streamlit as st
    name=f'{record.context.shift_start_date}_야간근무표'
    try:
        st.download_button('🖼 현장 양식 이미지 다운로드',render_png(record,cp),file_name=f'{name}.png',
                           mime='image/png',key=key,width='stretch')
    except Exception as exc:st.warning(f'근무표 이미지를 만들지 못했습니다: {exc}')
    try:                                             # 업로드 옆 좁은 칸이라 위아래로 쌓는다
        st.download_button('📊 현장 양식 엑셀 다운로드',render_xlsx(record,cp),file_name=f'{name}.xlsx',
                           mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                           key=f'{key}_xlsx',width='stretch')
    except Exception as exc:st.warning(f'근무표 엑셀을 만들지 못했습니다: {exc}')
