"""편성표 + 저장 일보 → 보안근무명령서. 정보 없는 칸은 비운다."""
from datetime import timedelta
from pathlib import Path
from io import BytesIO
from html import escape
import re
import openpyxl
from openpyxl.cell.cell import MergedCell
from . import daily_report as dr

TEMPLATE = Path(__file__).parent / 'templates/security_order.xlsx'


def note_parts(text):
    return [(m[1].strip(), int(m[2]), m[3] or '') for m in
            re.finditer(r'([^,()\n]+?)\s*(\d+)\s*명\s*(?:\(([^)]*)\))?', text or '')]


def order_kind(kind):
    """일보 비고 종류 → 원본 명령서 표기."""
    compact = kind.replace(' ', '')
    return {'야간지원': '야간조 지원', '예비군훈련': '훈련'}.get(compact, kind.strip())


def full_name(short, roster, warnings):
    matches = [p['이름'] for p in roster if p['이름'] == short or p['이름'].endswith(short)]
    matches = list(dict.fromkeys(matches))
    if len(matches) == 1: return matches[0]
    warnings.append(f"이름 확인 필요: {short}")
    return ''


def person_kinds(name, rows):
    return [kind for row in rows for kind,_,names in note_parts(row.get('비고','')) if name in names.replace(' ','')]


def present(name, rows, shift='day', off=False):
    kinds = person_kinds(name,rows)
    if any(any(what == 'absent' for what,_ in dr.classify(k)) for k in kinds): return False
    if shift == 'day' and any('야간지원' in k.replace(' ','') for k in kinds): return False
    if off: return any(k.startswith('출근:') or '주간지원' in k.replace(' ','') or '대근' in k or '교육' in k for k in kinds)
    return True


def security_adjust(rowmap):
    """보안 근무인원 보정 = 다른 조에서 온 출근자(야간조 주간교육·주간지원 + 휴무조 출근).
    2026-10-13 사용자 확정: 다른 조 인원은 보안에 더한다. 야간지원 간 사람(최정식 등)도 그날 출근이라
    사고가 아니므로 빼지 않는다(주간 성명 칸에서만 빠짐). 원본 7/1 박정호 주간교육+휴무조 교육 → 18명."""
    came = sum(n for kind, n, _ in note_parts(rowmap[15]['비고'])
               if '주간교육' in kind.replace(' ', '') or '주간지원' in kind.replace(' ', ''))
    off = rowmap[16]
    return came + max(0, off['T/O'] - off['사고'])   # 휴무조 출근 = 쉬는 날 나온 사람(이유 무관)


def clean_cell(post, text):
    if post == '셔틀' and any(x in text for x in ('해지','설치')): return ''
    return re.sub(r'(?:CS)?교육\s*[:：]\s*', '', text, flags=re.I)


def build(day, store, sources, report=None):
    warnings = []
    roster = dr.active(dr.load_roster(store),day)
    report = report or dr.build(day,dr.entries_for(store,day),dr.load_roster(store),dr.entries_for(store,day-timedelta(days=1)))
    rows = report["rows"]
    rowmap = {r['행']:r for r in rows}
    off = dr.is_off_day(day)
    book = openpyxl.load_workbook(TEMPLATE)
    ws = book['휴일' if off else '평일']
    for other in list(book):
        if other != ws: book.remove(other)
    ws.title = f'{day:%y.%m.%d}'
    night_head, night_start, night_cp, night_total = (41,43,51,56) if off else (46,48,56,61)
    day_cp, accident_row, day_total = (21,29,40) if off else (24,34,45)
    # 기준일 템플릿의 근무자와 시간표를 제거한다. 고정 공지·안내·I1 틀만 별도 보관.
    fixed_i1 = {c.coordinate:c.value for row in ws.iter_rows(min_row=20,max_row=21,min_col=7,max_col=34)
                for c in row if not isinstance(c,MergedCell)} if not off else {}
    fixed_guidance = {c.coordinate:c.value for row in ws.iter_rows(min_row=29,max_row=33,min_col=7,max_col=34)
                      for c in row if not isinstance(c,MergedCell)} if not off else {}
    # 원본 시트의 병합·테두리·크기·인쇄 설정을 그대로 둔다. 값만 바꾼다.
    for r in list(range(6,day_total))+list(range(night_start,night_total)):
        for c in range(3,5):
            if not isinstance(ws.cell(r,c),MergedCell): ws.cell(r,c).value=None
        for c in range(7,35):
            cell=ws.cell(r,c)
            if not isinstance(cell,MergedCell) and r != accident_row and not (accident_row < r < day_total):
                cell.value=None
    ws['C1']=f"{day:%Y.%m.%d}({'월화수목금토일'[day.weekday()]}) ~ 익일 07:00"
    ws['C2']='김  성  덕 (D조)'
    for coord,value in {**fixed_i1,**fixed_guidance}.items(): ws[coord]=value
    merged_anchors = {(r,c):(m.min_row,m.min_col) for m in ws.merged_cells.ranges
                      for r in range(m.min_row,m.max_row+1) for c in range(m.min_col,m.max_col+1)}
    time_values = {}

    def put_time(r,c,value,post):
        # 원본의 세로 병합이 대기/휴게 등 두 POST를 묶으면 각 값과 POST를
        # 한 병합 칸에 함께 남겨 값 누락·덮어쓰기를 막는다.
        anchor=merged_anchors.get((r,c),(r,c))
        time_values.setdefault(anchor,{})[(r,c)]=(post,value or '')
    teams=report['teams']
    source_cp_counts={}
    for shift,head,start,cp_start,end in [('day',4,6,day_cp,day_total),('night',night_head,night_start,night_cp,night_total)]:
        source=sources.get(shift)
        if not source:
            warnings.append(f"{'주간' if shift=='day' else '야간'} 편성표 없음: 이름·시간표 미작성")
            continue
        warnings.extend(source.get('basis_warnings',[]))
        expected=teams['주간조' if shift=='day' else '야간조'][0]
        if source['team'] and source['team'] != expected:
            warnings.append(f"{shift} 편성표 {source['team']}조 / 일보 {expected}조: 조 확인 필요")
        names=[]
        numbered=set(source['names'].values())
        if shift=='day':
            for p in roster:
                if p['소속']=='주간고정' and p['이름'][-2:] not in numbered and not p['직책'].startswith('통제'):
                    if present(p['이름'],rows,off=off): names.append((p['이름'],p['직책']))
        for control in source['controls']:
            matches=[p for p in roster if p['직책']=='통제'+control]
            if len(matches)==1:
                # 편성표 머리에 통제 표기가 남아 있어도 그날 연차·휴무면 성명에서 뺀다 (10/8 통제E 최정식 연차)
                fixed_staff=matches[0]['소속']=='주간고정'   # 주말·공휴일 미출근 규칙은 주간고정만
                if shift=='day' and not present(matches[0]['이름'],rows,off=off and fixed_staff):
                    warnings.append(f"통제{control} {matches[0]['이름']}: 일보상 출근하지 않아 명령서 성명에서 뺐습니다")
                    continue
                names.append((matches[0]['이름'],'통제'+control))
            else: warnings.append('통제'+control+' 이름 확인 필요')
        # 명단에 별도 통제가 없어도 주간조 조장 표기를 추정해서 추가하지 않는다.
        seen_numbers={}
        for number,short in sorted(source['names'].items(),key=lambda x:int(x[0])):
            if short in seen_numbers:
                warnings.append(f"{'주간' if shift=='day' else '야간'} 편성표 {seen_numbers[short]}번·{number}번 이름이 모두 '{short}'입니다. "
                                "아래 '편성표 번호 이름 고치기'에서 실제 이름으로 바꾸세요.")
            seen_numbers.setdefault(short,number)
            names.append((full_name(short,roster,warnings),f'대원({number})'))
        if source.get('name_fixes'):
            warnings.append(f"{'주간' if shift=='day' else '야간'} 편성표 이름 고침 적용: "
                            +', '.join(f'{k}번 {v}' for k,v in sorted(source['name_fixes'].items(),key=lambda x:int(x[0]))))
        if len(names)>cp_start-start:
            warnings.append(f"{shift} 보안 명단이 양식 칸보다 많아 초과 인원 확인 필요")
        for r,(name,role) in zip(range(start,cp_start),names): ws.cell(r,3,name); ws.cell(r,4,role)
        for i,short in source['cp'].items():
            ws.cell(cp_start+int(i)-1,3,full_name(short,roster,warnings))
            ws.cell(cp_start+int(i)-1,4,'CP'+i)
        source_cp_counts[shift]=len(source['cp']) if source['cp'] else None
        if not source['cp']: warnings.append(f'{shift} CP 명단 없음: CP 시간표 확인 필요')
        # POST 행은 성명 행과 독립이다.
        for r in range(start,end):
            label=re.sub(r'\s+','',str(ws.cell(r,5).value or ''))
            post=(re.search(r'\(([^)]+)\)',label)[1] if re.search(r'\(([^)]+)\)',label) else label).upper()
            post={'지하층순찰':'순찰','휴게실':'휴게','순회/셔틀':'순회'}.get(post,post)
            if shift=='day' and post=='I1' and not off: continue
            if not post or post=='사고내용': continue
            for h,values in source['cells'].items():
                h=int(h)
                col=7+2*(h-6 if shift=='day' else (h-18)%24)
                if not 7<=col<=33: continue
                value=clean_cell(post,values.get(post,''))
                if not value and post=='CP' and source_cp_counts[shift] in (2,3):
                    n=source_cp_counts[shift]
                    if shift=='day' and 8<=h<=17:
                        pattern=(['CP1.CP2','CP1.CP3','CP2.CP3']*4 if off and n==3 else
                                 ['CP1.CP2','CP2.CP3','CP3','CP1.CP2','CP2.CP3','CP1.CP3','CP1','CP2.CP3','CP1.CP3','CP1.CP2'] if n==3 else
                                 ['CP1.CP2','CP1','CP2','CP1.CP2','CP1','CP2','CP1','CP1','CP2','CP1'])
                        value=pattern[h-8]
                    elif shift=='night' and n==3:
                        from .cp_schedule import hourly_label
                        value=hourly_label((h-18)%24)
                put_time(r,col,value,post)
        if shift=='day' and not off and source_cp_counts[shift] in (2,3):
            put_time(15,29,f"CP{source_cp_counts[shift]}",'셔틀') # 17시 셔틀
        if shift=='night' and source_cp_counts[shift]==3: put_time(end-1,33,'CP3','셔틀')
    for (r,c),items in time_values.items():
        values=list(items.values())
        unique=list(dict.fromkeys(value for _,value in values if value))
        if len(set(value for _,value in values))>1:
            ws.cell(r,c).value='\n'.join(f'{post}: {value}' for post,value in values if value)
        else: ws.cell(r,c).value=unique[0] if unique else None
    # 안내·주차 성명은 해당 날짜 출근자만. 고정 시간 틀은 사용자 확정 사항.
    for group,start,capacity in [('안내',29 if not off else 24,5),('주차',36 if not off else 31,4)]:
        persons=[p for p in roster if p['소속']==group and present(p['이름'],rows,off=off)]
        if len(persons)!=rowmap[20 if group=='안내' else 21]['T/O']-rowmap[20 if group=='안내' else 21]['사고']:
            warnings.append(group+' 출근 명단과 일보 인원 불일치: 확인 필요')
        for i,p in enumerate(persons[:capacity]):
            ws.cell(start+i,3,p['이름']); ws.cell(start+i,4,p['직책'] if group=='안내' and i==0 else f'대원({i if group=="안내" else i+1})')
    # CP를 일보 조 인원에서 분리. CP 명단이 없으면 해당 합계를 임의로 계산하지 않는다.
    cp_day=source_cp_counts.get('day')
    # 보안 = 주간고정·주간조 출근 − CP + 다른 조에서 주간에 온 사람(교육·지원). CP는 보안에 넣지 않는다(주말도, 2026-10-13 사용자 확정).
    security=(rowmap[12]['T/O']-rowmap[12]['사고']+rowmap[13]['T/O']-rowmap[13]['사고']-cp_day+security_adjust(rowmap)) if cp_day is not None else None
    # 원본 명령서 작성 방식(7~9월 대조, 2026-10-07): 결원은 합쳐 맨 앞, 같은 종류는 한 묶음(이름은 ','),
    # 빠진 사람 줄 다음 줄에 지원·반차·교육 등 출근자 비고, '이사님'은 빼고 이름만.
    cp_names={p['이름'] for p in roster if p['직책']=='CP'}
    vacancy={12:0,13:0,16:0}; absent={}; attending={}; cp_merged={}
    for r in rows:
        if r['행'] not in (12,13,16): continue
        for kind,n,names in note_parts(r['비고']):
            kind=order_kind(kind)
            people=[re.sub(r'\s*이사님$','',x.strip()) for x in names.split(',') if x.strip()]
            if kind=='결원':
                vacancy[r['행']]+=n; continue
            for person in people:
                group=cp_merged if person in cp_names else absent if any(w=='absent' for w,_ in dr.classify(kind)) else attending
                if person not in group.setdefault(kind,[]): group[kind].append(person)
    phrase=lambda merged:[f'{k} {len(v)}명({",".join(v)})' for k,v in merged.items() if v]
    first=([f'결원 {sum(vacancy.values())}명'] if sum(vacancy.values()) else [])+phrase(absent)
    off_people=rowmap[16]['사고']-vacancy[16]+((rowmap[12]['사고']-vacancy[12]) if off else 0)
    security_line=f"○ 보안 : - 주간조: {teams['주간조']} / 야간조: {teams['야간조']}\n            - "
    if security is not None: security_line+=f'근무인원: {security}명 / '
    security_line+=f"휴무인원: {off_people}명"
    for line in (first,phrase(attending)):
        if line: security_line+='\n            - '+', '.join(line)
    accident=security_line+'\n\n○ CP : '+(f'- 근무인원: {cp_day}명' if cp_day is not None else '')
    cp_notes=phrase(cp_merged)
    if cp_notes: accident+=', '+', '.join(cp_notes)
    for group,r in [('안내',20),('주차',21)]:
        row=rowmap[r]; n=row['T/O']-row['사고']
        accident+=f'\n\n○ {group} : - '+(f'휴무 {row["사고"]}명' if off and n==0 else f'근무인원: {n}명')
        if row['비고']: accident+=', '+row['비고']
    ws.cell(accident_row,6,accident)
    total=security+cp_day+sum(rowmap[r]['T/O']-rowmap[r]['사고'] for r in (20,21)) if security is not None else None
    ws.cell(day_total,3).value=total
    ws.cell(day_total,7).value=f'○ 근무인원: {total}명' if total is not None else None
    night_n=rowmap[15]['T/O']-rowmap[15]['사고']
    ws.cell(night_total,3).value=night_n
    ws.cell(night_total,7).value=f"○ 근무인원: {night_n}명"+(f" ○{rowmap[15]['비고']}" if rowmap[15]['비고'] else '')
    return book, ws, list(dict.fromkeys(warnings))


def excel_bytes(book):
    out=BytesIO(); book.save(out); return out.getvalue()


def sheet_html(ws):
    spans={}; hidden=set()
    for merged in ws.merged_cells.ranges:
        if merged.max_col>34: continue
        spans[(merged.min_row,merged.min_col)]=(merged.max_row-merged.min_row+1,merged.max_col-merged.min_col+1)
        hidden.update((r,c) for r in range(merged.min_row,merged.max_row+1) for c in range(merged.min_col,merged.max_col+1)
                      if (r,c)!=(merged.min_row,merged.min_col))
    last=56 if ws.max_row<61 else 61
    parts=["<div style='overflow-x:auto'><table style='border-collapse:collapse;table-layout:fixed;width:1400px;background:white;color:#111;font-size:12px'>"]
    for r in range(1,last+1):
        parts.append(f"<tr style='height:{max(20,ws.row_dimensions[r].height or 20)}px'>")
        for c in range(1,35):
            if (r,c) in hidden: continue
            rs,cs=spans.get((r,c),(1,1)); cell=ws.cell(r,c)
            value=escape(str(cell.value if cell.value is not None else '')).replace('\n','<br>')
            bg='#eee' if cell.fill.fgColor.type=='rgb' and cell.fill.fgColor.rgb not in ('00000000','00FFFFFF','FFFFFFFF') else '#fff'
            borders=[]
            for side in ('top','bottom','left','right'):
                edge=getattr(cell.border,side)
                width=2 if edge and edge.style in ('medium','thick','double') else 1
                line='double' if edge and edge.style=='double' else 'dashed' if edge and edge.style in ('dashed','dotted') else 'solid'
                borders.append(f'border-{side}:{width}px {line} #111' if edge and edge.style else f'border-{side}:none')
            parts.append(f"<td rowspan='{rs}' colspan='{cs}' style='{';'.join(borders)};padding:2px;white-space:normal;overflow-wrap:anywhere;vertical-align:middle;text-align:{cell.alignment.horizontal or 'center'};background:{bg}'>{value}</td>")
        parts.append('</tr>')
    return ''.join(parts)+ '</table></div>'
