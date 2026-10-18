"""주말·공휴일 주간 현장 양식 이미지·엑셀 (2026-10-12 사용자 요청). 평일 주간과 같은 방식.

양식은 원본 편성표의 휴일 주간 시트(26년09월19일(주))를 서식 그대로 복사한 templates/weekend_day_form.xlsx
(tools/make_day_form.ps1). 이미지는 엑셀이 그 양식을 그린 빈 그림(weekend_form_blank.png)과 칸 위치·글꼴
(weekend_form_geometry.json, tools/export_day_form_geometry.ps1) 위에 글자만 쓴다(day_form_draw). 원본 휴일 시트에는
번호 색 규칙이 없어 칸을 칠하지 않는다. 엑셀은 양식에 같은 칸 내용만 넣는다. 이미지·엑셀 칸 계산은 build_spec 한 곳.
"""
from __future__ import annotations

import re
from datetime import date
from io import BytesIO
from pathlib import Path

TEMPLATES = Path(__file__).resolve().parent / "templates"
FORM = TEMPLATES / "weekend_day_form.xlsx"
BLANK = TEMPLATES / "weekend_form_blank.png"
GEOMETRY = TEMPLATES / "weekend_form_geometry.json"
WEEKDAYS = "월화수목금토일"
COLUMNS = ("G1", "G2", "G3", "I1", "L1", "L2", "CP", "순찰", "대기", "휴게")
ALIASES = {"I1(데스크)": "I1", "보안실/대기": "대기"}
# 엑셀에 같은 인원 휴일 형태가 없을 때 (원본 9/13·9/19 형태): CP 칸 순서, 나머지 CP는 휴게
CP_CYCLE = {3: ("CP1.CP2", "CP1.CP3", "CP2.CP3"), 2: ("CP1.CP2", "CP1", "CP2"), 1: ("CP1",)}


def available() -> bool:
    from .day_form_draw import FONT
    return FORM.exists() and BLANK.exists() and GEOMETRY.exists() and FONT[False].exists()


def _norm(value) -> str:
    text = re.sub(r"\s+", "", str(value or "")).upper()
    return ALIASES.get(text, text)


def layout(form: Path = FORM) -> dict:
    """양식 칸: 제목 날짜 칸, 명단 두 줄의 칸들, 머리글(포스트)별 열, 시간별 행."""
    import openpyxl
    ws = openpyxl.load_workbook(form).worksheets[0]
    header = next(r for r in range(1, 20) if _norm(ws.cell(r, 1).value) == "근무시간")
    columns = {}
    for c in range(2, ws.max_column + 1):
        label = _norm(ws.cell(header, c).value)
        if label in COLUMNS and label not in columns:
            columns[label] = c
    names = (header - 2, header - 1)
    starts = sorted(m.min_col for m in ws.merged_cells.ranges
                    if m.min_row == names[0] and m.max_row == names[0] and 2 <= m.min_col
                    and m.min_col <= _print_last(ws))
    hours = {}
    for r in range(header + 1, header + 20):
        m = re.fullmatch(r"(\d{1,2}):00", str(ws.cell(r, 1).value or "").strip())
        if m and int(m[1]) <= 17:
            hours[int(m[1])] = r
    date_cell = next(c.coordinate for row in ws.iter_rows(min_row=1, max_row=header)
                     for c in row if isinstance(c.value, str) and "근무조" in c.value)
    return {"names": names, "slots": starts or sorted(columns.values()), "columns": columns, "hours": hours,
            "date_cell": date_cell}


def _print_last(ws) -> int:
    from openpyxl.utils import column_index_from_string
    area = str(ws.print_area or "")
    m = re.search(r"\$?([A-Z]+)\$?\d+$", area)
    return column_index_from_string(m[1]) if m else ws.max_column


def _cell(row: int, col: int) -> str:
    from openpyxl.utils import get_column_letter
    return f"{get_column_letter(col)}{row}"


def _cp_hours(store, day: date, cp: dict) -> list[str]:
    from .calendar_info import is_off_day
    from .cp_schedule import standard_pattern
    if not cp:
        return [""] * 10
    found = standard_pattern(Path(store), "day", list(cp), day, keep=is_off_day)
    if found:
        return found["hours"]
    numbers = sorted(cp, key=int)
    relabel = {str(i + 1): n for i, n in enumerate(numbers)}
    cycle = CP_CYCLE.get(len(numbers), ("",))
    return [re.sub(r"CP(\d)", lambda m: f"CP{relabel[m[1]]}", cycle[t % len(cycle)]) for t in range(10)]


def cells_for(schedule, store, cp: dict | None, review: dict | None = None):
    """(명단 위 줄, 아래 줄, {시: {머리글: 글자}}). review = 업로드 원본 칸(손대지 않은 저장본)이면 그 칸 그대로."""
    from .cp_schedule import short
    from .schedule_image_import import normalize_cell
    day = date.fromisoformat(str(schedule.day))
    control = "E" if str(schedule.names.get("A", "")).startswith("통제E") else "A"
    if review and len(review.get("rows") or []) == 10:
        names = {int(k): v for k, v in (review.get("mapping") or {}).items() if str(k).isdigit() and v}
        cps = {k[2:]: v for k, v in (review.get("extra_names") or {}).items() if k.startswith("CP") and v} or (cp or {})
        controls = [c for c in review.get("controls") or [] if c in "ADE"] or [control]
        hours = {}
        for t, row in enumerate(review["rows"]):
            hour = {}
            for key, value in row.items():
                label = _norm(key)
                if label not in COLUMNS or not value:
                    continue
                text = normalize_cell(str(value))
                if control == "E":                       # 가져올 때 통제 자리로 바꾼 E를 원본 표기로
                    text = re.sub(r"(?<![A-Z])A(?![A-Z0-9])", "E", text)
                hour[label] = text
            hours[8 + t] = hour
    else:
        numbers = sorted((p for p in schedule.grid if p != "A"), key=int)
        names = {int(p): schedule.names.get(p, p) for p in numbers}
        cps = cp or {}
        controls = [control] if "A" in schedule.grid else []
        cp_text = _cp_hours(store, day, cps)
        hours = {}
        for t in range(10):
            col = {p: row[t] for p, row in schedule.grid.items()}

            def who(label):
                lead = [control] if col.get("A") == label else []
                return lead + [p for p in numbers if col[p] == label]
            g2 = who("G2")
            hour = {"G1": who("G1"), "G2": g2[:1], "G3": g2[1:], "I1": who("I1"), "L1": who("L1"),
                    "순찰": who("순찰"), "대기": who("대기"), "휴게": who("휴게")}   # 대기: 2026-10-18 보안실 대기
            on = {n for n in re.findall(r"CP(\d)", cp_text[t])}
            hour = {k: ".".join(v) for k, v in hour.items() if v}
            hour["CP"] = cp_text[t]
            rest = [f"CP{n}" for n in sorted(cps, key=int) if n not in on]
            hour["휴게"] = ".".join(x for x in (hour.get("휴게", ""), ".".join(rest)) if x)
            hours[8 + t] = {k: v for k, v in hour.items() if v}
    tags = [f"{n}{names[n]}" for n in sorted(names)]
    top = [f"통제{c}" for c in controls] + tags
    bottom = [f"CP{k}{short(v)}" for k, v in sorted(cps.items())]
    return top, bottom, hours


def build_spec(schedule, store, cp: dict | None, review: dict | None = None, team: str = "A",
               form: Path = FORM) -> dict:
    lay = layout(form)
    clear, values = [], {}
    for r in lay["names"]:
        clear += [_cell(r, c) for c in lay["slots"]]
    for r in lay["hours"].values():
        clear += [_cell(r, c) for c in lay["columns"].values()]
    d = date.fromisoformat(str(schedule.day))
    values[lay["date_cell"]] = f"{d.year}년 {d.month:02d}월 {d.day:02d}일 {WEEKDAYS[d.weekday()]}요일 근무조 : {team}조"
    top, bottom, hours = cells_for(schedule, store, cp, review)
    for r, names in zip(lay["names"], (top, bottom)):
        for c, name in zip(lay["slots"], names):
            values[_cell(r, c)] = name
    for hour, row in lay["hours"].items():
        for label, text in hours.get(hour, {}).items():
            col = lay["columns"].get(label)
            if col and text:
                values[_cell(row, col)] = text
    return {"clear": clear, "values": values}


def draw(schedule, store, cp: dict | None, review: dict | None = None) -> bytes:
    from .day_form_draw import draw as paint
    return paint(build_spec(schedule, store, cp, review)["values"], blank_path=BLANK, geometry_path=GEOMETRY,
                 colored=False)


def to_xlsx(schedule, store, cp: dict | None, review: dict | None = None, team: str = "A") -> bytes:
    """원본 휴일 주간 시트와 같은 엑셀 (양식 서식 그대로, 이미지와 같은 칸 내용만)."""
    import openpyxl
    from openpyxl.cell.cell import MergedCell
    spec = build_spec(schedule, store, cp, review, team)
    wb = openpyxl.load_workbook(FORM)
    ws = wb.worksheets[0]
    d = date.fromisoformat(str(schedule.day))
    ws.title = f"{d:%y}년{d:%m}월{d:%d}일(주)"
    for address in spec["clear"]:
        if not isinstance(ws[address], MergedCell):
            ws[address] = None
    for address, text in spec["values"].items():
        ws[address] = int(text) if text.isdigit() else text   # 원본처럼 번호 하나는 숫자로
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
