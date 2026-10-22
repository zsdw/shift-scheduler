"""조별 조회의 현장 양식 출력. 기존 A조와 같은 템플릿·렌더러를 사용한다."""
from datetime import date
from types import SimpleNamespace
from io import BytesIO
from team_schedule_analysis import people,hours


def original_notes(source):
    """같은 업로드 원본의 본표 아래 특이사항을 보존한다."""
    if not source.get('_original_path'):return {}
    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string
    wb=load_workbook(source['_original_path'],read_only=True,data_only=True)
    try:
        if source.get('sheet') not in wb.sheetnames:return {}
        ws=wb[source['sheet']]
        # 원본 시간표 마지막 행 다음의 본표 영역만 읽는다(우측 보조 양식 제외).
        last=0
        for row in ws.iter_rows(min_col=1,max_col=1,max_row=35):
            v=str(row[0].value or '').strip()
            if v in ('07:00','17:00'):last=row[0].row
        if not last:return {}
        limit=column_index_from_string('BP' if source['shift']=='night' else 'BU')
        return {c.coordinate:str(c.value) for row in ws.iter_rows(min_row=last+1,max_row=min(last+5,ws.max_row),max_col=limit)
                for c in row if isinstance(c.value,str) and c.value.strip() and c.value.strip()!='목차'}
    finally:wb.close()


def night_cells(source):
    from night_shift.schedule_image import HEADERS
    names=[('통제'+c) for c in source.get('controls',[]) if c in people(source)]
    names += [n+v for n,v in source['names'].items()]
    rows=[]
    for h in hours(source):
        cells=source['cells'].get(str(h),{})
        values=[]
        for p in HEADERS[1:]:
            value=cells.get(p,'')
            if p=='휴 게':value=cells.get('휴게',cells.get('휴 게',''))
            if p=='보안실':value=cells.get('보안실',cells.get('대기',''))
            if p=='셔틀':value=cells.get('셔틀',cells.get('순회',''))
            values.append(str(value or ''))
        rows.append([f'{24 if h==0 else h:02}:00',*values])
    return {'names':(names+['']*7)[:7],
            'cps':([f'CP{n}{name}' for n,name in source.get('cp',{}).items()]+['']*3)[:3],
            'rows':rows,'notes':source.get('field_notes',([],[]))}


def day_spec(source):
    from night_shift.calendar_info import is_off_day
    from night_shift import day_excel_image as weekday,weekend_form as weekend
    module=weekend if is_off_day(date.fromisoformat(source['day'])) else weekday
    lay=module.layout();groups=lay.get('slots') or sorted(lay['columns'].values())
    from openpyxl.utils import get_column_letter as letter
    cell=lambda c,r:f'{letter(c)}{r}'
    values={};clear=[]
    for r in lay['names']:clear.extend(cell(c,r) for c in groups)
    for r in lay['hours'].values():clear.extend(cell(c,r) for c in lay['columns'].values())
    day=date.fromisoformat(source['day'])
    values[lay['date_cell']]=f'{day:%Y년 %m월 %d일} {"월화수목금토일"[day.weekday()]}요일 근무조 : {source["team"]}조'
    tags=['통제'+c for c in source.get('controls',[])]+[n+v for n,v in source['names'].items()]
    cp=[f'CP{n}{v}' for n,v in source.get('cp',{}).items()]
    if module is weekend:top,bottom=tags,cp
    else:top,bottom=tags[:len(groups)],tags[len(groups):]+cp
    for r,names in zip(lay['names'],(top,bottom)):
        for c,name in zip(groups,names):values[cell(c,r)]=name
    for h,r in lay['hours'].items():
        for p,v in source['cells'].get(str(h),{}).items():
            p=p.replace(' ','')
            if p=='보안실':p='대기'
            c=lay['columns'].get(p)
            if c:values[cell(c,r)]=str(v or '')
    return module,values,clear


def export(source,kind):
    from team_schedule_analysis import normalize_source
    source=normalize_source(source)
    if source['shift']=='night':
        from night_shift import schedule_image as night
        record=SimpleNamespace(context=SimpleNamespace(shift_start_date=date.fromisoformat(source['day'])))
        cells=night_cells(source)
        notes=original_notes(source)
        import re
        left=[];right=[]
        for text in notes.values():
            if text.strip()=='특이사항':continue
            for line in text.splitlines():
                parts=re.split(r'\s{3,}',line.strip(),maxsplit=1)
                left.append(parts[0])
                if len(parts)>1:right.append(parts[1])
        if left:cells['notes']=(left[:5],right[:2])
        return (night.render_png if kind=='png' else night.render_xlsx)(record,team=source['team'],cells=cells)
    module,values,clear=day_spec(source)
    values.update(original_notes(source))
    if kind=='png':
        from night_shift.day_form_draw import draw
        from night_shift import weekend_form
        return (draw(values,blank_path=module.BLANK,geometry_path=module.GEOMETRY,colored=False)
                if module is weekend_form else draw(values))
    from openpyxl import load_workbook
    from openpyxl.cell.cell import MergedCell
    wb=load_workbook(module.FORM);ws=wb.active
    for address in clear:
        if not isinstance(ws[address],MergedCell):ws[address]=None
    for address,value in values.items():
        if not isinstance(ws[address],MergedCell):ws[address]=int(value) if value.isdigit() else value or None
    out=BytesIO();wb.save(out);wb.close();return out.getvalue()
