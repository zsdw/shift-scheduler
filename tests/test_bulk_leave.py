"""2026-10-04: 편성표 엑셀 일괄 가져오기 · 연차일정 기본 상태 · 상태 '기타'(근무 제외 + 비고)."""
import io
import os
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from streamlit.testing.v1 import AppTest

from night_shift import bulk_import, history_store, leave_schedule
from night_shift.balance_notes import balance_reasons
from night_shift.service import calculate_all_stats

APP = str(Path(__file__).resolve().parents[1] / "app.py")
# 10/3 야간 현장 양식 (표시번호 1만년 2원종 3제욱 4광희 5성호 6준호, A=통제)
NIGHT = {  # 시간: (G1, L3, I1, 순찰, 보안실)
    18: ("2", "3", "5", "", "4"), 19: ("1", "4", "6", "", "3"), 20: ("6", "2", "5", "", "A.1"),
    21: ("1", "3", "2", "A", "6"), 22: ("3", "6", "5", "4", "1"), 23: ("5", "2", "1", "6", "3"),
    0: ("3", "A", "2", "", "4"), 1: ("4", "2", "5", "3", "1"), 2: ("A", "6", "1", "5", "2"),
    3: ("1", "4", "A", "2", "6"), 4: ("6", "1", "3", "", "4"), 5: ("5", "3", "6", "", "1"),
    6: ("4", "5", "A", "", "3"), 7: ("1", "6", "4", "", "5")}
DAY_NAMES = ["도운", "준형", "용주", "준호", "성호", "광희", "만년", "제욱", "원종"]
DAY_LABELS = ["G1", "G2", "G3", "L1", "L2", "C1", "C2", "순찰", "휴게"]


def night_sheet(wb, title, team="A", manager="통제A", change=None):
    ws = wb.create_sheet(title)
    ws.cell(3, 1, f"2026년 10월 03일 토요일 근무조 : {team}조")
    for col, name in zip(range(6, 42, 5), [manager, "1만년", "2원종", "3제욱", "4광희", "5성호", "6준호"]):
        ws.cell(5, col, name)
    ws.cell(5, 1, "근무자명")
    heads = ["근무시간", "G1", "L3", "I1", "G3", "G2", "G2(1)", "G2(2)", "CP", "셔틀", "순찰", "보안실", "휴 게"]
    for i, head in enumerate(heads):
        ws.cell(7, 1 + 5 * i, head)
    for r, hour in enumerate([*range(18, 24), *range(0, 8)]):
        g1, l3, i1, patrol, room = NIGHT[hour]
        if manager == "통제E":
            g1, l3, i1, patrol, room = (v.replace("A", "E") for v in (g1, l3, i1, patrol, room))
        if change and change[0] == hour:
            g1 = change[1]
        ws.cell(8 + r, 1, f"{(hour if hour else 24):02d}:00")
        for col, value in ((6, g1), (11, l3), (16, i1), (51, patrol), (56, room)):
            if value:
                ws.cell(8 + r, col, value)


def day_sheet(wb, title):
    ws = wb.create_sheet(title)
    ws.cell(3, 1, "2026년 10월 07일 수요일 근무조 : A조")
    ws.cell(5, 1, "근무자명")
    ws.cell(5, 5, "통제A")
    for i, name in enumerate(DAY_NAMES):
        ws.cell(5 + i // 5, 10 + 5 * (i % 5), f"{i + 1}{name}")
    heads = {"근무시간": 1, "G1": 5, "G2": 10, "G3": 15, "4C": 20, "L1": 25, "L2": 30, "C1": 35, "C2": 40,
             "셔틀": 55, "순찰": 60, "대기": 65, "휴게": 69}
    for head, col in heads.items():
        ws.cell(7, col, head)
    for t in range(10):
        ws.cell(8 + t, 1, f"{t + 8:02d}:00")
        for i in range(9):
            label = DAY_LABELS[(i + t) % 9]
            col = heads[label]
            old = ws.cell(8 + t, col).value
            ws.cell(8 + t, col, f"{old}.{i + 1}" if old else str(i + 1))


def workbook(change=None):
    wb = Workbook()
    wb.remove(wb.active)
    night_sheet(wb, "26년10월03일(야)", change=change)
    night_sheet(wb, "26년10월03일(주)", team="C")                  # 다른 조 → 무시
    night_sheet(wb, "26년10월04일(야)", manager="통제E")            # 통제E가 통제 자리
    day_sheet(wb, "26년10월07일(주)")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


LEAVE = """10월 03일(토)

A조 이성호 사원 - 연차
A조 김광희 사원 - 예비군 훈련
A조 백원석 조장 - 연차
C조 홍길동 사원 - 연차

10월 07일(수)

D조 이준형 사원 - 연차
A조 김광희 사원 - 교육
A조 이준호 선임 - 반차
A조 유원종 사원 - 입사
"""


class LeaveScheduleTests(unittest.TestCase):
    def test_parse_and_classify(self):
        entries = leave_schedule.parse(LEAVE)
        night, notices = leave_schedule.defaults_for(entries, date(2026, 10, 3), "night",
                                                     ["만년", "원종", "제욱", "광희", "성호", "준호"])
        self.assertEqual(night, {"성호": ("연차", ""), "광희": ("기타", "예비군 훈련"), "A": ("연차", "")})
        day, notices = leave_schedule.defaults_for(entries, date(2026, 10, 7), "day", ["준형", "광희", "준호", "원종"])
        self.assertEqual(day, {"준형": ("연차", ""), "광희": ("기타", "교육")})       # 입사 줄은 상태 아님
        self.assertTrue(any("반차" in n for n in notices))                          # 전·후반 모르는 반차는 안내만
        self.assertEqual(leave_schedule.classify("오후 반차", "night")[0], None)     # 주간 표현은 야간에 단정 안 함
        self.assertEqual(leave_schedule.classify("후반부 반차", "night")[0], leave_schedule.NIGHT_LATE)

    def test_other_shows_in_notes(self):
        from night_shift.models import Mode, NightType
        from night_shift.service import HistorySelection, ShiftRequest, SolveOptions, StaffInput, run_schedule
        staff = [StaffInput("A", "통제", True, False, None, True)]
        staff += [StaffInput(p, n, p != "4", True, int(p) if p != "4" else None, False)
                  for p, n in {"1": "만년", "2": "원종", "3": "제욱", "4": "광희", "5": "성호", "6": "준호"}.items()]
        result = run_schedule(ShiftRequest(date(2026, 10, 3), NightType.N4, Mode.FM, tuple(staff)),
                              HistorySelection(), SolveOptions(rules_version="2.20", seconds=20))
        text = " ".join(balance_reasons(result.record, result.employee_stats, others={"4": "예비군"}))
        self.assertIn("기타(근무 제외): 광희(예비군)", text)
        self.assertNotIn("연차: 광희", text)


class BulkImportTests(unittest.TestCase):
    def test_night_newcomer_from_personnel_notice_preserves_joint_source(self):
        from night_shift.daily_report import save_roster
        wb=Workbook();wb.remove(wb.active);night_sheet(wb,'26년 10월 03일 (야)')
        ws=wb.worksheets[0];ws.cell(6,6,'7지훈')
        for r in range(8,22):ws.cell(r,6,str(ws.cell(r,6).value)+'.7')
        stream=io.BytesIO();wb.save(stream)
        with tempfile.TemporaryDirectory() as folder:
            store=Path(folder)
            save_roster(store,[{'이름':'이지훈','소속':'C조','직책':'대원','시작':'2026-10-02'}])
            leave_schedule.save_text(store,'10월 2일\nC조 이지훈 사원 - 입사')
            plans=bulk_import.plan(stream.getvalue(),store,store/'day')
            self.assertEqual(plans[0].status,'new',plans[0].reason)
            self.assertEqual(plans[0].source_review['newcomers'],['7'])
            self.assertEqual(bulk_import.apply(plans,store,store/'day'),1)
            data=history_store.load(store,date(2026,10,3))
            self.assertEqual(data['source_review']['rows'][0]['G1'],'2.7')
            again=bulk_import.plan(stream.getvalue(),store,store/'day')
            self.assertEqual(again[0].status,'same')


    def test_force_upload_replaces_final_and_preserves_filename(self):
        from night_shift.day_shift import load_day, save_day
        from history_ui import upload_saved_label
        with tempfile.TemporaryDirectory() as night, tempfile.TemporaryDirectory() as day:
            kwargs=dict(file_name='최종편성표.xlsx',day_roster=DAY_NAMES,force=True)
            plans=bulk_import.plan(workbook(),Path(night),Path(day),**kwargs)
            self.assertEqual(bulk_import.apply(plans,Path(night),Path(day)),3)
            original=load_day(Path(day),date(2026,10,7))['schedule']
            original_grid={p:list(row) for p,row in original.grid.items()}
            worker=next(iter(original.grid))
            original.grid[worker][0]='휴게'
            save_day(Path(day),original,'수동 수정')
            history_store.save_edits(Path(night),date(2026,10,3),[{'type':'swap','a':['1',0],'b':['2',0]}])
            again=bulk_import.plan(workbook(),Path(night),Path(day),**kwargs)
            self.assertEqual({p.status for p in again},{'changed','same'})
            self.assertEqual(bulk_import.apply(again,Path(night),Path(day)),2)
            for data in (load_day(Path(day),date(2026,10,7)),history_store.load(Path(night),date(2026,10,3))):
                self.assertEqual(data['source_file'],'최종편성표.xlsx')
                self.assertEqual(data['final_save_method'],'upload')
                self.assertIn('업로드 파일로',upload_saved_label(data))
            self.assertEqual(load_day(Path(day),date(2026,10,7))['schedule'].grid,original_grid)
            self.assertEqual(history_store.load(Path(night),date(2026,10,3))['edits'],[])
            self.assertTrue(history_store.list_trash(Path(day)))
            self.assertTrue(history_store.list_trash(Path(night)))

    def test_plan_apply_same_changed(self):
        with tempfile.TemporaryDirectory() as night, tempfile.TemporaryDirectory() as day:
            plans = bulk_import.plan(workbook(), Path(night), Path(day), day_roster=DAY_NAMES + ["성욱"],
                                     start=date(2026, 9, 1))
            status = {(p.day, p.shift): (p.status, p.reason) for p in plans}
            self.assertEqual(len(plans), 3)                                     # C조 시트는 무시
            self.assertEqual(status[date(2026, 10, 3), "night"][0], "new")
            self.assertEqual(status[date(2026, 10, 4), "night"][0], "new")     # 통제E 야간도 읽힘
            self.assertEqual(status[date(2026, 10, 7), "day"][0], "new", status)
            self.assertEqual(bulk_import.apply(plans, Path(night), Path(day)), 3)
            saved = history_store.actual_record(Path(night), date(2026, 10, 3))
            self.assertEqual(next(r for r in saved.assignments if r.employee_id == "1" and r.slot.index == 2)
                             .barrier_install_flag, True)                       # 20시 보안실 1번 = 설치
            self.assertIn("통제E(지원)", history_store.load(Path(night), date(2026, 10, 4))["note"])
            self.assertTrue(calculate_all_stats(saved))
            again = bulk_import.plan(workbook(), Path(night), Path(day), day_roster=DAY_NAMES + ["성욱"], start=date(2026, 9, 1))
            self.assertEqual({p.status for p in again}, {"same"})               # 다시 올려도 다시 만들지 않음
            changed = bulk_import.plan(workbook(change=(18, "3")), Path(night), Path(day), day_roster=DAY_NAMES + ["성욱"],
                                       start=date(2026, 9, 1))
            self.assertEqual({(p.day, p.status) for p in changed if p.status != "same"},
                             {(date(2026, 10, 3), "skip")} | set())             # 3번이 G1·L3 중복 → 읽지 않음
            from night_shift.day_shift import load_day
            schedule = load_day(Path(day), date(2026, 10, 7))["schedule"]
            # 2026-10-11: 근무표에 없는 명단 인원은 연차일정에 있을 때만 연차로 넣는다(추측해서 연차로 만들지 않음)
            self.assertNotIn("성욱", {p.name for p in schedule.people})


class DefaultStatusUiTests(unittest.TestCase):
    def test_night_and_day_defaults_pc_and_mobile(self):
        for view in ("pc", "mobile"):
            with self.subTest(view=view), tempfile.TemporaryDirectory() as folder, \
                    patch.dict(os.environ, {"NIGHT_SHIFT_HISTORY_DIR": folder}):
                leave_schedule.save_text(Path(folder), LEAVE)
                at = AppTest.from_file(APP, default_timeout=120)
                at.query_params["view"] = view
                at.run()
                at.session_state["shift_date"] = date(2026, 10, 3)         # A조 야간
                at.run()
                self.assertFalse(at.exception)
                statuses = at.session_state["night_statuses"]
                self.assertEqual((statuses["5"], statuses["4"]), ("연차", "기타"))
                self.assertEqual(at.session_state["manager_state"], "A 결근 (준호 대행)")
                note = at.text_input(key="night_note_2026-10-03_4")
                self.assertEqual(note.value, "예비군 훈련")
                note.set_value("예비군(오전)").run()
                self.assertEqual(leave_schedule.load_notes(Path(folder))["2026-10-03"]["night"], {"광희": "예비군(오전)"})
                at.session_state["shift_date"] = date(2026, 10, 7)         # A조 주간
                at.run()
                self.assertFalse(at.exception)
                self.assertEqual(at.text_input(key="day_note_2026-10-07_광희").value, "교육")
                if view == "mobile":
                    self.assertEqual(at.session_state["day_statuses_2026-10-07"].get("준형"), "연차")


if __name__ == "__main__":
    unittest.main()
