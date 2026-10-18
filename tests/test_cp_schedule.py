"""CP 근무표 (2026-10-05 사용자 요청): '주1'·'야2' = 그날 주간 CP1·야간 CP2. 날짜를 고르면 현장 양식 이미지
(둘째 줄 CP1이름…, 시간별 CP 칸 CP1.CP3→CP1.CP2→CP2.CP3 반복, 07시 셔틀 CP3)와 주간 화면에 기본 표시."""
import io
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook

from night_shift import cp_schedule

NAMES = {"이덕영": {3: "야2", 4: "야1", 1: "주1"}, "안철룡": {3: "야3", 4: "야2", 1: "주2"},
         "권혁준": {3: "야1", 4: "야3", 1: "주3"}, "김준하": {20: "연차", 3: "주1"}}


def workbook() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet["J1"] = "CP 10월 근무"
    sheet.cell(3, 1, None)
    for d in range(1, 32):
        sheet.cell(3, 2 + d, d)
    row = 5
    for team, name in (("A조", "이덕영"), (None, "안철룡"), (None, "권혁준"), ("B조", "김준하")):
        sheet.cell(row, 1, team)
        sheet.cell(row, 2, name)
        for d, value in NAMES[name].items():
            sheet.cell(row, 2 + d, value)
        row += 1
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


class CpScheduleTests(unittest.TestCase):
    def test_parse_excel_and_store(self):
        entries = cp_schedule.parse_excel(workbook(), 2026)
        self.assertEqual(entries["2026-10-03"]["night"], {"2": "이덕영", "3": "안철룡", "1": "권혁준"})
        self.assertEqual(entries["2026-10-03"]["day"], {"1": "김준하"})
        self.assertNotIn("2026-10-20", entries)                            # 연차 = 근무 없음
        with tempfile.TemporaryDirectory() as folder:
            cp_schedule.save(Path(folder), entries)
            self.assertEqual(cp_schedule.for_day(Path(folder), date(2026, 10, 4), "night")["1"], "이덕영")
        self.assertTrue(cp_schedule.is_cp_workbook(workbook()))

    def test_duplicate_number_rejected(self):
        with self.assertRaises(ValueError):
            cp_schedule.parse_rows(2026, 10, {"가": {1: "야1"}, "나": {1: "야1"}})

    def test_hourly_pattern(self):
        self.assertEqual([cp_schedule.hourly_label(t) for t in range(4)],
                         ["CP1.CP3", "CP1.CP2", "CP2.CP3", "CP1.CP3"])
        self.assertEqual(cp_schedule.hourly_label(13), "CP1.CP2")
        self.assertEqual(cp_schedule.short("이덕영"), "덕영")

    def test_field_image_draws_cp(self):
        from night_shift.models import Mode, NightType
        from night_shift.service import HistorySelection, ShiftRequest, SolveOptions, StaffInput, run_schedule
        from night_shift.schedule_image import render_png
        staff = (StaffInput("A", "통제", True, False, None, True),
                 *(StaffInput(str(p), f"직원{p}", True, True, p, False) for p in range(1, 7)))
        result = run_schedule(ShiftRequest(date(2026, 10, 6), NightType.N1, Mode.FM, staff), HistorySelection(),
                              SolveOptions(rules_version="2.20", seconds=15))
        plain = render_png(result.record)
        with_cp = render_png(result.record, {"1": "권혁준", "2": "이덕영", "3": "안철룡"})
        self.assertNotEqual(plain, with_cp)
        # 원본 양식 엑셀 (2026-10-11): 이미지와 같은 칸 내용, 원본 시트 병합 그대로
        from openpyxl import load_workbook
        from night_shift.field_exporter import TEMPLATE
        from night_shift.schedule_image import field_cells, render_xlsx
        cp = {"1": "권혁준", "2": "이덕영", "3": "안철룡"}
        sheet = load_workbook(io.BytesIO(render_xlsx(result.record, cp))).active
        cells = field_cells(result.record, cp)
        self.assertEqual(sheet.title, "26년10월06일(야)")
        self.assertIsNone(sheet.freeze_panes)                              # 원본처럼 틀 고정 없음
        self.assertEqual(sheet["A3"].value, "2026년 10월 06일 화요일 근무조 : A조")
        self.assertEqual((sheet["F5"].value, sheet["F6"].value, sheet["AT7"].value), ("통제A", "CP1혁준", "셔틀"))
        self.assertEqual(sheet["AO8"].value, "CP1.CP3")
        self.assertEqual(sheet["AT21"].value, "CP3")                       # 07시 셔틀
        cols = ["A", "F", "K", "P", "U", "Z", "AE", "AJ", "AO", "AT", "AY", "BD", "BI"]
        for t, row in enumerate(cells["rows"]):
            got = ["" if sheet[f"{c}{8 + t}"].value is None else str(sheet[f"{c}{8 + t}"].value) for c in cols]
            self.assertEqual(got, row)
        self.assertTrue(sheet["F22"].value.startswith("흡연장 파라솔 - "))
        self.assertEqual({str(m) for m in sheet.merged_cells.ranges},
                         {str(m) for m in load_workbook(TEMPLATE).active.merged_cells.ranges})


if __name__ == "__main__":
    unittest.main()
