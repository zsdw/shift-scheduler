import io,os,tempfile,unittest
from pathlib import Path
from datetime import date
from types import SimpleNamespace
from unittest.mock import patch
from openpyxl import load_workbook
from tests.test_cp_schedule import workbook
from night_shift import cp_schedule as cp,order_source

class CpUploadTests(unittest.TestCase):
    def test_month_replace_empty_dates_and_reference_override(self):
        with tempfile.TemporaryDirectory() as folder:
            store=Path(folder)
            cp.save(store,{'2026-09-01':{'day':{'1':'이덕영'},'night':{}},'2026-10-20':{'day':{},'night':{'1':'과거이름'}}})
            raw=workbook();cp.save_uploaded(store,raw,'cp.xlsx',2026)
            self.assertEqual(cp.for_day(store,date(2026,10,20),'night'),{})
            self.assertTrue(cp.for_day(store,date(2026,9,1),'day'))
            sources={'night':{'cp':{'1':'이전'},'cells':{'18':{'CP':'CP1.CP2','셔틀':'CP3'}}}}
            result=order_source.cp_sources(sources,store,date(2026,10,3))
            self.assertEqual(result['night']['cp']['1'],'권혁준')
            self.assertEqual(result['night']['cells']['18']['CP'],'')
            empty=order_source.cp_sources(sources,store,date(2026,10,20))
            self.assertEqual(empty['night']['cp'],{})
            self.assertEqual(empty['night']['cells']['18']['셔틀'],'')
            self.assertEqual(sources['night']['cp'],{'1':'이전'})
            book=load_workbook(io.BytesIO(raw));sheet=book.active
            for row in sheet.iter_rows(min_row=5):
                for cell in row[2:]:cell.value='연차'
            output=io.BytesIO();book.save(output);book.close()
            cp.save_uploaded(store,output.getvalue(),'empty.xlsx',2026)
            self.assertEqual(cp.for_day(store,date(2026,10,3),'night'),{})
            self.assertEqual(len(list((store/'source_files').iterdir())),1)

    def test_leave_tab_upload_pc_mobile_and_reject_bad_file(self):
        from streamlit.testing.v1 import AppTest
        raw=workbook()
        for mobile in (False,True):
            with self.subTest(mobile=mobile),tempfile.TemporaryDirectory() as folder,patch.dict(os.environ,{'NIGHT_SHIFT_HISTORY_DIR':folder}),patch('mobile_ui.is_mobile',return_value=mobile),patch('streamlit.file_uploader',return_value=SimpleNamespace(name='CP.xlsx',getvalue=lambda:raw)):
                at=AppTest.from_string('from history_ui import store_dir\nfrom leave_view_ui import upload_box\nupload_box(store_dir(),"test_cp")',default_timeout=30).run()
                at.number_input(key='cp_year').set_value(2026).run()
                at.button(key='cp_save').click().run()
                self.assertFalse(at.exception,[e.value for e in at.exception])
                store=Path(folder)
                self.assertEqual(cp.for_day(store,date(2026,10,4),'night')['1'],'이덕영')
                before=(store/cp.CP_FILE).read_bytes()
                with self.assertRaises(Exception):cp.save_uploaded(store,b'bad','bad.xlsx',2026)
                self.assertEqual((store/cp.CP_FILE).read_bytes(),before)

if __name__=='__main__':unittest.main()
