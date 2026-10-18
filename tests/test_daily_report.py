"""근태관리 일보 자동 작성 (2026-10-06 사용자 요청·확정 기준)."""
import unittest
from datetime import date
from io import BytesIO

import openpyxl

from night_shift import daily_report as dr
from night_shift.leave_schedule import parse

ROSTER = [dict(r) for r in dr.DEFAULT_ROSTER]


def rows(report):
    return {r["구분"].split(" ")[0]: r for r in report["rows"]}


class DailyReportTests(unittest.TestCase):
    def test_rotation_follows_a_team_calendar(self):
        self.assertEqual(dr.teams_for(date(2026, 10, 7)), {"주간조": "A조", "야간조": "B조", "휴무조": "C조"})
        self.assertEqual(dr.teams_for(date(2026, 9, 10)), {"주간조": "C조", "야간조": "A조", "휴무조": "B조"})
        self.assertEqual(dr.teams_for(date(2026, 9, 12)), {"주간조": "B조", "야간조": "C조", "휴무조": "A조"})

    def test_oct7_matches_hand_made_report(self):
        entries = [{"team": "D", "name": "최정식", "kind": "연차"}, {"team": "A", "name": "이성호", "kind": "연차"},
                   {"team": "", "name": "김은예", "kind": "연차"}]
        r = rows(dr.build(date(2026, 10, 7), entries, ROSTER))
        self.assertEqual((r["주간고정"]["사고"], r["주간고정"]["비고"]), (1, "연차 1명(최정식)"))
        self.assertEqual((r["주간조"]["사고"], r["주간조"]["비고"]), (1, "연차 1명(이성호)"))
        self.assertEqual((r["야간조"]["사고"], r["휴무조"]["사고"]), (0, 10))
        self.assertEqual((r["안내"]["사고"], r["주차"]["사고"]), (1, 0))
        # 작성자 = 전날 밤 야간조 조장 (원본 대조: 10/5 일보 = 10/4 A조 두 번째 야간 백원석)
        self.assertEqual(dr.build(date(2026, 10, 7), entries, ROSTER)["writer"], "이채용")   # 10/6 야간 C조
        self.assertEqual(dr.build(date(2026, 10, 5), [], ROSTER)["writer"], "백원석")        # 10/4 야간 A조
        self.assertEqual(dr.build(date(2026, 10, 2), [], ROSTER)["writer"], "김현")          # 10/1 야간 B조
        absent = dr.build(date(2026, 9, 27), [], ROSTER, [{"team": "B", "name": "김현", "kind": "연차"}])
        self.assertEqual(absent["writer"], "")                                               # 조장 연차 → 확인 필요
        self.assertTrue(any("작성자" in c for c in absent["checks"]))

    def test_attended_not_counted_support_counts_as_present(self):
        # 반차·교육은 출근(비고만), 야간지원은 원래 소속에서 출근 — 사고 아님 (사용자 확정)
        entries = [{"team": "A", "name": "한만년", "kind": "오후 반차"}, {"team": "D", "name": "윤여진", "kind": "신임경비교육"},
                   {"team": "D", "name": "최정식", "kind": "A조 야간지원"}]
        r = rows(dr.build(date(2026, 10, 1), entries, ROSTER))
        self.assertEqual(r["주간고정"]["사고"], 0)
        self.assertIn("야간지원 1명(최정식)", r["주간고정"]["비고"])
        self.assertIn("교육 1명(윤여진)", r["주간고정"]["비고"])
        self.assertEqual(r["주간조"]["사고"], 0)
        self.assertIn("오후반차 1명(한만년)", r["주간조"]["비고"])

    def test_note_wording_follows_original_report(self):
        # 2026-10-06 사용자 요청: 사고내용(비고)을 원본 일보 기재 양식대로
        entries = [{"team": "D", "name": "김수환", "kind": "연차"}, {"team": "D", "name": "최정식", "kind": "연차"},
                   {"team": "D", "name": "최찬우", "kind": "연차"}]
        r = rows(dr.build(date(2026, 9, 10), entries, ROSTER))
        self.assertEqual(r["주간고정"]["비고"], "결원 1명, 연차 3명(최찬우 이사님, 최정식, 김수환)")   # 명단 순서·호칭
        entries = [{"team": "C", "name": "송승현", "kind": "연차"}, {"team": "C", "name": "허승회", "kind": "후반야 반차"},
                   {"team": "C", "name": "박정호", "kind": "연차 및 퇴사"}]
        r = rows(dr.build(date(2026, 9, 18), entries, ROSTER))
        self.assertEqual(r["야간조"]["비고"], "연차 2명(박정호, 송승현), 반차 1명(허승회), 퇴사 1명(박정호)")
        self.assertEqual(dr.classify("전반부 반차"), [("present", "전반야 반차")])
        self.assertEqual(dr.classify("하계휴가"), [("absent", "휴가")])
        self.assertEqual(dr.classify("예비군 훈련"), [("absent", "훈련")])
        self.assertEqual(dr.classify("입사 및 주간조 교육"), [("present", "입사 및 주간교육")])
        self.assertEqual(dr.classify("신임경비교육"), [("present", "신임경비교육")])

    def test_weekend_fixed_staff_off_except_support(self):
        entries = [{"team": "D", "name": "김성덕", "kind": "A조 주간지원"}]
        r = rows(dr.build(date(2026, 9, 26), entries, ROSTER))
        self.assertEqual(r["주간고정"]["사고"], 10)                       # 11 − 주간지원 1
        self.assertIn("결원 1명", r["주간고정"]["비고"])                  # 9/26은 윤여진 입사 전
        self.assertEqual((r["안내"]["사고"], r["주차"]["사고"]), (5, 3))

    def test_vacancy_from_roster_dates_and_unknown_name(self):
        report = dr.build(date(2026, 9, 17), [{"team": "C", "name": "홍길동", "kind": "연차"}], ROSTER)
        r = rows(report)
        self.assertEqual(r["주간고정"]["사고"], 1)                       # 결원 1 (윤여진 입사 전)
        self.assertTrue(any("홍길동" in c for c in report["checks"]))

    def test_excel_fills_template(self):
        report = dr.build(date(2026, 10, 7), [{"team": "A", "name": "이성호", "kind": "연차"}], ROSTER)
        ws = openpyxl.load_workbook(BytesIO(dr.to_excel(report))).active
        self.assertEqual(ws.title, "26.10.07")
        self.assertEqual(ws["G6"].value, "2026년 10월 07일(수)")
        self.assertEqual(ws["M6"].value, "이 채 용 (인)")                 # 10/6 야간조(C) 조장
        self.assertEqual((ws["G13"].value, ws["H13"].value, ws["J13"].value), (10, 1, "연차 1명(이성호)"))
        self.assertEqual(ws["I13"].value, "=IMSUB(G13,H13)")               # 원본 수식 유지

    def test_sheet_view_matches_excel_formulas(self):
        # 2026-10-06 사용자 요청: 화면도 실제 엑셀과 같은 모양, 계·소계·합계는 원본 수식과 같게
        report = dr.build(date(2026, 10, 7), [{"team": "A", "name": "이성호", "kind": "연차"},
                                               {"team": "", "name": "김은예", "kind": "연차"}], ROSTER)
        v = dr.sheet_values(report)
        self.assertEqual(v[13][:3], (10, 1, 9))
        self.assertEqual(v[14][:3], (21, 1, 20))
        self.assertEqual(v[17][:3], (41, 11, 30))
        self.assertEqual(v[24][:3], (49, 12, 37))          # 소계 + 안내 + 주차 (위 주차·정산·아래 안내 12는 제외)
        html = dr.sheet_html(report)
        for text in ("H D S", "사 업 자 명", "2026년 10월 07일(수)", "이 채 용 (인)", "합 계", "연차 1명(이성호)"):
            self.assertIn(text, html)
        self.assertNotIn("관리/시스템", html)                # 화면에서 뺌 (2026-10-06 사용자 요청)
        # 사고내용 칸을 세로로 나눠 계산 근거 (2026-10-06 사용자 요청)
        r = rows(report)
        self.assertEqual(r["주간조"]["근거"], "연차 1 = 1")
        self.assertEqual(r["야간조"]["근거"], "빠진 사람 없음 = 0")
        self.assertEqual(r["휴무조"]["근거"], "휴무조 전원 10")
        self.assertIn("사고 계산 근거", html)
        self.assertIn("소계 + 안내 + 주차", html)
        weekend = rows(dr.build(date(2026, 9, 26), [{"team": "D", "name": "김성덕", "kind": "A조 주간지원"}], ROSTER))
        self.assertEqual(weekend["주간고정"]["근거"], "주말·공휴일 휴무 전원 11 − 출근 1(지원·대근·교육) = 10")
        # 휴무조인데 교육으로 출근하면 사고에서 뺌 (원본 7/1~3·8/26: 휴무조 사고 9)
        off = rows(dr.build(date(2026, 10, 7), [{"team": "C", "name": "이지훈", "kind": "주간조 교육"}], ROSTER))
        self.assertEqual(off["휴무조"]["사고"], 9)
        edited = [dict(x, 사고=0) if x["행"] == 13 else x for x in report["rows"]]
        self.assertIn("직접 고침 (계산값 1)", dr.sheet_html(report, edited))
        members = dr.team_members(date(2026, 10, 7), ROSTER)
        self.assertEqual(members["B조"]["조장"], "김현")
        self.assertIn("이지훈", members["C조"]["대원"])
        self.assertEqual(members["C조"]["CP"], ["이석산", "조백호", "송승현"])
        self.assertIn("윤여진", members["주간고정"]["대원"])

    def test_leave_parser_reads_titles_with_nim(self):
        parsed = parse("09월 26일(토)\n\nD조 김성덕 팀장님 - A조 주간지원\n")
        self.assertEqual(parsed[date(2026, 9, 26)][0]["name"], "김성덕")


class SavedReportTests(unittest.TestCase):
    def test_save_load_overlay_and_trash(self):
        # 2026-10-06 사용자 요청: 고친 사고·비고를 날짜별 저장 (근무명령서가 이 값을 씀)
        import tempfile
        from pathlib import Path
        day = date(2026, 10, 7)
        report = dr.build(day, [{"team": "A", "name": "이성호", "kind": "연차"}], ROSTER)
        with tempfile.TemporaryDirectory() as store:
            self.assertIsNone(dr.load_saved(store, day))
            edited = [dict(r, 사고=2, 비고="연차 1명(이성호), 휴무 1명(최정식)") if r["행"] == 12 else dict(r)
                      for r in report["rows"]]
            dr.save_report(store, day, edited, "이채용", report["rows"])
            saved = dr.load_saved(store, day)
            self.assertEqual(saved["writer"], "이채용")
            rows, changed = dr.apply_saved(report, saved)
            self.assertEqual(next(r for r in rows if r["행"] == 12)["사고"], 2)
            self.assertEqual(changed, [])
            later = dr.build(day, [{"team": "A", "name": "이성호", "kind": "연차"},
                                   {"team": "B", "name": "김현", "kind": "연차"}], ROSTER)       # 연차일정이 바뀜
            self.assertEqual(dr.apply_saved(later, saved)[1], ["야간조 (B조)"])
            dr.save_report(store, day, rows, "이채용", report["rows"])                        # 덮어쓰기 → 휴지통
            dr.delete_saved(store, day)
            self.assertIsNone(dr.load_saved(store, day))
            self.assertEqual(len(list((Path(store) / "daily_reports" / ".trash").glob("*.json"))), 2)


class WritingDayTests(unittest.TestCase):
    def test_default_day_is_report_written_tonight(self):
        # D일 일보는 D-1일 밤에 씀: 저녁·낮이면 내일, 새벽(야간 근무 중)이면 오늘
        from datetime import datetime
        from daily_report_ui import writing_day
        self.assertEqual(writing_day(datetime(2026, 10, 6, 22, 0)), date(2026, 10, 7))
        self.assertEqual(writing_day(datetime(2026, 10, 7, 3, 0)), date(2026, 10, 7))
        self.assertEqual(writing_day(datetime(2026, 10, 7, 10, 0)), date(2026, 10, 8))


class DailyReportPageTests(unittest.TestCase):
    def test_separate_tab_page_renders_and_downloads(self):
        # 2026-10-06 사용자 요청: 근무표와 분리된 맨 위 '근태관리 일보' 탭
        import os, tempfile
        from streamlit.testing.v1 import AppTest
        with tempfile.TemporaryDirectory() as store:
            (__import__("pathlib").Path(store) / "leave_schedule.txt").write_text(
                "10월 07일(수)\n\nA조 이성호 사원 - 연차\n", encoding="utf-8")
            os.environ["NIGHT_SHIFT_HISTORY_DIR"] = store
            from unittest.mock import patch
            try:   # 2026-10-11부터 일보는 편성표 자료가 있는 날만 → 빈 검사 저장소에서는 자료 있는 날로 고정
                with patch("night_shift.order_source.document_available", return_value=True), \
                        patch("night_shift.order_source.latest_file", return_value=None):   # 바탕화면 실제 편성표 안 읽음
                    app = AppTest.from_string(
                        "from datetime import date\nimport streamlit as st\nimport daily_report_ui\n"
                        "st.session_state.setdefault('daily_report_day', date(2026, 10, 7))\ndaily_report_ui.page()\n").run()
            finally:
                os.environ.pop("NIGHT_SHIFT_HISTORY_DIR", None)
            self.assertFalse(app.exception)
            self.assertEqual(app.title[0].value, "근태관리 일보")
            self.assertTrue(any("연차 1명(이성호)" in m.value for m in app.markdown))
            self.assertFalse(any("일보 저장" in b.label for b in app.button))
            members = next(d.value for d in app.dataframe if "조장" in d.value.columns)   # 조별 근무자 칸
            self.assertEqual(list(members["조"].str[:4]), ["주간고정", "A조 (", "B조 (", "C조 ("])


class LeaveViewTests(unittest.TestCase):
    def test_leave_tab_shows_month_and_people(self):
        # 2026-10-06 사용자 요청: 연차일정을 항상 한눈에 보는 탭 (PC 달력)
        import os, tempfile
        from pathlib import Path
        from streamlit.testing.v1 import AppTest
        import leave_view_ui
        self.assertEqual(leave_view_ui.tone("오후 반차"), "반차")
        self.assertEqual(leave_view_ui.tone("A조 야간지원"), "지원")
        self.assertEqual(leave_view_ui.tone("연차"), "사고")
        with tempfile.TemporaryDirectory() as store:
            Path(store, "leave_schedule.txt").write_text("10월 07일(수)\n\nA조 이성호 사원 - 연차\n"
                                                         "10월 09일(금)\n\n주차 이정민 사원 - 오후 반차\n", encoding="utf-8")
            os.environ["NIGHT_SHIFT_HISTORY_DIR"] = store
            try:
                app = AppTest.from_string("import streamlit as st\nimport leave_view_ui\n"
                                          "st.session_state.setdefault('leave_month', (2026, 10))\n"
                                          "leave_view_ui.page()\n").run()
            finally:
                os.environ.pop("NIGHT_SHIFT_HISTORY_DIR", None)
            self.assertFalse(app.exception)
            html = " ".join(m.value for m in app.markdown)
            self.assertIn("이성호 연차", html)
            self.assertIn("이정민 오후 반차", html)
            people = app.dataframe[0].value
            self.assertEqual(set(people["이름"]), {"이성호", "이정민"})


if __name__ == "__main__":
    unittest.main()
