"""2026-10-21 사용자 사례: 편성표 명단 칸 '8.세명'(번호 뒤 점)을 못 읽어 이세명이 근태일보에서 미출근으로 나왔다."""
import unittest
from io import BytesIO

import openpyxl

from night_shift import bulk_import, order_source


def book() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "26년10월21일(주)"
    ws["A3"] = "2026년 10월 21일 수요일 근무조 : C조"
    ws["A5"], ws["E5"], ws["J5"] = "근무자명", "7석규", "8.세명"
    ws["O5"], ws["T5"] = "9 용주", "10,여진"
    ws["A7"], ws["E7"] = "근무시간", "G1"
    ws["A8"], ws["E8"] = "08:00", 8
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


class RosterDotNameTests(unittest.TestCase):
    def test_bulk_import_names(self):
        numbered, _ = bulk_import._names([["7석규", "8.세명", "9·용주", "10-여진", "11 지훈"]])
        self.assertEqual(numbered, {"7": "석규", "8": "세명", "9": "용주", "10": "여진", "11": "지훈"})

    def test_order_source_names(self):
        names = order_source.parse_uploaded(book(), "x.xlsx")["2026-10-21"]["day"]["names"]
        self.assertEqual(names, {"7": "석규", "8": "세명", "9": "용주", "10": "여진"})


if __name__ == "__main__":
    unittest.main()
