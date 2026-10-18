import io
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch
import openpyxl
from night_shift import daily_report as dr
from night_shift import order_source as source
from night_shift import security_order as order


class SecurityOrderTests(unittest.TestCase):
    def test_security_count_adjust(self):
        # 2026-10-13 사용자 확정: 다른 조 출근자는 보안에 더한다. 야간지원 간 사람도 출근이라 빼지 않는다. CP는 보안에 넣지 않음.
        row=lambda to,acc,note='':{'T/O':to,'사고':acc,'비고':note}
        # 원본 7/1: 야간조 '입사 1명 및 주간교육 1명(박정호)' + 휴무조 교육 1명 → +2 (11+9-4+2=18)
        rm={12:row(11,0),13:row(10,1),15:row(10,0,'입사 1명 및 주간교육 1명(박정호)'),16:row(10,9,'결원 1명, A조 편입 및 교육 1명(박인철)')}
        self.assertEqual(order.security_adjust(rm),2)
        # 주간고정 최정식 야간조 지원은 출근(사고 아님) → 빼지 않음, 휴무조 야간지원도 출근 → +1
        rm={12:row(11,2,'결원 1명, 연차 1명(최찬우), 야간조 지원 1명(최정식)'),13:row(10,1),15:row(10,0),16:row(10,9,'야간지원 1명(홍길동)')}
        self.assertEqual(order.security_adjust(rm),1)

    def test_pc_mobile_page_and_date_change(self):
        from streamlit.testing.v1 import AppTest
        root=Path(__file__).resolve().parents[1]
        for mobile in (False,True):
            at=AppTest.from_file(str(root/'app.py'),default_timeout=15)
            at.session_state['app_tab']='📝 보안근무명령서'
            at.session_state['order_day']=date(2026,10,7)
            with patch('mobile_ui.is_mobile',return_value=mobile):
                at.run()
                self.assertEqual(len(at.exception),0)
                # 저장된 프로그램 근무표가 있으면 그것이 주간 기준 (2026-10-07 사용자 확정), 없으면 편성표 시트
                self.assertTrue(any('A조' in m.value and ('26년10월07' in m.value or '프로그램 근무표' in m.value) for m in at.caption))
                # 자료끼리 다르면(2026-10-11부터 모든 날짜 대조) 기준 선택이 먼저, 같으면 미리보기·다운로드
                preview=any('보안팀 근무 명령서' in m.value for m in at.markdown)
                self.assertTrue(preview or at.radio or any('기준' in b.label for b in at.button))
                if preview:self.assertEqual(len(at.get('download_button')),1)
                at.button(key='order_next').click().run()
                self.assertEqual(len(at.exception),0)
                self.assertEqual(at.session_state['order_day'],date(2026,10,8))

    def test_all_teams_duplicate_and_raw_cells(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'편성표.xlsx'
            w=openpyxl.Workbook(); w.remove(w.active)
            for title,number in [('26년09월22일(주)',1),('26년09월22일(주) 수정',2),('26년09월22일(야)',3)]:
                s=w.create_sheet(title); s['A3']='2026년 9월 22일 근무조 : C조'
                s['F5']='1석규'; s['F6']='CP1백호'
                s.append([]); s['A7']='근무시간'; s['F7']='G1'; s['K7']='휴 게'
                s['A8']='08:00' if '(주)' in title else '18:00'; s['F8']=number
            w.save(path)
            result=source.read_day(path,date(2026,9,22),Path(folder))
            self.assertEqual(result['day']['team'],'C')
            self.assertEqual(result['day']['cells']['8']['G1'],'2')
            self.assertEqual(result['night']['cells']['18']['G1'],'3')
            self.assertEqual(result['day']['cp'],{'1':'백호'})

    def test_live_report_ignores_archive_and_missing_source_blank(self):
        with tempfile.TemporaryDirectory() as folder:
            store=Path(folder); day=date(2026,10,7)
            r=dr.build(day,[],dr.load_roster(store)); rows=r['rows']
            rows[0]['사고']=2; rows[0]['비고']='야간지원 1명(최정식)'
            dr.save_report(store,day,rows,'작성자')
            b,s,warnings=order.build(day,store,{})
            self.assertIsNone(s['G6'].value)
            self.assertIsNone(s['C6'].value)
            self.assertIsNone(s['C45'].value)
            self.assertNotIn('야간조 지원 1명(최정식)',s['F34'].value)
            from night_shift.leave_schedule import save_manual
            save_manual(store, '최정식', 'D', day, day, '야간지원')
            b,s,warnings=order.build(day,store,{})
            self.assertIn('야간조 지원 1명(최정식)',s['F34'].value)
            self.assertIsNone(s['G49'].value)
            self.assertEqual(s['C2'].value,'김  성  덕 (D조)')
            self.assertEqual(len([x for x in warnings if '편성표 없음' in x]),2)
            exported=openpyxl.load_workbook(io.BytesIO(order.excel_bytes(b)))
            self.assertEqual(exported.active['F34'].value,s['F34'].value)
            self.assertNotIn('백원석',order.sheet_html(s))
            # 안내의 반시간 고정값은 병합/테두리 보정으로 바뀌지 않는다.
            self.assertEqual(exported.active['K31'].value,'1')
            self.assertEqual(exported.active['L31'].value,'2')

    def test_entire_original_frame_preserved_after_export(self):
        template=openpyxl.load_workbook(order.TEMPLATE)
        with tempfile.TemporaryDirectory() as folder:
            for day in (date(2026,10,7),date(2026,10,3)):
                original=template['휴일' if dr.is_off_day(day) else '평일']
                for sources in ({},{'day':{'team':dr.teams_for(day)['주간조'][0], 'names':{'1':'제욱'},
                       'cp':{'1':'혁준','2':'덕영','3':'철룡'},'controls':['A'],
                       'cells':{'8':{'G1':'1'},'10':{'대기':'5','휴게':'6'}}}}):
                    book,_,_=order.build(day,Path(folder),sources)
                    generated=openpyxl.load_workbook(io.BytesIO(order.excel_bytes(book))).active
                    self.assertEqual(set(map(str,generated.merged_cells.ranges)),set(map(str,original.merged_cells.ranges)))
                    self.assertEqual((generated.max_row,generated.max_column),(original.max_row,original.max_column))
                    self.assertEqual(str(generated.print_area).split('!')[-1],str(original.print_area).split('!')[-1])
                    for row in original:
                        for cell in row:
                            actual=generated[cell.coordinate]
                            self.assertEqual(cell._style,actual._style,(day,cell.coordinate,'전체 서식'))
                            self.assertEqual(type(cell),type(actual),(day,cell.coordinate,'병합 상태'))
                    self.assertEqual({k:v.height for k,v in original.row_dimensions.items()},
                                     {k:v.height for k,v in generated.row_dimensions.items()})
                    self.assertEqual({k:v.width for k,v in original.column_dimensions.items()},
                                     {k:v.width for k,v in generated.column_dimensions.items()})
                    if sources and not dr.is_off_day(day):
                        self.assertEqual(generated['O18'].value,'대기: 5\n휴게: 6')

    def test_source_transfer_weekday_and_holiday(self):
        with tempfile.TemporaryDirectory() as folder:
            store=Path(folder)
            for day in [date(2026,10,7),date(2026,10,3)]:
                s={'day':day.isoformat(),'shift':'day','team':dr.teams_for(day)['주간조'][0],
                   'names':{'1':'제욱'},'cp':{'1':'혁준','2':'덕영','3':'철룡'},'controls':['A'],
                   'cells':{'8':{'G1':'1','셔틀':'5(해지)','순찰':'CS교육:1'},'17':{'G1':'A'}}}
                b,ws,_=order.build(day,store,{'day':s})
                self.assertEqual(ws['K6'].value,'1')
                target=17 if not dr.is_off_day(day) else 13
                self.assertEqual(ws.cell(target,11).value,'1')
                self.assertEqual(ws['C2'].value,'김  성  덕 (D조)')
                self.assertIn('김제욱',[c.value for row in ws for c in row if c.column==3])

    def test_folder_leave_updates_and_preserves_previous(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder); store=root/'history'; store.mkdir()
            (store/'leave_schedule.txt').write_text('old',encoding='utf-8')
            p=root/'연차일정(2026).txt'; p.write_text('10월 7일\n최정식 파트장 - 야간지원',encoding='utf-8')
            source.sync_leave(store,root)
            self.assertEqual(dr.entries_for(store,date(2026,10,7))[0]['kind'],'야간지원')
            self.assertEqual(len(list((store/'.trash').glob('*.txt'))),1)
            source.sync_leave(store,root)
            self.assertEqual(len(list((store/'.trash').glob('*.txt'))),1)


class AccidentWordingTests(unittest.TestCase):
    def test_original_order_wording(self):
        from night_shift.security_order import order_kind
        self.assertEqual(order_kind('야간지원'), '야간조 지원')
        self.assertEqual(order_kind('예비군 훈련'), '훈련')
        self.assertEqual(order_kind('연차'), '연차')


class OrderNameTests(unittest.TestCase):
    def test_absent_control_and_duplicate_number_fixed(self):
        from night_shift.leave_schedule import save_text
        with tempfile.TemporaryDirectory() as folder:
            store=Path(folder); day=date(2026,10,8)
            save_text(store,'10월 08일(목)\nD조 최정식 파트장 - 연차\n')
            src={'team':'A','controls':['D','E','A'],'names':{'1':'만년','2':'광희','6':'만년'},'cp':{},'cells':{},'sheet':'x'}
            b,s,warnings=order.build(day,store,{'day':src})
            names=[s.cell(r,3).value for r in range(6,24) if s.cell(r,3).value]
            self.assertNotIn('최정식',names)                       # 연차인 통제E는 성명에서 뺀다
            self.assertTrue(any('통제E 최정식' in w for w in warnings))
            self.assertTrue(any("1번·6번 이름이 모두 '만년'" in w for w in warnings))
            source.save_name_fix(store,day,'day',{'6':'여진'})
            fixed=source.apply_name_fixes({'day':src},store,day)
            b,s,warnings=order.build(day,store,fixed)
            row={s.cell(r,4).value:s.cell(r,3).value for r in range(6,24)}
            self.assertEqual(row['대원(6)'],'윤여진')
            self.assertEqual([n for n in [s.cell(r,3).value for r in range(6,24)] if n=='윤여진'],['윤여진'])


class ProgramSourceTests(unittest.TestCase):
    def test_saved_program_schedule_overrides_excel(self):
        from night_shift.day_shift import DayPerson, DaySchedule, save_day
        with tempfile.TemporaryDirectory() as folder:
            store=Path(folder); day=date(2026,10,8)
            people=tuple(DayPerson(n,n,'출근',i+1) for i,n in enumerate(('만년','광희','제욱','원종','도운','성욱','용주','준호','여진')))
            grid={p.pid:['G1' if p.number==1 else '휴게']*10 for p in people}
            save_day(store/'day',DaySchedule(day,people,grid))
            source.save_name_fix(store,day,'day',{'6':'여진'})          # 엑셀용 이름 고침은 프로그램 근무표에 쓰지 않음
            program=source.saved_day_source(store,day)
            self.assertTrue(program['program'])
            self.assertEqual(program['names']['6'],'성욱')
            self.assertEqual(program['cells']['8']['G1'],'1')
            with patch.object(source,'latest_file',return_value=None):
                found=source.sources_for(store,day)
            self.assertEqual(found['day']['names']['6'],'성욱')
            b,s,warnings=order.build(day,store,found)
            names=[s.cell(r,3).value for r in range(6,24) if s.cell(r,3).value]
            self.assertIn('진기운',names)                                # 통제D는 출근이면 성명에 포함
            self.assertEqual(s['K6'].value,'1')                          # 08시(K열) 1층(G1) = 1번
