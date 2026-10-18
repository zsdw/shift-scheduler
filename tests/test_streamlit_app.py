import os
import tempfile

# Keep UI tests independent of the user's confirmed history folder.
os.environ["NIGHT_SHIFT_HISTORY_DIR"] = tempfile.mkdtemp(prefix="night_shift_history_")

import unittest
from dataclasses import replace
from datetime import date
from io import BytesIO
from itertools import combinations
from pathlib import Path
from unittest.mock import patch

from streamlit.testing.v1 import AppTest
from openpyxl import load_workbook

from app import exclusion_payload, exclusion_payloads, staff_for_input
from night_shift.excel_exporter import export_xlsx
from night_shift.field_exporter import export_field_xlsx
from night_shift.models import AssignmentKind
from night_shift.presentation import assignment_view


class StreamlitAppTests(unittest.TestCase):
    @staticmethod
    def schedule_identity(result):
        """Identity ignores display numbers and UI row order."""
        return tuple(sorted(
            (row.employee_id, row.slot.index,
             row.displayed_assignment.kind.value, row.displayed_assignment.label,
             row.actual_post_or_work_type.kind.value, row.actual_post_or_work_type.label,
             row.posture_or_mode.value if row.posture_or_mode else None,
             row.actual_work_time_value, row.actual_rest_time_value, row.fatigue_value,
             row.counts_for_continuous_work_restriction,
             row.barrier_install_flag, row.barrier_remove_flag)
            for row in result.record.assignments
        ))

    @staticmethod
    def schedule_structure(result):
        """Compare the grid itself; permuting names/numbers does not change it."""
        def row_identity(employee_id):
            return tuple(
                (row.slot.index, row.displayed_assignment.kind.value,
                 row.displayed_assignment.label,
                 row.actual_post_or_work_type.kind.value,
                 row.actual_post_or_work_type.label,
                 row.posture_or_mode.value if row.posture_or_mode else None,
                 row.actual_work_time_value, row.actual_rest_time_value,
                 row.fatigue_value, row.counts_for_continuous_work_restriction,
                 row.barrier_install_flag, row.barrier_remove_flag)
                for row in sorted(
                    (item for item in result.record.assignments
                     if item.employee_id == employee_id),
                    key=lambda item: item.slot.index)
            )
        manager_id = next(p.employee_id for p in result.employee_stats.values() if p.is_manager)
        general_rows = sorted(
            (row_identity(p.employee_id)
             for p in result.employee_stats.values() if not p.is_manager),
            key=repr,
        )
        return row_identity(manager_id), tuple(general_rows)

    @staticmethod
    def materially_changed_people(left, right):
        def by_person(result):
            rows = {}
            for row in result.record.assignments:
                rows[row.employee_id, row.slot.index] = (
                    row.displayed_assignment.kind.value,
                    row.displayed_assignment.label,
                    row.actual_post_or_work_type.kind.value,
                    row.actual_post_or_work_type.label,
                    row.posture_or_mode.value if row.posture_or_mode else None,
                    row.actual_work_time_value, row.actual_rest_time_value,
                    row.fatigue_value, row.counts_for_continuous_work_restriction,
                    row.barrier_install_flag, row.barrier_remove_flag,
                )
            return rows
        left_rows, right_rows = by_person(left), by_person(right)
        employee_ids = {employee_id for employee_id, _ in left_rows}
        return sum(
            sum(left_rows[employee_id, slot] != right_rows[employee_id, slot]
                for slot in range(14)) >= 3
            for employee_id in employee_ids
        )

    def test_day_shift_date_only_shows_notice(self):
        app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "app.py"),
                                default_timeout=90).run()
        app.session_state["shift_date"] = date(2026, 9, 1)
        app.run()   # 평일 주간 -> 주간 근무표 화면
        self.assertFalse(app.exception)
        self.assertIn("day_generate", [button.key for button in app.button])
        self.assertNotIn("generate_button", [button.key for button in app.button])
        app.session_state["shift_date"] = date(2026, 9, 13)  # 일요일 주간 -> 안내만
        app.run()
        self.assertTrue(any("주말·공휴일 주간근무" in item.value for item in app.info))
        self.assertNotIn("day_generate", [button.key for button in app.button])
        app.session_state["shift_date"] = date(2026, 9, 5)                    # 휴무
        app.run()
        self.assertTrue(any("A조 휴일" in item.value for item in app.info))
        self.assertNotIn("generate_button", [button.key for button in app.button])

    def test_proxy_manager_mapping(self):
        staff = staff_for_input(False, {str(i): i != 2 for i in range(1, 7)})
        self.assertFalse(staff[0].present)
        self.assertTrue(staff[-1].present)
        self.assertTrue(staff[-1].acts_as_manager)
        self.assertFalse(staff[-1].is_general_worker)
        self.assertEqual(staff[-1].display_number, 6)

    def test_half_day_input_and_controller_label(self):
        presence = {str(i): i != 1 for i in range(1, 7)}
        staff = staff_for_input(
            True, presence, {str(i): i for i in range(1, 7)},
            {"1": "전반부 반차 (01시 출근)"},
        )
        self.assertEqual(staff[0].name, "통제")
        half_day = next(person for person in staff if person.employee_id == "1")
        self.assertTrue(half_day.present)
        self.assertEqual(half_day.available_slot_indices, tuple(range(7, 14)))

    def test_live_v213_flow_and_stale_result_clearing(self):
        app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "app.py"),
                                default_timeout=90).run()
        self.assertFalse(app.exception)
        app.session_state["shift_date"] = date(2026, 9, 28)
        app.run()
        self.assertEqual(app.selectbox(key="night_type_select").value, "N1")   # 월요일 야간 자동
        app.selectbox(key="night_type_select").set_value("N3")
        app.run()
        self.assertTrue(any("미생성" in item.value for item in app.markdown))
        self.assertTrue(app.button(key="generate_button").disabled)        # 5-7: N3는 추가 남문 확인 전 잠금
        app.radio(key="extra_south_2026-09-28_N3").set_value("추가 없음").run()
        app.button(key="generate_button").click().run()
        self.assertFalse(app.exception)
        self.assertTrue(any("미확정" in item.value for item in app.markdown))
        result = app.session_state["last_result"]
        self.assertTrue(result.valid)
        self.assertEqual(result.solve_result.status, "FEASIBLE")
        self.assertEqual(result.validation_errors, ())
        self.assertEqual(len(result.record.assignments), 98)
        self.assertTrue(app.success)
        self.assertEqual([tab.label for tab in app.tabs],
                         ["근무표", "직원 분석", "전체 비교", "직원 상세분석", "검증 결과", "월간 기록"])
        summary_tables = [frame.value for frame in app.dataframe
                          if "실제 업무 연속시간" in frame.value.columns]
        self.assertEqual(len(summary_tables), 1)
        self.assertIn("연속근무 제한 적용시간", summary_tables[0].columns)
        self.assertTrue(any("하드룰 검증 통과" in item.value for item in app.success))
        self.assertTrue(any("소프트룰" in item.value for item in app.subheader))
        visible_numbers = [stats.display_number for stats in result.employee_stats.values()
                           if not stats.is_manager]
        self.assertEqual(sorted(visible_numbers), list(range(1, 7)))
        colors = {assignment_view(row)[2] for row in result.record.assignments
                  if row.displayed_assignment.label in ("G1", "L3", "I1")}
        self.assertGreaterEqual(len(colors), 3)
        # Generated schedules expose the field-layout PNG and the original-form Excel (2026-10-11).
        self.assertEqual([b.label for b in app.download_button],
                         ["🖼 현장 양식 이미지 다운로드", "📊 현장 양식 엑셀 다운로드"])
        field_payload = export_field_xlsx(result)
        self.assertTrue(field_payload.startswith(b"PK"))
        field_book = load_workbook(BytesIO(field_payload), data_only=True)
        self.assertEqual(field_book.sheetnames,
                         ["현장 야간근무표", "직원별 요약", "검증 결과", "이력 분석"])
        self.assertEqual(len(field_book._external_links), 0)
        field = field_book.worksheets[0]
        self.assertEqual(field["A1"].value, "야간 근무 편성표")
        self.assertEqual(field["A8"].value, "18:00")
        self.assertEqual(field["A14"].value, "24:00")
        self.assertEqual(field["AT7"].value, "차단봉")
        self.assertEqual(field["AY7"].value, "순찰")
        self.assertEqual(field["BI7"].value, "휴게")
        self.assertEqual(field.page_setup.fitToWidth, 1)
        self.assertEqual(field.page_setup.orientation, "landscape")
        post_columns = {"G1": "F", "L3": "K", "I1": "P", "G3": "U",
                        "G2": "Z", "G2(1)": "AE", "G2(2)": "AJ"}
        other_columns = {AssignmentKind.PATROL: "AY", AssignmentKind.BARRIER_ONLY: "AT",
                         AssignmentKind.SECURITY_STANDBY: "BD", AssignmentKind.REST: "BI"}
        for row in result.record.assignments:
            person = result.employee_stats[row.employee_id]
            token = str(person.display_number) if person.display_number is not None else "A"
            assignment = row.displayed_assignment
            column = (post_columns[assignment.label] if assignment.kind == AssignmentKind.POST
                      else other_columns[assignment.kind])
            self.assertIn(token, field[f"{column}{8 + row.slot.index}"].value.split("."))
            if row.barrier_install_flag or row.barrier_remove_flag:
                self.assertIn(token, field[f"AT{8 + row.slot.index}"].value.split("."))
        self.assertTrue(all(field[f"AO{row}"].value is None for row in range(8, 22)))
        self.assertEqual(field_book["검증 결과"]["B3"].value, 0)

        payload = export_xlsx(result)
        self.assertTrue(payload.startswith(b"PK"))
        workbook = load_workbook(BytesIO(payload), data_only=True)
        self.assertEqual(workbook.sheetnames,
                         ["근무표", "직원별 요약", "검증 결과", "이력 분석"])
        schedule = workbook["근무표"]
        self.assertEqual(schedule["B2"].value.date(), date(2026, 9, 28))
        self.assertEqual((schedule["E2"].value, schedule["H2"].value, schedule["N2"].value),
                         ("N3", "FM", "v2.21"))   # 엔진 분기값 2.20 = 룰북 v2.21 표기
        self.assertEqual(schedule["B3"].value, "하드룰 검증 통과 · 위반 0건")
        self.assertEqual([schedule.cell(5, col).value for col in range(2, 16)],
                         [f"{(18 + i) % 24:02d}시" for i in range(14)])
        people = sorted(result.employee_stats.values(),
                        key=lambda p: (not p.is_manager, p.display_number or 0))
        assignments = {(row.employee_id, row.slot.index): row for row in result.record.assignments}
        summary = workbook["직원별 요약"]
        for excel_row, person in enumerate(people, start=6):
            expected_label = "통제" if person.employee_id == "A" else f"{person.display_number} {person.name}"
            self.assertEqual(schedule.cell(excel_row, 1).value, expected_label)
            for index in range(14):
                label, notes, _, _ = assignment_view(assignments[person.employee_id, index])
                self.assertEqual(schedule.cell(excel_row, index + 2).value,
                                 label + ("\n" + " · ".join(notes) if notes else ""))
            summary_row = excel_row - 1
            self.assertEqual(summary.cell(summary_row, 2).value, person.actual_work_time_value)
            self.assertEqual(summary.cell(summary_row, 3).value, person.actual_rest_time_value)
            self.assertEqual(summary.cell(summary_row, 4).value, person.fatigue_value)
            self.assertEqual(summary.cell(summary_row, 10).value, person.max_actual_work_run)
            self.assertEqual(summary.cell(summary_row, 12).value, person.max_rest_run)
        verification = workbook["검증 결과"]
        self.assertEqual((verification["B2"].value, verification["B3"].value,
                          verification["B4"].value),
                         ("하드룰 검증 통과", 0, result.solve_result.status))
        self.assertEqual(workbook["이력 분석"]["B2"].value, 0)
        self.assertEqual(schedule.page_setup.fitToWidth, 1)
        self.assertEqual(schedule.page_setup.orientation, "landscape")
        with self.assertRaises(ValueError):
            export_xlsx(replace(result, validation_errors=({"rule": "test", "detail": "invalid"},)))
        with self.assertRaises(ValueError):
            export_field_xlsx(replace(result, validation_errors=({"rule": "test", "detail": "invalid"},)))

        app.button(key="regenerate_0").click().run(timeout=180)
        self.assertFalse(app.exception)
        self.assertEqual(len(app.session_state["schedule_results"]), 2)
        second = app.session_state["schedule_results"][1]
        self.assertNotEqual(self.schedule_identity(result), self.schedule_identity(second))
        self.assertEqual(len(exclusion_payloads([result, second])), 2)
        self.assertNotIn("display_number", exclusion_payload(result)[0])
        self.assertIs(app.session_state["last_result"], second)
        self.assertEqual(app.session_state["selected_schedule_index"], 1)

        # A third click excludes every prior employee-by-slot schedule. Merely
        # rotating visible numbers cannot count as a new solution.
        app.button(key="regenerate_1").click().run(timeout=180)
        self.assertFalse(app.exception)
        generated = app.session_state["schedule_results"]
        self.assertEqual(len(generated), 3)
        self.assertEqual(len({self.schedule_identity(item) for item in generated}), 3)
        self.assertEqual(len({self.schedule_structure(item) for item in generated}), 3)
        for left, right in combinations(generated, 2):
            self.assertGreaterEqual(self.materially_changed_people(left, right), 3)

        # 후보 선택은 이름(key)으로 — 화면에 다른 선택칸(남문 G2 추가 운영 등)이 앞에 생겨 radio[0]이 다른 칸이 됨
        app.radio(key="selected_schedule_index").set_value(0).run()
        self.assertIs(app.session_state["last_result"], result)
        self.assertEqual(app.session_state["selected_schedule_index"], 0)
        selected_field = load_workbook(BytesIO(export_field_xlsx(app.session_state["last_result"])),
                                       data_only=True).worksheets[0]
        self.assertEqual(
            [[cell.value for cell in row] for row in selected_field.iter_rows(max_row=22, max_col=68)],
            [[cell.value for cell in row] for row in field.iter_rows(max_row=22, max_col=68)],
        )

        # 야간 근무자 표(이름·상태·역할): 체크박스 없이 상태만 바꾼다
        self.assertFalse([c for c in app.checkbox if str(c.key).startswith("present_")])
        app.session_state["night_statuses"] = {**app.session_state["night_statuses"],
                                               "1": "전반부 반차 (01시 출근)"}
        app.run()
        self.assertEqual(app.session_state["night_statuses"]["1"], "전반부 반차 (01시 출근)")
        self.assertIsNone(app.session_state["last_result"])
        self.assertFalse(app.success)

        with patch("night_shift.service.run_schedule", side_effect=RuntimeError("시험 실패")):
            app.button(key="generate_button").click().run()
        self.assertFalse(app.exception)
        self.assertIsNone(app.session_state["last_result"])
        self.assertTrue(app.error)


if __name__ == "__main__":
    unittest.main()
