"""휴일 주간 현장 양식 이미지·엑셀 (2026-10-12): 원본 휴일 시트 양식 그대로, 이미지·엑셀 같은 칸."""
import io
import tempfile
import unittest
from pathlib import Path

from night_shift import weekend_form as wf
from night_shift.weekend_day import WeekendSchedule

# 원본 9/19 휴일 시트(통제E 지원, CP 3명)와 같은 근무
GRID = {'1': ['G1', 'I1', '휴게', 'G1', 'G2', '휴게', 'I1', 'L1', '순찰', 'G2'],
        '2': ['휴게', 'G1', 'G2', '휴게', 'G1', 'L1', '순찰', 'G1', 'L1', '휴게'],
        '3': ['G2', '휴게', 'G1', 'G2', '휴게', 'I1', 'L1', '순찰', 'I1', 'L1'],
        '4': ['I1', 'L1', '순찰', 'I1', 'L1', '휴게', 'G1', 'G2', '휴게', 'G1'],
        '5': ['L1', '순찰', 'I1', 'L1', '휴게', 'G1', 'G2', '휴게', 'G1', 'I1'],
        'A': ['휴게', 'G2', 'L1', '순찰', 'I1', 'G2', '휴게', 'I1', 'G2', '휴게']}
NAMES = {'A': '통제E(지원)', '1': '준호', '2': '광희', '3': '원종', '4': '제욱', '5': '만년'}
CP = {'1': '안철룡', '2': '권혁준', '3': '이덕영'}


class WeekendFormTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.store = Path(self.tmp.name)
        self.s = WeekendSchedule(day='2026-09-19', names=NAMES, grid=GRID)

    def tearDown(self):
        self.tmp.cleanup()

    def test_cells_match_original_sheet(self):
        v = wf.build_spec(self.s, self.store, CP)["values"]
        self.assertEqual(v["A3"], "2026년 09월 19일 토요일 근무조 : A조")
        self.assertEqual([v.get(c) for c in ("F5", "K5", "AE5", "F6", "P6")],
                         ["통제E", "1준호", "5만년", "CP1철룡", "CP3덕영"])
        self.assertEqual([v.get(f"{c}8") for c in ("F", "K", "U", "Z", "AJ", "AY")],
                         ["1", "3", "4", "5", "CP1.CP2", "E.2.CP3"])        # 원본 08시 줄
        self.assertEqual([v.get(f"{c}9") for c in ("K", "AJ", "AO", "AY")], ["E", "CP1.CP3", "5", "3.CP2"])

    def test_two_people_on_g2_go_to_g2_and_g3(self):
        grid = {**GRID, '1': ['G2'] + GRID['1'][1:]}
        v = wf.build_spec(WeekendSchedule(day='2026-09-19', names=NAMES, grid=grid), self.store, CP)["values"]
        self.assertEqual((v.get("K8"), v.get("P8")), ("1", "3"))

    def test_image_and_excel(self):
        from PIL import Image
        import openpyxl
        image = Image.open(io.BytesIO(wf.draw(self.s, self.store, CP)))
        with Image.open(wf.BLANK) as blank:
            self.assertEqual(image.size, blank.size)
        ws = openpyxl.load_workbook(io.BytesIO(wf.to_xlsx(self.s, self.store, CP))).worksheets[0]
        self.assertEqual(ws.title, "26년09월19일(주)")
        self.assertEqual((ws["F8"].value, ws["AY8"].value, ws["K5"].value), (1, "E.2.CP3", "1준호"))
        self.assertGreater(len(ws.merged_cells.ranges), 200)                  # 원본 양식 병합 그대로


class CompactSheetTests(unittest.TestCase):
    """폰 주간·휴일 표: 가로 스크롤 없이 (최소 너비 없음), 내용은 PC 표와 같음 (2026-10-12)."""

    def test_weekend_compact(self):
        import re
        from weekend_day_ui import sheet_html
        s = WeekendSchedule(day='2026-09-19', names=NAMES, grid=GRID)
        pc, phone = sheet_html(s), sheet_html(s, compact=True)
        self.assertIn("wsheet dsheet compact", phone)
        self.assertIn("min-width:0", phone)
        strip = lambda h: re.sub(r"<style>.*?</style>|<[^>]+>", "", h)
        self.assertEqual(strip(pc), strip(phone))

    def test_day_compact_breaks_only_text_dots(self):
        from day_ui import _break_dots
        html = "<td style='color:rgba(0,0,0,.3)'>CP1.CP2</td>"
        self.assertEqual(_break_dots(html), "<td style='color:rgba(0,0,0,.3)'>CP1.<wbr>CP2</td>")


if __name__ == "__main__":
    unittest.main()
